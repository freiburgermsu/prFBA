#!/usr/bin/env python
"""Build the prFBA update deck from the repositories' own result files.

Every number on a slide is read from the JSON/CSV the pipeline produced --
``taxacc_region_summary.json``, ``gene_capture_by_region.json``,
``comparison_summary.json``, ``annotation_ceiling.json``,
``depth_probe_results.json``, the MAG arm's ``mag_summary_full.json`` /
``crossover.json`` / ``gtdb_summary_2000.json``, and the two-panel comparison's
``panel_comparison.json`` / ``mag_panel_comparison.json`` / ``space_deciles.json``
-- so the deck cannot drift from the data it reports.

The one exception is the block of constants below: a handful of MAG-arm figures
that the pipeline reports only in prose. Each names the file it came from.

    python presentations/make_deck.py --out presentations/prFBA_updates_2026-10.pptx

Record depth: the default is the depth-100 standard (``--depth 100 --rule union``). The
amplicon-arm summaries (taxacc/gene-capture, PICRUSt2 comparison, three-panel comparison,
MAG-vs-16S crossover, annotation dialect, alignment wall time) are read from the depth-100
rerun trees (region_validation_d100/RUNBOOK_d100.md), and the hand-entered self-recovery /
emptied constants from ``--constants`` (default presentations/deck_constants_d100.json).
``--depth 20`` points the same inputs at the earlier depth-20 trees, but the slide text now
describes the depth-100 result; the deck as cited at depth 20 is
presentations/prFBA_updates_2026-10.pptx at commit 007deec.

Style (since 2026-10-27): short takeaway titles, phrases rather than sentences, one dominant
visual per slide, and the full explanation, caveats and sources in the speaker notes.
"""
import argparse
import json
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
RV = os.path.join(REPO, "region_validation")
CMP = os.path.join(RV, "picrust2_comparison", "data")
MAG = os.path.join(REPO, "mag_validation")
TP = os.path.join(REPO, "two_panel")
SR = os.path.join(RV, "shared_reference")

# --- record-depth switch (see module docstring; --depth 20 keeps every value below) ----
# Only the amplicon-arm inputs that depend on the record depth move; the MAG arm, the
# k-fold benchmark, the depth probe, the annotation ceiling / dialect and the two-panel
# files are depth-independent and keep their sources.
DEPTH = 20
AMP_RV = RV                                                    # "tax" / "cap"
CMP_SUMMARY = os.path.join(CMP, "comparison_summary.json")      # "cmp"
TP3_JSON = os.path.join(REPO, "three_panel", "three_panel_comparison.json")  # "tp3"
MAGX_JSON = os.path.join(MAG, "out", "crossover.json")         # "magx" (16S side depth-dependent)
DIALECT_JSON = os.path.join(CMP, "annotation_dialect.json")     # "dialect"
ALIGN_STAT = ("4.5 h", "to align all 60,461 unique benchmark amplicons against the full "
              "reference set")                                  # s04 stat (depth-20 run)
DEPTH_CONSTANTS = ("SELF_TOP20", "SELF_SAT", "SELF_CRED", "EMPTIED")

def configure_depth(depth, rule="union", constants=None):
    """Point the depth-dependent inputs at the depth-100 rerun trees (no-op at depth 20)."""
    global DEPTH, AMP_RV, CMP_SUMMARY, TP3_JSON, ALIGN_STAT, MAGX_JSON, DIALECT_JSON
    if depth == 20:
        return
    if depth != 100:
        raise SystemExit("--depth must be 20 or 100")
    DEPTH = depth
    sfx = "_d100" + ("_legacy" if rule == "legacy" else "")
    AMP_RV = os.path.join(REPO, "region_validation" + sfx)
    CMP_SUMMARY = os.path.join(REPO, "region_validation_d100", "picrust2_comparison", "data",
                               "comparison_summary.json")
    TP3_JSON = os.path.join(REPO, "three_panel", "three_panel_comparison%s.json" % sfx)
    MAGX_JSON = os.path.join(MAG, "out", "crossover%s.json" % sfx)
    DIALECT_JSON = os.path.join(REPO, "region_validation_d100", "picrust2_comparison", "data",
                                "annotation_dialect.json")
    ms_path = os.path.join(REPO, "region_validation_d100", "data", "hits", "raw",
                           "method_selection.json")
    if not os.path.exists(ms_path):
        raise SystemExit("--depth %d: no depth-%d alignment yet (%s)" % (depth, depth, ms_path))
    ms = json.load(open(ms_path))
    if ms["params"]["k2"] != depth:
        raise SystemExit("region_validation_d100 alignment has k2=%s, not %d"
                         % (ms["params"]["k2"], depth))
    ALIGN_STAT = ("%.1f h" % (ms["wall_seconds_total"] / 3600.0),
                  "to align all {:,} unique benchmark amplicons at depth {} against the full "
                  "reference set".format(ms["n_asvs"], depth))
    path = constants or os.path.join(REPO, "presentations", "deck_constants_d%d.json" % depth)
    vals = json.load(open(path))
    missing = [k for k in DEPTH_CONSTANTS
               if not isinstance(vals.get(k), dict)
               or any(vals[k].get(p) is None for p in ("original", "kmeans", "random"))]
    if missing:
        raise SystemExit("%s: fill in %s (each {original, kmeans, random}) before building the "
                         "depth-%d deck" % (path, ", ".join(missing), depth))
    for k in DEPTH_CONSTANTS:
        globals()[k] = {p: vals[k][p] for p in ("original", "kmeans", "random")}
    print("[make_deck] DEPTH=%d rule=%s: tax/cap <- %s, cmp <- %s, tp3 <- %s, constants <- %s"
          % (depth, rule, AMP_RV, CMP_SUMMARY, TP3_JSON, path))

# --- palette: teal/mint for the biology and for prFBA; amber for PICRUSt2 where the two
# methods are compared, elsewhere for a loss or non-answer, and for the NEW chip; the three
# panels keep the colours of their embedding figure (PANEL_COLORS) -----------------------
INK = RGBColor(0x10, 0x29, 0x2E)
DEEP = RGBColor(0x05, 0x34, 0x3D)
TEAL = RGBColor(0x02, 0x80, 0x90)
SEAFOAM = RGBColor(0x00, 0xA8, 0x96)
MINT = RGBColor(0x02, 0xC3, 0x9A)
AMBER = RGBColor(0xE8, 0x87, 0x1A)      # PICRUSt2 / loss / NEW chip
AMBER_TXT = RGBColor(0xC9, 0x6E, 0x0A)  # amber readable on the light tint
SEA_TXT = RGBColor(0x00, 0x7A, 0x6C)    # seafoam readable on the light tint
MUTED = RGBColor(0x5B, 0x71, 0x79)
LIGHT = RGBColor(0xEF, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY = "Cambria", "Calibri"

W, H = 13.333, 7.5
M = 0.6                      # left/right margin
TITLE_Y, TITLE_H = 0.42, 0.85
BODY_TOP = 1.5               # every content slide starts its body here


def load():
    j = lambda p: json.load(open(p))  # noqa: E731
    return {
        "tax": j(os.path.join(AMP_RV, "data", "taxacc_region_summary.json")),
        "cap": j(os.path.join(AMP_RV, "data", "gene_capture_by_region.json")),
        "cmp": j(CMP_SUMMARY),
        "ceil": j(os.path.join(CMP, "annotation_ceiling.json")),
        "probe": j(os.path.join(CMP, "depth_probe", "depth_probe_results.json")),
        "dialect": j(DIALECT_JSON),
        # the MAG arm: -n 100 is the audited RESULTS_v2 basis, used everywhere
        # except the cross-panel slide, which is depth-matched at -n 500
        "mag": j(os.path.join(MAG, "out", "mag_summary_full.json")),
        "magx": j(MAGX_JSON),
        "gtdb": j(os.path.join(MAG, "out_gtdb", "gtdb_summary_2000.json")),
        # the two-panel comparison
        "pan": j(os.path.join(TP, "panel_comparison.json")),
        "panmag": j(os.path.join(TP, "mag_panel_comparison.json")),
        "dec": j(os.path.join(TP, "space_deciles.json")),
        "knn": j(os.path.join(TP, "knn_density.json")),
        # the shared-reference k-fold benchmark and its phenotype arm
        "kg": j(os.path.join(SR, "data", "summary_genus.json")),
        "kf": j(os.path.join(SR, "data_family", "summary_family.json")),
        # the phenotype arm: amendment A-C12's pinned rerun (SPEC section 14) once both of
        # its passes have written a summary, else the pre-registered run of 2026-10-18; the
        # pre-registered pair stays available as *_pre for the C12 diagnosis and fix slides
        **phenotype_inputs(j),
        # pre-registered run against the amendment in use (gate, growth, reversals)
        "c12cmp": next((j(os.path.join(SR, "data", f)) for f in
                        ("c13_vs_prereg_compare_gf.json", "c12fix_compare_gf.json")
                        if os.path.exists(os.path.join(SR, "data", f))), None),
        # the three-panel comparison: taxonomy-stratified / k-means / uniform random
        "tp3": j(TP3_JSON),
        "tp3g": j(os.path.join(REPO, "three_panel", "three_panel_mag_gtdb.json")),
        "tp3p": j(os.path.join(REPO, "three_panel", "three_panel_mag_panelarm.json")),
        "dec3": j(os.path.join(REPO, "three_panel", "space_deciles_three.json")),
        "sel3": {k: j(p)["genomes"] for k, p in (
            ("original", os.path.join(RV, "data", "selection_10k.json")),
            ("kmeans", os.path.join(REPO, "kmeans_panel", "selection_kmeans.json")),
            ("random", os.path.join(REPO, "random_panel", "selection_random.json")))},
        # why the MAG arm has to exist, and what it costs to build one
        "ssu": j(os.path.join(CMP, "mag_gtdb_ssu_summary.json")),
        "ret": j(os.path.join(CMP, "mag_retention_summary.json")),
        "feed": j(os.path.join(CMP, "mag_direct_feed_summary.json")),
        "hmm": j(os.path.join(CMP, "src16s_hmm_check_summary.json")),
        "aln_ms": j(os.path.join(AMP_RV, "data", "hits", "raw", "method_selection.json")),
        "aln_gpu": j(os.path.join(REPO, "gpu_align_stats.json")),
        "aln_emb": j(os.path.join(REPO, "old_asv_comparison", "kmer_augmented_search.json")),
        "aln_cos": j(os.path.join(REPO, "density_stats.json"))["global_pairwise_cosine"]["mean"],
        "bvbrc_raw": j(os.path.join(REPO, "manifest.json"))["n_input_sequences"],
        "sr_leg": j(os.path.join(REPO, "region_validation_d100_legacy", "data", "self_recovery_diagnosis.json")),
        "sr_uni": j(os.path.join(REPO, "region_validation_d100", "data", "self_recovery_diagnosis.json")),
        "sr_k2": j(os.path.join(REPO, "region_validation_d100", "data", "hits", "raw", "method_selection.json"))["params"]["k2"],
        "calib": (lambda sc: {"meta": sc["meta"]["bootstrap"], "ec": {k: v for k, v in sc["primary_controlled"]["exact_tie|d100"]["ec"].items() if k in ("all|all=all", "single|all=all", "multi|all=all", "n2|all=all", "n3|all=all", "n4+|all=all")}})(j(os.path.join(SR, "data", "calibration", "scores.json"))),
        "calib_ok": j(os.path.join(SR, "data", "calibration", "assertions.json"))["all_required_zero"],
        "mcn": j(os.path.join(AMP_RV, "data", "mcnemar_regions.json")),
        "c2": j(os.path.join(SR, "data_c2", "summary_incself.json")),
        "ctl": j(os.path.join(SR, "data", "controls.json")),
        "kplan": j(os.path.join(SR, "data", "fold_plan.json"))["meta"],
        "vts": j(os.path.join(SR, "data", "template_direction_audit", "validation_summary.json")),
        "taxd": {k: j(os.path.join(REPO, t, "data", "taxacc_region_summary.json"))
         for k, t in (("d20", "region_validation"),
                      ("L100", "region_validation_d100_legacy"),
                      ("U100", "region_validation_d100"))},
        "capd": {k: j(os.path.join(REPO, t, "data", "gene_capture_by_region.json"))
         for k, t in (("d20", "region_validation"),
                      ("L100", "region_validation_d100_legacy"),
                      ("U100", "region_validation_d100"))},
        "mcnd": {k: j(os.path.join(REPO, t, "data", "mcnemar_regions.json"))
         for k, t in (("d20", "region_validation"),
                      ("L100", "region_validation_d100_legacy"))},
        "cmpd": {"d20": j(os.path.join(CMP, "comparison_summary.json")),
         "d100": j(os.path.join(REPO, "region_validation_d100", "picrust2_comparison",
                                "data", "comparison_summary.json"))},
        "absd": {k: (lambda sel: (
            sum(1 for a in sel if a != "_meta"),
            sum(1 for a, v in sel.items() if a != "_meta"
                and v["exclude_0.987"]["n_selected"] == 0),
            sum(1 for a, v in sel.items() if a != "_meta"
                and v["exclude_0.987"]["n_selected"] == 0
                and v["exclude_0.987"].get("best_identity") in (0.0, None))))(
            j(os.path.join(REPO, t, "picrust2_comparison", "data", "prfba_selection.json")))
         for k, t in (("d20", "region_validation"), ("d100", "region_validation_d100"))},
        "tdv": j(os.path.join(SR, "data", "template_direction_audit", "validation_summary.json")),
        "magp": j(MAGX_JSON.replace("crossover", "paired_comparison")),
        "accmap": j(os.path.join(REPO, "reference_maps", "gtdb_r226_acc_to_bvbrc.stats.json")),
        "tp3_emptied": (lambda p: j(p) if os.path.exists(p) else None)(TP3_JSON.replace("three_panel_comparison", "emptied_strata")),
        "random_meta": j(os.path.join(REPO, "random_panel", "selection_random.json"))["meta"],
        "sel_legacy_abst": sum(1 for v in j(os.path.join(os.path.dirname(CMP_SUMMARY), "prfba_selection_legacy.json")).values() if isinstance(v, dict) and "exclude_0.987" in v and not v["exclude_0.987"]["reps"]),
    }


# --- numbers the pipeline reports only in prose; source named on each ---------
# mag_validation/out/RESULTS_v2.md, "the AF gate is almost everything"
# RESULTS_v2.md's table counts qualifying hits over all 9,997 MAGs but divides by the 9,837
# scoreable ones; recounted on one basis (scoreable organisms, c100_x0 hits after the
# scorer's exclusion) on 2026-10-18. It is the share with a qualifying surviving hit, before
# selection, not the resolve rate (0.308).
AF_SWEEP = [
    ("as shipped — ANI \u2265 83, AF \u2265 15%", 0.311),
    ("ANI floor 83 \u2192 0", 0.316),
    ("AF gate 15% \u2192 5%", 0.379),
    ("AF gate 15% \u2192 2%", 0.452),
    ("both relaxed — ANI \u2265 80, AF \u2265 2%", 0.567),
    ("any surviving hit at all (ceiling)", 0.655),
]
# GTDB r226's skani sketch holds 143,614 species representatives; 732,475 is the accession
# count of the bac120 + ar53 TAXONOMY files, which RESULTS_v2.md's "73x" divided by mistake.
GTDB_FOLD = 14                 # 143,614 / 10,000-genome panel (14.4x), verified 2026-10-18
PANEL_KM_SPECIES = 4935        # two_panel/RESULTS.md section 1
PANEL_KM_PHYLA = 142           # two_panel/RESULTS.md section 1
PANEL_OG_PHYLA = 206           # two_panel/RESULTS.md section 1
PANEL_OVERLAP = 804            # TWO_PANEL_RUNBOOK.md part 2
# mag_validation/synth_mags.py: module constants and the docstring's justification
FRAG_MEAN_KB, FRAG_MIN_KB = 50, 2.5
PANEL_CLOSED_PCT, PANEL_LT10_PCT = 19.8, 40.4
# --- the three-panel analysis (prFBA_paper sec:twopanel; verified 2026-10-18) ---------
# self-recovery, amplicon level, from each panel's data/hits/asv_top20_alignment_hits.json
# (676 MB per panel, so not loaded here): source in its own top 20; misses where the top
# 20 are all exact ties; misses where the source's sequence is credited to another genome
SELF_TOP20 = {"original": 96.7, "kmeans": 69.0, "random": 61.2}
SELF_SAT = {"original": 3.3, "kmeans": 15.6, "random": 24.6}
SELF_CRED = {"original": 0.0, "kmeans": 15.4, "random": 14.2}
EMPTIED = {"original": 14.3, "kmeans": 40.4, "random": 49.9}   # full-length, tau = 0.987
GTDB_BEST_ANI = {"original": 89.5, "kmeans": 91.4, "random": 92.1}
# --- the synthetic-genome re-evaluation (prFBA_paper sec:recon) ----------------------
BENCH_N1, DEPLOY_N1 = 86.8, 62.5     # % of answered queries whose selection is one genome
XPHYLUM_DEPLOY = (31.7, 59.0)        # % of n=2 / n>=3 deployment merges spanning phyla
XPHYLUM_BENCH = 6.4                  # same, controlled benchmark (all multi-genome sets)
# mean per-genome EC-F1 difference, prFBA exact-tie minus PICRUSt2, over the 755 paired
# ATP-informative genomes (phenotype_atp rows, deduplicated; recomputed 2026-10-18)
PHENO_EC_GAP = -0.0284
# --- the claim-level audits of the manuscript (prFBA_paper f7fdc38, d2d7516) ----------
# the second audit's record holds 48 flags, one overturned on recheck ("no correction needed")
AUDIT_CLAIMS, AUDIT_DEFECTS = (417, 178), (114, 47)
# the verification of this deck and the manuscript patch, 2026-10-18: five auditors, each
# flag re-derived by a skeptic (claims checked, defects confirmed)
DECK_AUDIT = (598, 54)
# the second pass, over the fixes themselves (claims checked, distinct defects confirmed)
FIX_AUDIT = (377, 14)
# faults our own runs and reproduction checks caught (the phenotype gating and duplicate-row
# faults were claim-audit catches, so they are listed with those, not here)
PIPELINE_CAUGHT = [
    "A gene-family fetcher with no User-Agent: every request refused, fds leaked",
    "A three-panel scorer that counted any non-abstention as a call at every rank "
    "(genus coverage overstated by up to 0.19)",
    "A density-decile comparison that read its stored array in reverse",
    "A reproduction check that passed while comparing zero cells (a wrong key path)"]


# --- constants added in the brevity/depth-100 rewrite (2026-10-27); source on each ---
# --- the deployment build in KBase (s02, s06) ----------------------------------------
# newdata_build_54d1150.log: '[select] ... over 3276 ASVs', '[build] DONE: built=3224
# abstain=52'; newdata_upload_54d1150.log: 'GenomeSet ... -> 261948/3226/1 (3224 genomes)'
KBASE_ASVS, KBASE_BUILT, KBASE_ABSTAIN, KBASE_WS = 3276, 3224, 52, 261948
# the real-amplicon deployment run behind DEPLOY_N1 / XPHYLUM_DEPLOY, made on 20-hit
# records before the 100-hit standard: prFBA.tex sec:recon (3,753 answered + 197 abstaining)
DEPLOY_ASVS = 3950
# all 100 retained hits tied at the best score at depth 100: distinct sequences, of all
# distinct sequences, median species spanned; and sequences whose whole 500-candidate
# shortlist is tied (prFBA.tex sec:align and sec:limits; region_validation_d100/
# D100_VS_D20.md section 2G)
TIE_SAT_D100 = (453, 60461, 19)
TIE_ALL500 = 95
# region_validation/shared_reference/C12_FIX_PLAN.md, "The run" table: wall-time estimates
# (1-2 min per genome per worker on 16 workers; 13-15 min per genome on 8-10 workers)
C12_RUN_HOURS = {"atp": "2–3 h", "gf": "4–6 h"}
# region_validation/shared_reference/EGC_C12_REPORT.md, "the five-reaction fix": each set to
# its physiological direction '>' (the catalase loop is closed by rxn09180, fixed at rxn00006)
C12_FIX_RXNS = ("rxn00006", "rxn10966", "rxn00381", "rxn00084", "rxn08167")
# --- the depth-100 standard and the selection-rule choice (prFBA_paper sec:depth, sec:select)
# distinct sequences whose 500 shortlisted references all tie at the top score, of all
# 60,461: recounted 2026-10-27 from region_validation_d100/data/hits/raw/gpu_scores.npy
# (208 have >= 200 tied), as prFBA.tex sec:align states; the 120 MB array is not loaded here
SHORTLIST_ALL_TIED = (95, 60461)
# stage 3 since 2026-10-27 (region_validation_d100/GPU_CAND_200.md): the GPU score ranks
# the shortlist and Biopython only aligns the top k2. Hits on which the GPU and Biopython
# scores agree (gpu_cand_check.py: all of the depth-100 records'); per-call cost (ms) of a
# Biopython re-score and of an alignment traceback (gpu_cand_timing.py, idle machine); and
# the stage-3 wall (s, 40 workers) of the depth-100 run and of the traceback-only re-run
BITEXACT_HITS = 6046100
RESCORE_MS = (4.3, 16.0)
STAGE3_S = (7128, 3425)
# mean genomes selected per self-inclusion answer (self_mode include, outcome ok), column
# n_selected of region_validation_d100{,_legacy}/data/gene_capture_cells.csv (V4: 4.57 / 1.43)
INCL_N_SELECTED = {"union": 2.88, "legacy": 1.32}
# 95% genome-clustered bootstrap CI (1,000 replicates) of prFBA minus PICRUSt2 population EC
# F1 at depth 100: region_validation_d100/D100_VS_D20.md section 2D (no repo JSON carries it)
GAP_CI_D100 = {"union": (-0.021, -0.011), "legacy": (-0.007, 0.003)}
# the five-reaction template fix that closes the C12 leaks (ModelSEEDTemplates 21a8084):
# region_validation/shared_reference/C12_FIX_PLAN.md
FIVE_RXN_FIX = ("rxn00006", "rxn10966", "rxn00381", "rxn00084", "rxn08167")
# --- the MAG-arm slides (s16-s20, s18b): figures the pipeline reports only in prose ------
# mag_validation/synth_mags.py docstring: 8 synthetic variants sharing the first contig
# header '>src_0' came back from skani as 5 query results
SKANI_COLLIDE = (8, 5)
# mag_validation/out/RESULTS_v2.md "Marker reach is the dominant cause" (prFBA_paper
# sec:magval): organisms with a same-genus congener in the panel; median AF_query (%) of
# their best surviving hit; share of them (%) below the 15% AF gate
MAG_REACH = (6011, 20.3, 45.1)
# RESULTS_v2.md, same section: MAG resolve rate in the amplicon arm's best-non-self 16S
# identity bins 0.945-0.97 and 0.97-0.987 (the amplicon side is read from D["cmp"])
BAND_MAG = (0.082, 0.202)
# the 2% AF-gate counterfactual (RESULTS_v2.md "The gate is not a defect"; prFBA_paper
# sec:magval): organisms a 2% gate newly answers with a non-empty prediction, and their
# mean gene-capture F1 (per-MAG cells recomputed 2026-10-18; the gate is MAG-only, so
# depth-independent)
AF2_NEW = (1367, 0.59946)
# how many of those 1,367 exceed the full-length 16S level of the pairing in use, keyed by
# crossover*.json's FullLength16S_amp at 100%: 369 (27.0%) at depth 20 (0.7092), 365 at
# depth 100 under the union (0.71), 435 (31.8%, prFBA_paper sec:magval) at depth 100
# under the legacy reducer (0.6912); recounted from the same per-MAG cells 2026-10-27
AF2_BEAT_FL = {0.7092: 369, 0.71: 365, 0.6912: 435}
# mag_validation/out/DEPTH_PROBE.md: MAG queries both capped at -n 100 and emptied by
# self-exclusion, of all 99,970 (contaminated grid included); and the direct test --
# capped queries and uncapped controls re-searched at -n 2000, selections changed (all
# abstention -> resolution)
MAG_CAPPED_EMPTIED = (156, 99970)
# the same, over the 49,985 uncontaminated 30-100% records the MAG slides report
# (prFBA_paper sec:depth: "93 of the 49,985 uncontaminated 30--100% records ... or 0.19%")
MAG_CAPPED_EMPTIED_X0 = (93, 49985)
MAG_DEPTH_PROBE = (447, 200, 11)
# --- the checking since the last deck (s24_audit); prose-only figures, source on each -----
# region_validation_d100/D100_VS_D20.md section 5: "Confirmed: 85 of 91 claims checked" (the
# verifiers' raw verdicts, workflow record wf_13e52561-c3b.json, are 75 confirmed, 15 partially,
# 1 refuted; the report's adjudicated count is the one quoted)
D100_REPORT_CHECK = (85, 91)
# findings per verification round of the manuscript's depth-100 revision, all applied in
# prFBA_paper 735e208, 7589686 and eba93af ("four wording and precision slips"); per-round
# counts from the workflow records wf_61e46fab-aa4.json, wf_eb9b676d-0e8.json and
# wf_0918b912-346.json (~/.claude/projects/-home-freiburger-Documents-prFBA/<session>/workflows)
MS_D100_ROUNDS = (67, 32, 4, 9)   # 4th: the union rewrite, region_validation_d100/UNION_VERIFICATION.md
# review rounds of the ModelSEEDTemplates PR texts: prFBA c569aff ("after two review rounds";
# review found the six-reaction fix set NadX rxn05119/rxn05117 '<', blocking NAD synthesis)
# and 5011c64 ("a third independent review")
TEMPLATE_PR_ROUNDS = 3
# region_validation_d100/D100_VS_D20.md section 5, "Mismatched source taxa": FullLength16S
# exclude rows per k-means/random tree whose src_taxon_id differs from benchmark_truth.json,
# and the largest effect on the emptied share (points)
SRC_TAXON_MISMATCH = (23, 33, 0.06)
# prFBA_paper/prFBA.tex sec:twopanel: Archaea among the 140,197 admissible organisms
POOL_ARCHAEA = 4945
# GTDB r226 species representatives in the skani sketch (the GTDB_FOLD numerator above)
GTDB_REPS = 143614

def sg(x, d=3):
    """Signed number with a typographic minus, so a line never breaks inside it.

    Rounds half away from zero on the stored decimal (-0.1235 -> -0.124), as the
    manuscript does; binary float formatting would give -0.123."""
    from decimal import Decimal, ROUND_HALF_UP
    v = Decimal(repr(x)).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP)
    return f"{v:+.{d}f}".replace("-", "\u2212")


def rd(x, d=3):
    """Unsigned counterpart of sg(): half away from zero on the stored decimal."""
    from decimal import Decimal, ROUND_HALF_UP
    return f"{Decimal(repr(x)).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP):.{d}f}"


def phenotype_inputs(j):
    """phatp/phgf/phgf_gated from A-C12's rerun when both summaries exist (c12fix True), else
    from the pre-registered run; *_pre always the pre-registered run. phgf_gated holds the
    truth genomes whose gapfilled model fails the sanity gate (C12 leg b), counted as
    genomes: meta.n_gated_out counts rows, one per (genome, pipeline)."""
    def tree(tag):
        d = os.path.join(SR, "data")
        return (os.path.join(d, "phenotype_atp" + tag, "phenotype_summary.json"),
                os.path.join(d, "phenotype_gf" + tag, "phenotype_summary.json"),
                os.path.join(d, "phenotype_gf" + tag, "per_genome.jsonl"))
    pre = tree("")
    # the latest amendment with both summaries: A-C13 (updated ModelSEEDpy), then A-C12
    use, tag, label = pre, "", None
    for t, lab in (("_c13", "A-C13"), ("_c12fix", "A-C12")):
        cand = tree(t)
        if os.path.exists(cand[0]) and os.path.exists(cand[1]):
            use, tag, label = cand, t, lab
            break
    return {"phatp": j(use[0]), "phgf": j(use[1]), "phgf_gated": gated_truth_genomes(use[2]),
            "phatp_pre": j(pre[0]), "phgf_pre": j(pre[1]),
            "phgf_gated_pre": gated_truth_genomes(pre[2]), "c12fix": bool(tag),
            "amend_tag": tag, "amend_label": label,
            # the amendments' wording: A-C13 only adds the ModelSEEDpy update to A-C12
            "amend_text": ("amendments A-C12 and A-C13" if label == "A-C13"
                           else "amendment A-C12")}


def gated_truth_genomes(path):
    out = set()
    with open(path) as fh:
        for line in fh:
            if '"truth"' in line and '"gated"' in line:
                r = json.loads(line)
                if r["arm"] == "truth" and r["status"] == "gated":
                    out.add(r["src_genome_id"])
    return sorted(out)


def composition(gs):
    """Panel composition, counting unclassified placeholders as their own genus bins."""
    import collections
    named = lambda v: bool(v) and not str(v).startswith("unclassified")   # noqa: E731
    gen = collections.Counter(e["genus"] for e in gs)
    v = sorted(gen.values()); n, tot = len(v), sum(v)
    gini = (2 * sum((i + 1) * x for i, x in enumerate(v))) / (n * tot) - (n + 1) / n
    arc = sum(1 for e in gs if e["domain"] == "Archaea")
    return dict(n=len(gs), archaea=arc, arc_pct=100 * arc / len(gs),
                phyla=len({e["phylum"] for e in gs}),
                families=len({e["family"] for e in gs if named(e["family"])}),
                genera=len({e["genus"] for e in gs if named(e["genus"])}),
                species=len({e["species"] for e in gs}), gini=gini, largest=max(v))


def tp3_cell(D, region, regime, endpoint):
    return next(c for c in D["tp3"]["cells"] if c["region"] == region
                and c["regime"] == regime and c["endpoint"] == endpoint)


REG_LABEL = {"FullLength16S": "Full-length", "V1-V2": "V1–V2", "V1-V3": "V1–V3",
             "V3-V4": "V3–V4", "V4": "V4", "V4-V5": "V4–V5", "V4-V5_944R": "V4–V5 944R",
             "V7-V9": "V7–V9"}
# the panels in three_panel/plot_space_three.py's colours, so each keeps its colour from the
# embedding figure to the charts (Okabe-Ito ochre, blue, green)
PANEL_COLORS = {"original": RGBColor(0xB0, 0x76, 0x00), "kmeans": RGBColor(0x00, 0x72, 0xB2),
                "random": RGBColor(0x00, 0x9E, 0x73)}


# --------------------------------------------------------------------- helpers
def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, *, size=15, bold=False, color=INK, font=BODY, space_after=6,
         align=PP_ALIGN.LEFT, italic=False, first=False, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    return p


def rect(slide, x, y, w, h, fill, *, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None,
         adj=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = adj
        except (IndexError, AttributeError):
            pass
    s.text_frame.word_wrap = True
    s.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    return s


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


# Slides whose content is new in this revision (2026-10-27) carry the amber NEW chip; the
# policy lives here, not in each slide function (main() sets CURRENT_SLIDE).
NEW_SLIDES = {"s06c_calibration", "s14_depth", "s14b_rule", "s17b_c12_cause", "s17c_c12_fix",
              "s17d_template_audit", "s18b_accession_map", "s24_audit"}
CURRENT_SLIDE = None


def title(slide, text, *, sub=None, dark=False, tag=None):
    """Content-slide title: identical position, size and alignment on every slide."""
    tag = "NEW" if CURRENT_SLIDE in NEW_SLIDES else None
    if tag:
        chip = rect(slide, W - M - 0.72, TITLE_Y + 0.05, 0.72, 0.32, AMBER)
        tf = chip.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, tag, size=11.5, bold=True, color=INK, align=PP_ALIGN.CENTER, first=True)
    tf = textbox(slide, M, TITLE_Y, W - 2 * M - (0.95 if tag else 0), TITLE_H)
    para(tf, text, size=30, bold=True, color=WHITE if dark else INK, font=HEAD, first=True,
         space_after=2)
    if sub:
        para(tf, sub, size=13.5, color=MINT if dark else MUTED, space_after=0)


_STATS = []   # (run, size) of the current slide's auto-sized stat values; see main()


def stat(slide, x, y, w, value, label, *, color=TEAL, vsize=None, h=1.45, tint=LIGHT):
    # one rule for every stat card: 32pt, or 28pt for values over 7 characters, and main()
    # then sets every auto-sized value on a slide to the smallest size among them
    auto = vsize is None
    if auto:
        vsize = 32 if len(value) <= 7 else 28
    card = rect(slide, x, y, w, h, tint)
    tf = card.text_frame
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = para(tf, value, size=vsize, bold=True, color=color, font=HEAD, first=True,
             space_after=1)
    if auto:
        _STATS.append((p.runs[0], vsize))
    para(tf, label, size=11.5, color=MUTED, space_after=0, line=0.92)
    return card


def bullets(tf, items, *, size=14.5, space=9, color=INK, head_color=TEAL):
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            head, rest = it
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(space)
            r = p.add_run()
            r.text = head
            r.font.size, r.font.bold, r.font.name = Pt(size), True, BODY
            r.font.color.rgb = head_color
            r2 = p.add_run()
            r2.text = rest
            r2.font.size, r2.font.name = Pt(size), BODY
            r2.font.color.rgb = color
        else:
            para(tf, it, size=size, color=color, space_after=space, first=(i == 0))


def hanging(p, indent=0.17):
    """Wrap wrapped lines under the text, not under the bullet glyph."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent * 914400)))
    pPr.set("indent", str(-int(indent * 914400)))
    return p


def footer(slide, text, *, dark=False):
    tf = textbox(slide, M, H - 0.6, W - 2 * M, 0.3)
    para(tf, text, size=11, color=MINT if dark else MUTED, first=True, space_after=0)


def style_chart(chart, *, legend=True, size=11):
    chart.has_title = False          # otherwise a single series auto-titles itself
    chart.font.size = Pt(size)
    chart.font.name = BODY
    chart.font.color.rgb = MUTED
    if legend:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        chart.legend.font.color.rgb = INK
    else:
        chart.has_legend = False


def chart_title(chart, text):
    """An in-chart heading, in one style across the deck."""
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    for r in tf.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.name = Pt(13), True, BODY
        r.font.color.rgb = INK


def series_color(series, color):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color


# ------------------------------------------------------------- visual vocabulary
# Shared by every slide since the brevity pass of 2026-10-27: phrases rather than
# sentences, one visual per slide, and the explanation in the speaker notes.

def notes(slide, text):
    """Speaker notes carry the full sentences, caveats and sources."""
    slide.notes_slide.notes_text_frame.text = text


def takeaway(slide, text, x, y, w, *, dark=False, size=18, color=None, h=0.6):
    """The slide's one-line message, large and in the accent colour."""
    tf = textbox(slide, x, y, w, h)
    para(tf, text, size=size, bold=True, font=HEAD,
         color=color or (MINT if dark else TEAL), first=True, space_after=0)
    return tf


def _bar_chart(slide, kind, x, y, w, h, cats, series, *, colors=None, fmt="0.000",
               vmin=None, vmax=None, label_size=None, cat_size=12, dark=False, legend=None,
               gap=60, axis=False):
    """Every bar prints its value, so the value axis and its gridlines are hidden unless
    axis=True. Labels are 14pt on one series, 12pt when series sit side by side."""
    cd = CategoryChartData()
    cd.categories = list(cats)
    for name, vals in series:
        cd.add_series(name, list(vals), number_format=fmt)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    many = len(series) > 1
    if label_size is None:
        label_size = 12 if many else 14
    style_chart(ch, legend=many if legend is None else legend, size=cat_size)
    txt = WHITE if dark else MUTED
    ch.font.color.rgb = txt
    if ch.has_legend:
        ch.legend.font.color.rgb = WHITE if dark else INK
    va, ca = ch.value_axis, ch.category_axis
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    va.has_major_gridlines = False
    va.visible = axis
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(cat_size - 1)
    va.tick_labels.font.color.rgb = txt
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(cat_size)
    ca.tick_labels.font.color.rgb = WHITE if dark else INK
    ch.plots[0].gap_width = gap
    if many:
        ch.plots[0].overlap = -10
    palette = colors or [TEAL, AMBER, SEAFOAM, MUTED]
    for i, sr in enumerate(ch.series):
        if not many and colors and len(colors) == len(cats):
            for j, pt in enumerate(sr.points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = colors[j]
        else:
            series_color(sr, palette[i % len(palette)])
        sr.has_data_labels = True
        dl = sr.data_labels
        dl.show_value = True
        dl.number_format, dl.number_format_is_linked = fmt, False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(label_size)
        dl.font.bold = True
        dl.font.color.rgb = WHITE if dark else INK
        dl.font.name = BODY
    return ch


def hbar(slide, x, y, w, h, cats, vals, **kw):
    """Horizontal bars in the deck's style; cats/vals listed top to bottom, colors (one
    per bar) likewise. Values print at the bar ends; no legend."""
    colors = kw.pop("colors", None)
    return _bar_chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, list(cats)[::-1],
                      [("v", list(vals)[::-1])],
                      colors=colors[::-1] if colors else None, **kw)


def vbar(slide, x, y, w, h, cats, series, **kw):
    """Columns in the deck's style: series = [(name, values), ...]; a legend appears
    when there is more than one series."""
    return _bar_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cats, series, **kw)


def flow(slide, x, y, w, h, steps, *, fill=LIGHT, color=INK, sub_color=MUTED,
         arrow=TEAL, size=14, gap=0.42, anchor=MSO_ANCHOR.MIDDLE, pad_top=0.06):
    """A left-to-right process: one rounded box per step, arrows between. A step is a
    short label or a (label, sub-label) pair. anchor=TOP with a pad_top keeps the headings
    level when one of them wraps."""
    n = len(steps)
    bw = (w - gap * (n - 1)) / n
    for i, st in enumerate(steps):
        head, sub = st if isinstance(st, tuple) else (st, None)
        bx = x + i * (bw + gap)
        box = rect(slide, bx, y, bw, h, fill)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Inches(0.1)
        tf.margin_top, tf.margin_bottom = Inches(pad_top), Inches(0.06)
        tf.vertical_anchor = anchor
        para(tf, head, size=size, bold=True, color=color, font=HEAD, align=PP_ALIGN.CENTER,
             first=True, space_after=2 if sub else 0)
        if sub:
            para(tf, sub, size=size - 3, color=sub_color, align=PP_ALIGN.CENTER,
                 space_after=0, line=0.95)
        if i < n - 1:
            aw = min(0.22, gap - 0.1)
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                        Inches(bx + bw + (gap - aw) / 2),
                                        Inches(y + h / 2 - 0.14), Inches(aw), Inches(0.28))
            ar.fill.solid()
            ar.fill.fore_color.rgb = arrow
            ar.line.fill.background()
            ar.shadow.inherit = False


def table(slide, x, y, w, rows, *, col_w=None, size=13, row_h=0.44, dark=False,
          emph_col=None, header_fill=None):
    """A compact comparison table; rows[0] is the header. Cells are short strings;
    emph_col (an index) is set bold in the accent colour."""
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                                 Inches(row_h * nr))
    tbl = shp.table
    widths = col_w or [w / nc] * nc
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    body_a = RGBColor(0x0A, 0x4A, 0x55) if dark else LIGHT
    body_b = RGBColor(0x08, 0x40, 0x4A) if dark else WHITE
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = ((header_fill or (TEAL if not dark else SEAFOAM)) if i == 0
                                     else (body_a if i % 2 else body_b))
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            head = i == 0
            col = (WHITE if head else
                   ((MINT if dark else TEAL) if j == emph_col else (WHITE if dark else INK)))
            para(tf, str(val), size=size - (1 if head else 0), bold=head or j == 0
                 or j == emph_col, color=col, first=True, space_after=0,
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    if not dark:            # a teal rule under the last row, so a white row keeps its edge
        from lxml import etree
        from pptx.oxml.ns import qn
        for j in range(nc):
            tcPr = tbl.cell(nr - 1, j)._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("a:lnB")):
                tcPr.remove(old)
            ln = etree.Element(qn("a:lnB"), w=str(Pt(1.25)), cap="flat", cmpd="sng",
                               algn="ctr")
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"),
                             val=str(TEAL))
            pos = sum(1 for ch in tcPr if ch.tag in (qn("a:lnL"), qn("a:lnR"), qn("a:lnT")))
            tcPr.insert(pos, ln)
    return tbl


# ------------------------------------------------------------------- the slides
def s01_title(prs, D):
    s = blank(prs)
    background(s, DEEP)
    k2 = D["aln_ms"]["params"]["k2"]
    tf = textbox(s, M, 1.35, W - 2 * M, 2.2)
    para(tf, "prFBA", size=60, bold=True, color=WHITE, font=HEAD, first=True, space_after=2)
    para(tf, "Probabilistic flux balance analysis for amplicon-resolved metabolic modeling",
         size=22, color=MINT, font=HEAD, space_after=0)
    # what this update brings, as three chips
    chips = (f"{k2}-hit records, the new standard" if k2 == 100 else f"{k2}-hit records",
             "Like for like, PICRUSt2 leads",
             "p(function) ranks, but overstates")
    gap = 0.3
    cw = (W - 2 * M - gap * (len(chips) - 1)) / len(chips)
    for i, t in enumerate(chips):
        c = rect(s, M + i * (cw + gap), 4.0, cw, 0.7, RGBColor(0x0A, 0x4A, 0x55), line=SEAFOAM)
        ctf = c.text_frame
        ctf.margin_left = ctf.margin_right = Inches(0.12)
        ctf.margin_top = ctf.margin_bottom = 0
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(ctf, t, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True,
             space_after=0)
    tf2 = textbox(s, M, 5.9, W - 2 * M, 0.4)
    para(tf2, "Andrew Freiburger  ·  Tyo & Henry labs  ·  Project update, October 2026",
         size=13, color=MINT, first=True, space_after=0)
    g = D["kg"]
    gate = ((" Its gapfill phenotype gate (C12), which failed in the pre-registered run, passes "
             f"under {D['amend_text']} (corrected templates, post hoc and disclosed), so that "
             "leave-genus-out run is valid; no EC number changed.")
            if D["c12_rescored"] and g["headline"].get("valid") else
            "" if g["headline"].get("valid", True) else (
        " That leave-genus-out run is formally marked invalid by its gapfill phenotype gate "
        "(C12) until the template-direction fix is adopted; no EC number depends on that gate."))
    lead = (f"This update reports every benchmark result at the new standard record depth: "
            f"each amplicon now keeps its {k2} best hits, from an edlib shortlist of "
            f"{D['aln_ms']['params']['prefilter_k']} candidates that is rescored in full; the "
            "previous deck used 20-hit records. Both real-ASV deployment runs (the KBase build "
            "and the weighting census) predate the standard and are labelled as 20-hit runs."
            if k2 == 100 else
            f"This build reports the amplicon arm at the earlier record depth of {k2} hits.")
    notes(s, lead + " The headline has not changed direction: on a "
             "controlled benchmark, where both methods share one reference set, one answer key "
             "and leave-genus-out folds, PICRUSt2 predicts EC content more accurately than "
             "prFBA, beyond the effect size fixed in advance." + gate +
             " New in this deck is a calibration of the per-function weight p(function): it "
             "ranks functions correctly but is overconfident, and thresholding it barely moves "
             "accuracy. Raising the record depth more than halved the self-exclusion abstention "
             "rate under the shipped exact-tie rule while leaving conditional quality roughly "
             "unchanged, so coverage and population figures remain lower bounds. Re-run at "
             "depth 100 under one selection rule, the three-panel comparison still shows "
             "conditional quality generalizing across panels while self-exclusion coverage "
             "does not.")
    return s


def s02_problem(prs, D):
    from pptx.enum.shapes import MSO_CONNECTOR
    s = blank(prs)
    title(s, "Amplicons name the members, not their metabolism",
          sub="16S surveys: more species per unit of DNA, but no genes")
    top, ch, cw = 1.7, 3.0, 3.5

    def card(x, head, caption):
        c = rect(s, x, top, cw, ch, LIGHT)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.16)
        para(tf, head, size=18, bold=True, color=INK, font=HEAD, first=True, space_after=0,
             align=PP_ALIGN.CENTER)
        p = para(tf, caption, size=17, bold=True, color=TEAL, font=HEAD, space_after=0,
                 align=PP_ALIGN.CENTER)
        p.space_before = Pt(140)
        return c

    # left: a census, drawn as a ranked-abundance profile
    card(M, "16S census", "who is present")
    bx, by = M + 0.42, top + 0.78
    for i, (frac, col) in enumerate(((1.0, TEAL), (0.74, SEAFOAM), (0.52, MINT),
                                     (0.34, SEAFOAM), (0.2, MUTED))):
        rect(s, bx, by + i * 0.33, (cw - 0.84) * frac, 0.22, col, adj=0.5)
    # right: a metabolic network
    rx = W - M - cw
    card(rx, "Metabolic model", "what they do")
    nodes = [(0.55, 1.0), (1.35, 0.78), (2.25, 1.05), (2.95, 0.82), (0.95, 1.85),
             (1.85, 1.78), (2.75, 1.98)]
    edges = [(0, 1), (1, 2), (2, 3), (0, 4), (1, 4), (1, 5), (2, 5), (5, 6), (3, 6), (4, 5)]
    for a, b in edges:
        (x1, y1), (x2, y2) = nodes[a], nodes[b]
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(rx + x1), Inches(top + y1),
                                    Inches(rx + x2), Inches(top + y2))
        ln.line.color.rgb = RGBColor(0x9F, 0xB4, 0xB8)
        ln.line.width = Pt(2)
    for i, (nx, ny) in enumerate(nodes):
        rect(s, rx + nx - 0.15, top + ny - 0.15, 0.3, 0.3, (TEAL, SEAFOAM, MINT)[i % 3],
             shape=MSO_SHAPE.OVAL)
    # the bridge
    ax, aw = M + cw + 0.3, W - 2 * M - 2 * cw - 0.6
    ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(top + 0.45), Inches(aw),
                            Inches(1.15))
    ar.fill.solid()
    ar.fill.fore_color.rgb = TEAL
    ar.line.fill.background()
    ar.shadow.inherit = False
    atf = ar.text_frame
    atf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(atf, "prFBA", size=24, bold=True, color=WHITE, font=HEAD, align=PP_ALIGN.CENTER,
         first=True, space_after=0)
    tf = textbox(s, ax, top + 1.72, aw, 0.6)
    para(tf, "ASV → co-optimal reference genomes", size=15, bold=True, color=INK,
         align=PP_ALIGN.CENTER, first=True, space_after=2)
    para(tf, "ambiguity kept as probabilities", size=15, color=MUTED, align=PP_ALIGN.CENTER,
         space_after=0)
    # the usual fix, and its costs
    mag = rect(s, ax + 0.35, top + ch - 0.46, aw - 0.7, 0.46, WHITE, line=MUTED, adj=0.5)
    mtf = mag.text_frame
    mtf.margin_top = mtf.margin_bottom = 0
    mtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(mtf, "vs MAGs: incomplete, costly, abundant-biased", size=12.5, color=MUTED,
         align=PP_ALIGN.CENTER, first=True, space_after=0)
    # the scale of this work
    sy, sw, sg_ = top + ch + 0.5, (W - 2 * M - 3 * 0.3) / 4, 0.3
    n_raw, n_amp = D["bvbrc_raw"], D["tax"]["_meta"]["n_total"]
    for i, (v, lab, col) in enumerate([
            (f"{n_raw / 1e6:.2f}M", "BV-BRC 16S records", TEAL),
            (f"{n_amp:,}", "benchmark amplicons", TEAL),
            (f"{KBASE_BUILT:,}", "genomes in KBase (20-hit build)", TEAL),
            (f"{len(D['tax']['regions'])}", "16S regions", TEAL)]):
        stat(s, M + i * (sw + sg_), sy, sw, v, lab, color=col, h=1.4)
    notes(s, "Amplicon surveys capture more species per unit of DNA than shotgun sequencing, but "
             "a 16S read carries no genes: 16S tells you who is present, but GEM modeling "
             "needs gene content; MAGs are incomplete, expensive, and biased to "
             "abundant members; amplicons are cheap but functionally mute. prFBA maps each ASV "
             "to reference genome(s) via alignment scoring, keeps co-optimal candidates, and "
             "preserves the ambiguity as probabilities that are sorted out in optimization: the "
             "uncertainty in mapping an amplicon to a genome is real information, so prFBA "
             f"represents it instead of discarding it. The reference set is the BV-BRC 16S pull "
             f"of {n_raw:,} records, collapsed to {D['aln_ms']['n_unique_refs']:,} unique "
             f"sequences of at least 200 bp. The benchmark behind most of this deck is "
             f"{n_amp:,} in-silico amplicons from {len(D['sel3']['original']):,} genomes, in "
             f"{len(D['tax']['regions'])} regions of the 16S gene, and a deployment build has put "
             f"{KBASE_BUILT:,} synthetic genomes into KBase.")
    return s


def s03_pipeline(prs, D):
    s = blank(prs)
    p = D["aln_ms"]["params"]
    pk, k2 = p["prefilter_k"], p["k2"]
    title(s, "Five auditable steps from amplicon to model",
          sub="Every selection traces back to a published identity threshold")
    steps = [("Amplicon", "ASV or in-silico PCR"),
             ("Alignment", f"{pk} → {k2}-hit record"),
             ("Selection", "exact ties, per species"),
             ("Synthetic genome", "function union, p = k/n"),
             ("prFBA", "p-weighted objective")]
    x, y, w, h, gap = M, 2.75, W - 2 * M, 1.6, 0.42
    flow(s, x, y, w, h, steps, size=16, gap=gap)
    bw = (w - gap * (len(steps) - 1)) / len(steps)
    for i in range(len(steps)):
        d = rect(s, x + i * (bw + gap) + bw / 2 - 0.27, y - 0.27, 0.54, 0.54, TEAL,
                 shape=MSO_SHAPE.OVAL, line=WHITE)
        dtf = d.text_frame
        dtf.margin_left = dtf.margin_right = dtf.margin_top = dtf.margin_bottom = 0
        dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(dtf, str(i + 1), size=16, bold=True, color=WHITE, font=HEAD,
             align=PP_ALIGN.CENTER, first=True, space_after=0)
    tf = takeaway(s, "Every hit carries its disposition and reason", M, y + h + 0.55, W - 2 * M)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    footer(s, "align_hits.py → select_references.py → build_synthetic_genomes.py → KBase")
    notes(s, "Each amplicon, a real ASV or an in-silico PCR product in the benchmarks, is "
             f"aligned in three stages: an edlib edit-distance prefilter keeps the {pk} closest "
             f"of {D['aln_ms']['n_unique_refs']:,} references, Smith-Waterman scores them, and an "
             f"exact Biopython rescore emits a {k2}-hit record. Selection keeps every reference "
             "tied at the best identity, collapses them to one genome per species, and abstains "
             "below the family floor of 0.865 identity. The selected genomes are merged by "
             "function into one synthetic genome whose per-function weight p = k/n feeds the "
             "prFBA objective. The pipeline is hardware-adaptive: the prefilter runs "
             "identically everywhere and only the scoring backend changes (the CUDA kernel, or "
             "Biopython on CPU), so outputs are comparable across machines. Every hit carries "
             "its disposition and the reason for it in a JSON audit, so a selection can be "
             "checked rather than trusted, and every threshold it applies is a published one.")
    return s


def s04_alignment(prs, D):
    s = blank(prs)
    ms, gs, emb = D["aln_ms"], D["aln_gpu"], D["aln_emb"]
    p, rs = ms["params"], ms["run_stats"]
    pk, k2, gc = p["prefilter_k"], p["k2"], p["gpu_cand"]
    nre = max(gc, k2)                 # the aligner rescores the larger of the two
    full = nre >= pk                  # this run rescored its whole shortlist
    title(s, f"Exact alignment against all {ms['n_unique_refs']:,} references",
          sub=f"A {k2}-hit record: the GPU ranks a {pk}-candidate shortlist, Biopython "
              f"aligns the top {k2}" if full else
              f"A {k2}-hit record from a {pk}-candidate shortlist, best {nre} rescored")
    flow(s, M, BODY_TOP + 0.1, W - 2 * M, 1.3, [
        ("edlib prefilter", f"{ms['n_unique_refs']:,} → top {pk}"),
        ("GPU Smith-Waterman", f"ranks all {pk}" if full else f"keeps best {gc}"),
        ("Biopython alignment", f"top {k2}: identity, coordinates" if full else
         f"{nre} rescored → {k2} kept")], size=16)
    tot = ms["wall_seconds_total"]
    e_s, g_s = rs["wall_seconds_edlib_prefilter"], rs["wall_seconds_gpu_score"]
    now_s = e_s + g_s + STAGE3_S[1]   # stage 3 as it runs now, timed on the same records
    if full:
        cats = ["edlib prefilter", "GPU Smith-Waterman", "Biopython"]
        old = [e_s / 3600, g_s / 3600, STAGE3_S[0] / 3600]
        new = [e_s / 3600, g_s / 3600, STAGE3_S[1] / 3600]
        ch = _bar_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, M, 3.35, 6.9, 3.3, cats[::-1],
                        [(f"re-score all {pk}, then align", old[::-1]),
                         (f"align the GPU's top {k2} (now)", new[::-1])],
                        colors=[RGBColor(0xA7, 0xB8, 0xBC), TEAL], fmt='0.0" h"', vmin=0,
                        vmax=6.0, gap=45)
        chart_title(ch, f"Hours per stage, depth-{k2} records")
    else:
        stages = [("edlib prefilter", e_s), ("GPU Smith-Waterman", g_s),
                  ("Biopython rescore", rs["wall_seconds_biopython_rescore"])]
        ch = hbar(s, M, 3.35, 6.9, 3.3, [c for c, _ in stages], [v / tot for _, v in stages],
                  colors=[TEAL, SEAFOAM, SEAFOAM], fmt="0%", vmin=0, vmax=1, gap=45)
        chart_title(ch, "Share of wall time, by stage")
    gpu = gs["gpu"].replace("NVIDIA GeForce ", "")
    sx, sw = 7.85, (W - M - 7.85 - 0.3) / 2
    for i, (v, lab, col) in enumerate([
            (f"{gs['exhaustive_gcups']}", f"GCUPS on one {gpu}", TEAL),
            (f"{gs['speedup_vs_cpu_exhaustive']}×", "faster than CPU Biopython", TEAL),
            ((f"{BITEXACT_HITS / 1e6:.1f}M", "depth-100 hits bit-exact to Biopython", SEA_TXT)
             if full else (gs["correctness_exact_match"], "control pairs bit-exact", SEA_TXT)),
            ((f"{now_s / 3600:.1f} h", f"for {ms['n_asvs']:,} unique amplicons, as run now",
              SEA_TXT) if full else
             (ALIGN_STAT[0], f"for {ms['n_asvs']:,} unique amplicons", SEA_TXT))]):
        stat(s, sx + (i % 2) * (sw + 0.3), 3.4 + (i // 2) * 1.72, sw, v, lab, color=col,
             h=1.42)
    sc = gs["scoring"]
    sgn = lambda v: f"{v:+d}".replace("-", "−")   # noqa: E731
    footer(s, f"Same scoring on every backend: match {sgn(sc['match'])} / mismatch "
              f"{sgn(sc['mismatch'])} / gap {sgn(sc['gap_open'])}, {sgn(sc['gap_extend'])}")
    pct = lambda k: 100 * rs[k] / tot                   # noqa: E731
    stage3 = (f". Since 2026-10-27 nothing is re-scored: the GPU score is bit-exact to "
              f"Biopython's on all {BITEXACT_HITS:,} hits of the depth-{k2} records, so the "
              "GPU ranks the shortlist, under the order the re-score itself used (score, then "
              f"edit distance, then reference index), and Biopython only aligns the top {k2}, "
              "because the score-only kernel cannot give identity, aligned length or "
              "coordinates. Each alignment's score is checked against the GPU's, with a "
              "re-score of the best 200 as the fallback. Ranked this way the stage reproduces "
              f"all {ms['n_asvs']:,} records of the depth-{k2} run byte for byte, with no "
              f"fallback, in {STAGE3_S[1]:,} s instead of {STAGE3_S[0]:,} s on 40 workers: "
              f"that run re-scored all {pk} candidates (about {RESCORE_MS[0]:.1f} ms each) "
              f"before aligning {k2} (about {RESCORE_MS[1]:.0f} ms each), which made the "
              f"Biopython stage {pct('wall_seconds_biopython_rescore'):.0f}% of its wall time; "
              "about three fifths of that stage's growth from depth 20 was the re-score. "
              f"Summed over its stages, the pipeline as it runs now takes about "
              f"{now_s / 3600:.1f} h for these amplicons, against {tot / 3600:.1f} h."
              if full else
              f"; it rescores the {nre} best GPU candidates and keeps {k2}.")
    notes(s, f"Stage 1 scans all {ms['n_unique_refs']:,} unique references with edlib in infix "
             f"mode and keeps the {pk} closest per amplicon; it runs identically on every "
             f"machine and took {pct('wall_seconds_edlib_prefilter'):.0f}% of the "
             f"{ALIGN_STAT[0]} that the depth-{k2} run needed for {ms['n_asvs']:,} unique "
             "amplicons. Stage 2 scores the shortlist with the GPU Smith-Waterman kernel, "
             f"which on an exhaustive sweep reached {gs['exhaustive_gcups']} GCUPS on one "
             f"{gpu}, {gs['speedup_vs_cpu_exhaustive']} times a 60-process Biopython CPU "
             f"baseline, and is bit-exact to Biopython on {gs['correctness_exact_match']} "
             "control pairs. Stage 3 re-aligns candidates with Biopython to recover identity "
             "and coverage" + stage3 + " Retrieval is exact rather than approximate because a "
             "nucleotide-transformer embedding put the true best hit first only "
             f"{emb['top1_correct_pct']['cosine']}% of the time, against "
             f"{emb['top1_correct_pct']['kmer8']}% for a k-mer retriever, and random 16S pairs "
             f"already average cosine {D['aln_cos']:.2f}. The CUDA kernel is now the standalone "
             "gpusw package, reusable outside this project.")
    return s


def s05_selection(prs, D):
    s = blank(prs)
    k2 = D["sr_k2"]                   # the depth of the self-recovery files read below
    title(s, "Selection keeps the ambiguity it finds",
          sub="Exact ties kept, one genome per species, abstention below family")
    regions = ["FullLength16S", "V1-V2", "V1-V3", "V3-V4", "V4", "V4-V5", "V4-V5_944R",
               "V7-V9"]
    L, U = D["sr_leg"]["per_region"], D["sr_uni"]["per_region"]
    takeaway(s, f"Source kept, when it is in its {k2}-hit record", M, BODY_TOP + 0.05,
             W - 2 * M)
    ch = vbar(s, M, BODY_TOP + 0.45, W - 2 * M, 3.5, [REG_LABEL[r] for r in regions],
              [("Legacy (ordered) reducer", [L[r]["selected_given_top20"] for r in regions]),
               ("Exact-tie union (shipped)", [U[r]["selected_given_top20"] for r in regions])],
              colors=[MUTED, TEAL], fmt="0.0%", vmin=0, vmax=1.12, gap=70)
    ch.value_axis.major_unit = 0.25
    ch.value_axis.tick_labels.number_format = "0%"
    ch.value_axis.tick_labels.number_format_is_linked = False
    # the identity ladder that sets the call
    ly, lh = 5.95, 0.62
    para(textbox(s, M, ly - 0.36, W - 2 * M, 0.3), "Best-hit identity sets the call",
         size=12, color=MUTED, first=True, space_after=0)
    bands = [("< 0.865  abstain", RGBColor(0xFB, 0xE9, 0xD2), AMBER_TXT, 2.6),
             ("≥ 0.865  family", RGBColor(0xD5, 0xEC, 0xEC), INK, 3.0),
             ("≥ 0.945  genus", SEAFOAM, INK, 3.0),
             ("≥ 0.987  species", TEAL, WHITE, W - 2 * M - 8.6)]
    x = M
    for lab, fill, col, bw in bands:
        b = rect(s, x, ly, bw - 0.06, lh, fill, adj=0.12)
        btf = b.text_frame
        btf.margin_top = btf.margin_bottom = 0
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(btf, lab, size=14, bold=True, color=col, align=PP_ALIGN.CENTER, first=True,
             space_after=0)
        x += bw
    n_in = U["ALL"]["self_in_top20"]
    footer(s, f"Original panel · self-inclusion · {k2}-hit records · {n_in:,} amplicons whose "
              "record holds the source or its taxon")
    prec = {}                          # V4 self-inclusion gene-capture precision, both rules
    for t, tree in (("u", "region_validation_d100"), ("l", "region_validation_d100_legacy")):
        with open(os.path.join(REPO, tree, "data", "gene_capture_by_region.json")) as fh:
            prec[t] = json.load(fh)["regions"]["V4"]["include"]["macro"]["precision"]["value"]
    nf = D["sr_uni"]["n_failures"]
    dedup = (f"; every one of the union's {nf:,} misses is species de-duplication keeping a "
             "conspecific assembly" if D["sr_uni"]["failure_terminal_stage_overall"]
             == {"species_dedup": nf} else f"; the union misses it {nf:,} times")
    notes(s, "Each amplicon is first classified by its best-hit identity against the "
             "Yarza/Stackebrandt bands (species 0.987, genus 0.945, family 0.865); below the "
             "family floor it abstains with a taxonomy-only record, so no false metabolism is "
             "injected, and a reliability gate drops hits with coverage under 0.90 or an "
             "SW-edlib identity discordance above 0.02. The shipped exact-tie union keeps every "
             "reference tied at the best identity, then one genome per BV-BRC species (the "
             "highest-identity assembly, the gene-richest on a further tie), so a densely "
             "sequenced species cannot flood the union, and gene count never changes the "
             "taxonomic call; the first survivor, the anchor, makes that call. On the original "
             f"panel with {k2}-hit records the union keeps the source genome for "
             f"{U['ALL']['selected_given_top20']:.1%} of the {n_in:,} amplicons whose record "
             "holds it or a same-taxon assembly "
             f"({U['FullLength16S']['selected_given_top20']:.1%} on full-length 16S), against "
             f"{L['ALL']['selected_given_top20']:.1%} for the earlier legacy (ordered) reducer, "
             f"whose order-dependent stages discarded it {D['sr_leg']['n_failures']:,} times"
             + dedup + ". Keeping ties has a cost with the source present: a "
             f"{k2}-hit record lets the union grow, and V4 gene-capture precision falls to "
             f"{prec['u']:.3f} against {prec['l']:.3f} for the legacy reducer. The slides "
             "that follow use the exact-tie union throughout, the rule users run; the legacy "
             "reducer stays available as an option. Tie sets can "
             f"also outgrow the record: for {TIE_SAT_D100[0]} of {TIE_SAT_D100[1]:,} distinct "
             f"sequences all 100 retained hits are tied at the best score (median "
             f"{TIE_SAT_D100[2]} species), and for {TIE_ALL500} the whole 500-candidate "
             "shortlist is tied, so their probabilities come from a truncated sample of the "
             "co-optimal genomes.")
    return s


def s06_synthetic(prs, D):
    s = blank(prs)
    title(s, "Each function is weighted by its share of references",
          sub="The synthetic genome: the union of the selected genomes' functions")
    # a worked example: three selected references, four functions
    has = [(1, 1, 1, 0), (1, 1, 0, 0), (1, 0, 1, 1)]
    k = [sum(c) for c in zip(*has)]
    n = len(has)
    frac = {1: "1/3", 2: "2/3", 3: "1"}
    x0, y0, cs, cg, lw = 2.1, 1.95, 0.85, 0.14, 1.35

    def cell(x, y, w, fill, text="", col=WHITE, size=15, bold=True):
        c = rect(s, x, y, w, cs, fill, adj=0.14)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if text:
            para(tf, text, size=size, bold=bold, color=col, align=PP_ALIGN.CENTER, first=True,
                 space_after=0)
        return c

    for r, row in enumerate(has):
        y = y0 + r * (cs + cg)
        cell(x0 - lw - 0.15, y, lw, LIGHT, f"genome {r + 1}", col=INK, size=14)
        for c, v in enumerate(row):
            cell(x0 + c * (cs + cg), y, cs, TEAL if v else LIGHT, "ABCD"[c] if v else "")
    yp = y0 + n * (cs + cg) + 0.38
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 - lw - 0.15), Inches(yp - 0.24),
                              Inches(lw + 0.15 + 4 * cs + 3 * cg), Inches(0.03))
    rule.fill.solid()
    rule.fill.fore_color.rgb = MUTED
    rule.line.fill.background()
    rule.shadow.inherit = False
    cell(x0 - lw - 0.15, yp, lw, DEEP, "synthetic", size=14)
    shade = {3: TEAL, 2: SEAFOAM, 1: MINT}
    for c, kc in enumerate(k):
        cell(x0 + c * (cs + cg), yp, cs, shade[kc], f"p = {frac[kc]}", size=14,
             col=WHITE if kc == n else INK)
    # the rule, the deposit and where the weight goes
    rx, rw = 7.0, W - M - 7.0
    fc = rect(s, rx, 1.95, rw, 1.55, DEEP)
    ftf = fc.text_frame
    ftf.margin_left = ftf.margin_right = Inches(0.25)
    ftf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ftf, "p(function) = k / n", size=28, bold=True, color=WHITE, font=HEAD, first=True,
         space_after=4)
    para(ftf, "k of the n selected genomes carry it", size=14, color=MINT, space_after=0)
    sw = (rw - 0.3) / 2
    stat(s, rx, 3.8, sw, f"{KBASE_BUILT:,}", "synthetic genomes in KBase", h=1.45)
    stat(s, rx + sw + 0.3, 3.8, sw, f"{KBASE_ABSTAIN}", f"of {KBASE_ASVS:,} ASVs abstained",
         color=AMBER_TXT, h=1.45)
    takeaway(s, "Reaction weight = max over its genes", rx, 5.55, rw)
    footer(s, f"KBase build on 20-hit records, before the 100-hit standard · {KBASE_ASVS:,} "
              f"real ASVs · workspace {KBASE_WS}")
    notes(s, "For each answered ASV the selected references are merged into one synthetic KBase "
             "Genome keyed on gene function: the first genome to carry a function creates the "
             "merged feature and later genomes only add a vote, so p(function) = k/n is the "
             "share of the n selected genomes that carry it; core functions get 1, accessory "
             "functions less, and a single reference gives 1 for every function. Each function "
             "enters the ModelSEED draft as a gene-associated reaction, the reaction takes the "
             "largest weight among its genes, and prFBA penalizes flux through low-weight "
             "reactions; the weight never gates or removes a reaction. The weight is an "
             "annotation prevalence, not a calibrated probability, and the calibration slide "
             "measures how well it serves as one. Each genome ships with provenance: git commit, "
             f"input hashes, per-stage parameters, and the abstention reason when no model was "
             f"built. The KBase build made {KBASE_BUILT:,} synthetic genomes and a GenomeSet in "
             f"workspace {KBASE_WS} from {KBASE_ASVS:,} real ASVs, {KBASE_ABSTAIN} of which "
             "abstained rather than guess; it ran on 20-hit records, before the 100-hit "
             "standard, so its built and abstained counts are unlikely to move but its unions "
             "and weights would. Whole genomes take the same path: a MAG enters "
             "selection through skani average nucleotide identity instead of 16S identity, with "
             "thresholds recalibrated for genomes (95 ANI same species, 85 to 95 putatively "
             "species-distinct) and an aligned-fraction guard, and only BV-BRC matches go on to "
             "reconstruction; the MAG arm has its own slides later.")
    return s


def s07_validation(prs, D):
    s = blank(prs)
    title(s, "Without the source, family transfers best",
          sub="Every amplicon's true genome and gene set are known")
    tax, meta = D["tax"]["regions"], D["tax"]["_meta"]
    cap = D["cap"]["_meta"]["headline_macro_recall_pct"]
    regs = list(tax)
    rule = "legacy reducer" if meta["select_mode"] == "legacy" else "exact-tie union"
    acc = lambda m, rank: [                                                            # noqa: E731
        tax[r]["anchor"][m]["all"]["all"]["per_rank"][rank]["correct"] for r in regs]
    rng = lambda v, d=2: f"{rd(min(v), d)}–{rd(max(v), d)}"                           # noqa: E731
    pct = lambda m: rng([v / 100 for v in cap[m].values()])                             # noqa: E731
    prose = lambda r: "full-length 16S" if r == "FullLength16S" else REG_LABEL[r]      # noqa: E731
    sel = json.load(open(os.path.join(AMP_RV, "data", "selection_10k.json")))
    flow(s, M, BODY_TOP, W - 2 * M, 0.82, [
        (f"{len(sel['genomes']):,} genomes", "taxonomy-stratified"),
        ("In-silico PCR", f"{len(regs)} regions"),
        (f"{meta['n_total']:,} amplicons", "unaltered pipeline"),
        ("Score vs truth", "NCBI lineage, PGFams")], size=15)
    # the grid sits on the flow's columns: labels under step 1, values under steps 2-4
    fg = 0.42
    fb = (W - 2 * M - 3 * fg) / 4
    y0, rh, gap, lw, cw = BODY_TOP + 1.75, 1.35, 0.2, fb, fb
    cx = [M + (j + 1) * (fb + fg) for j in range(3)]
    for x, head in zip(cx, ("Genus accuracy", "Gene capture", "Family accuracy")):
        chip = rect(s, x, y0 - 0.5, cw, 0.38, TEAL)
        tf = chip.text_frame
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, head, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True,
             space_after=0)
    for i, (head, sub, col, mode) in enumerate((
            ("Source present", None, TEAL, "include"),
            ("Source removed", "among answered", AMBER_TXT, "exclude_0.987"))):
        y = y0 + i * (rh + gap)
        lab = rect(s, M, y, lw, rh, LIGHT)
        tf = lab.text_frame
        tf.margin_left = tf.margin_right = Inches(0.18)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, head, size=19, bold=True, color=col, font=HEAD, first=True,
             space_after=2 if sub else 0)
        if sub:
            para(tf, sub, size=13, color=MUTED, space_after=0)
        for x, v in zip(cx, (rng(acc(mode, "Genus")), pct(mode), rng(acc(mode, "Family")))):
            card = stat(s, x, y, cw, v, "", color=col, h=rh, vsize=32)
            lbl = card.text_frame.paragraphs[1]._p          # value only: the column chip names it
            lbl.getparent().remove(lbl)
            card.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    mc = D["mcn"]["by"]["anchor|Genus|exclude_0.987"]["pairs"]
    ns = [p for p in mc if not p["significant"]]
    footer(s, f"Anchor predictor · {rule} · {DEPTH}-hit records · gene capture = macro recall")
    fam_gap = [f - g for f, g in zip(acc("exclude_0.987", "Family"), acc("exclude_0.987", "Genus"))]
    hp = [p["p_holm"] for p in ns]
    notes(s, (
        f"The in-silico validation runs the unaltered pipeline on amplicons whose true genome "
        f"and gene set are known: {len(sel['genomes']):,} taxonomy-stratified BV-BRC genomes "
        f"(deterministic, seed {sel['meta']['seed']}), in-silico PCR for {len(regs)} regions from "
        f"V1–V2 to V7–V9 with full-length 16S as the positive control, {meta['n_total']:,} "
        "amplicons in all. Each amplicon is scored two ways from one run: self-inclusion leaves "
        "the source organism in the database, the realistic upper bound, while self-exclusion "
        "removes the source genome, its taxon and every hit at identity 0.987 or above, then "
        "re-selects, as if the organism were novel. Taxonomic accuracy is scored against the "
        "NCBI lineage and gene capture (macro recall) against the source genome's PGFam set. "
        f"With the source present, genus accuracy is {rng(acc('include', 'Genus'))}, family "
        f"{rng(acc('include', 'Family'))} and gene capture {pct('include')} (macro recall) "
        "across the regions; "
        "with it removed, among amplicons that still answer, genus accuracy is "
        f"{rng(acc('exclude_0.987', 'Genus'))}, gene capture {pct('exclude_0.987')} and family "
        f"accuracy {rng(acc('exclude_0.987', 'Family'))}, {rd(100 * min(fam_gap), 0)}–"
        f"{rd(100 * max(fam_gap), 0)} points above genus: family-level metabolism transfers "
        "even for novel organisms, and genus is where short amplicons pay. The whole-population "
        "figures, which count every non-answer as a miss, are on the denominators slide. The "
        f"paired McNemar test (Holm across {len(mc)} pairs) finds genus accuracy different in "
        f"{len(mc) - len(ns)} of {len(mc)} region pairs; the {len(ns)} that do not differ are "
        + ", ".join(("and " if k and k == len(ns) - 1 else "")
                    + f"{prose(p['region_a'])} against {prose(p['region_b'])}"
                    for k, p in enumerate(ns))
        + (f" (Holm p {rd(min(hp), 2)}–{rd(max(hp), 2)})" if hp else "") + ". "
        f"Every figure uses the {rule} on {DEPTH}-hit records, the basis of the manuscript's "
        "tab:main (taxacc_region_summary.json, gene_capture_by_region.json and "
        "mcnemar_regions.json)."))
    return s


def s08_regions(prs, D):
    s = blank(prs)
    title(s, "Region costs genus accuracy, not gene capture",
          sub="Points lost when the source organism is removed, by region")
    tax, meta = D["tax"]["regions"], D["tax"]["_meta"]
    cap = D["cap"]["_meta"]["headline_macro_recall_pct"]
    rule = "legacy reducer" if meta["select_mode"] == "legacy" else "exact-tie union"
    regs = list(tax)
    v1 = [r for r in regs if r == "FullLength16S" or r.startswith("V1")]
    v3 = [r for r in regs if r not in v1]
    order = v1 + v3
    gen = lambda r, m: (                                                               # noqa: E731
        tax[r]["anchor"][m]["all"]["all"]["per_rank"]["Genus"]["correct"])
    g_lost = [100 * (gen(r, "include") - gen(r, "exclude_0.987")) for r in order]
    c_lost = [cap["include"][r] - cap["exclude_0.987"][r] for r in order]
    series = [("Genus accuracy", g_lost), ("Gene capture", c_lost)]
    ch = vbar(s, M, BODY_TOP + 0.62, W - 2 * M, 4.5, [REG_LABEL[r] for r in order], series,
              fmt="0.0", vmin=0, vmax=40, colors=[AMBER, TEAL])
    # group the regions by where the amplicon starts: the split is V1-anchored vs V3 onward
    cd = CategoryChartData(number_format="0.0")
    for grp, rs in (("Full-length + V1-anchored", v1), ("V3 onward", v3)):
        c = cd.add_category(grp)
        for r in rs:
            c.add_sub_category(REG_LABEL[r])
    for name, vals in series:
        cd.add_series(name, vals, number_format="0.0")
    ch.replace_data(cd)
    gv1, gv3 = g_lost[:len(v1)], g_lost[len(v1):]
    r0 = lambda v: f"{rd(min(v), 0)}–{rd(max(v), 0)}"                                  # noqa: E731
    prose = lambda r: "full-length 16S" if r == "FullLength16S" else REG_LABEL[r]      # noqa: E731
    flat = max(c_lost) - min(c_lost) <= 2                    # "flat": within two points
    takeaway(s, (f"Gene capture flat at ~{rd(sum(c_lost) / len(c_lost), 0)} points" if flat
                 else f"Gene capture {r0(c_lost)} points")
                + "; genus loss " + ("jumps from V3 on" if min(gv3) > max(gv1)
                                     else "varies by region"),
             M, BODY_TOP, W - 2 * M, h=0.5)
    footer(s, f"Anchor predictor · {rule} · {DEPTH}-hit records · loss = source present − "
              "source removed, among amplicons that answer")
    ex = [gen(r, "exclude_0.987") for r in order]
    notes(s, (
        "Each pair of bars is how much a measure falls when the source organism is removed: "
        "genus accuracy in points (self-inclusion minus self-exclusion, among amplicons that "
        "answer) and gene capture (macro recall) in percentage points. Gene capture loses "
        f"{rd(min(c_lost), 1)}–{rd(max(c_lost), 1)} points across the regions, while the genus "
        f"loss is {r0(gv1)} points for full-length 16S and the two V1-anchored regions and "
        f"{r0(gv3)} points for the {len(v3)} regions that start at V3 or later. V1–V2, the "
        "second-shortest region, generalizes like V1–V3, so the split is by where the amplicon "
        "starts, not by its length. In absolute terms, conditional genus accuracy with the source "
        f"removed runs from {rd(max(ex), 2)} ({prose(order[ex.index(max(ex))])}) down to "
        f"{rd(min(ex), 2)} ({prose(order[ex.index(min(ex))])}). Region therefore costs "
        "taxonomic resolution more than it costs raw functional recovery. Anchor predictor, "
        f"Bacteria and Archaea pooled, {rule}, {DEPTH}-hit records (manuscript fig:leakage); "
        "source taxacc_region_summary.json and gene_capture_by_region.json."))
    return s


def s09_denominators(prs, D):
    import csv
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    s = blank(prs)
    tax, meta = D["tax"]["regions"], D["tax"]["_meta"]
    rule = "legacy reducer" if meta["select_mode"] == "legacy" else "exact-tie union"
    regs = list(tax)
    dd = os.path.join(AMP_RV, "data")
    # genus outcome of every amplicon under self-exclusion; the truth decides scoreability
    truth = json.load(open(os.path.join(dd, "benchmark_truth.json")))
    named = lambda v: bool(v) and not str(v).startswith("unclassified")  # noqa: E731
    k = dict(correct=0, wrong=0, unnamed=0, abstained=0, unscoreable=0)
    with open(os.path.join(dd, "asv_region_concordance.csv")) as fh:
        for r in csv.DictReader(fh):
            if r["self_mode"] != "exclude_0.987" or r["predictor"] != "anchor":
                continue
            if not named(truth[r["amplicon"]]["truth_lineage"].get("Genus")):
                k["unscoreable"] += 1
            elif r["abstained"] == "1":
                k["abstained"] += 1
            else:
                k[{"1": "correct", "0": "wrong"}.get(r["correct_Genus"], "unnamed")] += 1
    n = k["correct"] + k["wrong"] + k["unnamed"] + k["abstained"]
    gr = lambda r: (                                                                   # noqa: E731
        tax[r]["anchor"]["exclude_0.987"]["all"]["all"]["per_rank"]["Genus"])
    assert n == sum(gr(r)["n"] for r in regs), "genus-scoreable count disagrees with the summary"
    rng = lambda v, d=2: f"{rd(min(v), d)}–{rd(max(v), d)}"                             # noqa: E731
    cond = rng([gr(r)["correct"] for r in regs])
    popu = rng([gr(r)["correct_population"] for r in regs])
    unsc = [gr(r)["n_truth_undefined"] / (gr(r)["n"] + gr(r)["n_truth_undefined"]) for r in regs]
    title(s, "The honest denominator counts every non-answer",
          sub=f"Genus, source removed: {n:,} scoreable amplicons, {len(regs)} regions pooled")

    hmax, yb, cw = 3.3, BODY_TOP + 3.45, 3.4
    gapx = (W - 2 * M - 3 * cw) / 2
    xs = [M + i * (cw + gapx) for i in range(3)]
    pc = lambda key: f"{rd(100 * k[key] / n, 0)}%"  # noqa: E731
    seg_fill = {"correct": TEAL, "wrong": AMBER, "unnamed": RGBColor(0xF6, 0xD3, 0xA8),
                "abstained": RGBColor(0xC9, 0xD6, 0xD9)}
    cols = [
        [("correct", "correct"), ("wrong", "wrong genus")],
        [("correct", f"correct · {pc('correct')}"), ("wrong", f"wrong genus · {pc('wrong')}"),
         ("unnamed", f"unnamed genus · {pc('unnamed')}"),
         ("abstained", f"abstained · {pc('abstained')}")],
    ]
    for x, segs in zip(xs, cols):
        y = yb
        for key, lab in segs:
            h = hmax * k[key] / n
            y -= h
            box = rect(s, x, y, cw, h, seg_fill[key], shape=MSO_SHAPE.RECTANGLE)
            tf = box.text_frame
            tf.margin_left = tf.margin_right = Inches(0.08)
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(tf, lab, size=14 if key in ("correct", "wrong") else 13,
                 bold=key == "correct", color=WHITE if key == "correct" else INK,
                 align=PP_ALIGN.CENTER, first=True, space_after=0)
    h = hmax * k["unscoreable"] / n
    box = rect(s, xs[2], yb - h, cw, h, WHITE, shape=MSO_SHAPE.RECTANGLE, line=MUTED)
    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "not scoreable", size=14, color=MUTED, align=PP_ALIGN.CENTER, first=True,
         space_after=0)
    for x, v, lab, col in ((xs[0], cond, "conditional: ÷ answered", TEAL),
                           (xs[1], popu, "population: ÷ all scoreable", TEAL),
                           (xs[2], f"{rd(100 * min(unsc), 0)}–{rd(100 * max(unsc), 0)}%",
                            "excluded: no named genus in truth", MUTED)):
        stat(s, x, yb + 0.25, cw, v, lab, color=col, h=1.2)
    footer(s, f"Anchor predictor · {rule} · {DEPTH}-hit records · accuracy ranges by region; "
              "columns pooled, drawn to one scale")
    gcap = D["cap"]["_meta"]
    pop_cap = gcap["headline_macro_recall_pct_population"]["exclude_0.987"]
    res_cap = gcap["headline_macro_recall_pct"]["exclude_0.987"]
    oc = gcap["outcome_counts_total"]["exclude_0.987"]
    ab_all = sum(tax[r]["abstention"]["exclude_0.987"]["all"]["all"]["n_abstained"] for r in regs)
    cov = [gr(r)["coverage"] for r in regs]
    notes(s, (
        "Self-exclusion accuracy has two denominators, and both count only amplicons the truth "
        f"can score. Conditional accuracy ({cond} by region) divides correct genus calls by the "
        "amplicons that still resolve to a reference named at genus. Whole-population accuracy "
        f"({popu}) divides by every genus-scoreable amplicon and counts two kinds of non-answer "
        f"as misses: the {rd(100 * k['abstained'] / n, 1)}% of genus-scoreable amplicons that "
        "receive no taxonomic call once their species-level relatives are removed "
        f"({rd(100 * ab_all / meta['n_total'], 1)}% of all {meta['n_total']:,} amplicons), and a "
        f"further {rd(100 * k['unnamed'] / n, 1)}% whose selected reference has no named genus, "
        f"so genus coverage is {rng(cov)} by region. Amplicons whose own truth has no named NCBI "
        f"genus ({rd(100 * min(unsc), 0)}–{rd(100 * max(unsc), 0)}% by region, "
        f"{k['unscoreable']:,} pooled) are unscoreable and leave both denominators instead of "
        "counting as wrong. That exclusion is one of two scorer corrections, which moved the "
        "published generalization numbers in opposite directions: counting unscoreable ranks as "
        "prediction failures had made population accuracy pessimistic, and dropping abstentions "
        "from the gene-capture mean instead of scoring them zero had inflated capture. With "
        "abstentions scored zero, population gene capture is "
        f"{rd(min(pop_cap.values()), 0)}–{rd(max(pop_cap.values()), 0)}% against "
        f"{rd(min(res_cap.values()), 0)}–{rd(max(res_cap.values()), 0)}% among answered "
        f"amplicons ({rd(100 * oc['abstained'] / (oc['ok'] + oc['abstained']), 1)}% of scoreable "
        "amplicons keep no genome). Both corrections came from re-deriving each metric from its "
        "own definition, and both are now unit-checked in the scorer. Figures: "
        f"{rule}, {DEPTH}-hit records, from "
        "asv_region_concordance.csv, benchmark_truth.json, taxacc_region_summary.json and "
        "gene_capture_by_region.json."))
    return s


def s10_benchmark_design(prs, D):
    s = blank(prs)
    title(s, "Same amplicons, same truth, one shared projection",
          sub="PICRUSt2 against prFBA, each querying its own reference database")
    m, p2 = D["cmp"]["_meta"], D["cmp"]["_meta"]["picrust2"]
    n = D["cmp"]["by_region"]["ec"]["all"]["all"]["picrust2"]["n"]
    table(s, M, BODY_TOP, W - 2 * M, [
        ["", f"PICRUSt2 {p2['version']}", "prFBA"],
        ["Reference", f"{p2['n_reference_genomes']:,} genomes", "~315,000 BV-BRC genomes"],
        ["No close match", "always answers", "abstains below family"],
    ], col_w=[2.75, 4.69, 4.693], size=17, row_h=0.62)
    takeaway(s, "Truth and both predictions pass through one map", M, BODY_TOP + 2.35,
             W - 2 * M, h=0.5)
    flow(s, M, BODY_TOP + 2.95, W - 2 * M, 1.3, [
        ("Gene family", "prFBA and truth"),
        ("Product string", "BV-BRC"),
        ("EC number", f"{m['ec_universe_shared']:,} shared by both"),
        ("ModelSEED reaction", f"{m['rxn_universe_shared']:,} reached")], size=16)
    fb = (W - 2 * M - 3 * 0.42) / 4
    chip = rect(s, M + 2 * (fb + 0.42), BODY_TOP + 2.95 + 1.3 + 0.15, fb, 0.42,
                RGBColor(0xFB, 0xE9, 0xD2), adj=0.3)
    ctf = chip.text_frame
    ctf.margin_top = ctf.margin_bottom = 0
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(ctf, "PICRUSt2 enters here", size=13, bold=True, color=AMBER_TXT,
         align=PP_ALIGN.CENTER, first=True, space_after=0)
    footer(s, f"EC space: {n:,} scoreable amplicons; {m['n_amplicons'] - n:,} whose source genome "
              "has no truth EC set are excluded")
    notes(s, (
        f"PICRUSt2 {p2['version']} places each amplicon in a fixed reference phylogeny with "
        "EPA-ng and gappa after HMMER profile alignment, infers gene content by maximum-parsimony "
        f"ancestral-state reconstruction with castor over {p2['n_reference_genomes']:,} reference "
        "genomes, and always returns a gene complement. prFBA aligns each amplicon to named "
        "reference genomes, about 315,000 BV-BRC genomes behind 459,301 unique 16S sequences, "
        "unions the gene content those genomes actually carry, and abstains below the 0.865 "
        f"family floor. Both run on the same {m['n_amplicons']:,} benchmark amplicons and are "
        "scored against the same truth. prFBA predicts PGFam gene families and PICRUSt2 predicts "
        "EC numbers, so the truth and both predictions pass through one map: a family's BV-BRC "
        "product string yields its EC numbers, and EC numbers yield ModelSEED reactions through "
        f"the ModelSEED Biochemistry aliases; metrics use the {m['ec_universe_shared']:,} ECs both "
        f"systems can express and the {m['rxn_universe_shared']:,} reactions those reach. "
        "Family-level products lose some EC numbers, and because truth and prediction go through "
        f"the same map, that loss cancels. The {m['n_amplicons'] - n:,} amplicons whose source "
        "genome has no gene-family annotation, and so no truth EC set, are unscoreable and "
        "excluded. Two asymmetries are built in rather than controlled: each method queries its "
        "own database, roughly tenfold apart, and the truth is BV-BRC annotation, prFBA's own "
        "vocabulary; the controlled comparison later removes both. PICRUSt2 was installed "
        "without conda (uv venv, user-space R and castor, epa-ng, gappa and HMMER from source), "
        "and 60 of its 61 tests pass."))
    return s


def s11_results(prs, D):
    import csv
    import gzip
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    s = blank(prs)
    title(s, "More precise when it answers, narrowly behind overall",
          sub="As shipped, scored against BV-BRC truth: prFBA's own annotation")
    cm = D["cmp"]
    ec = cm["by_region"]["ec"]
    regs = [r for r in ec if r != "all"]
    allc = ec["all"]["all"]
    p2, pf = allc["picrust2"], allc["prfba_exclude"]
    rule = cm["_meta"]["prfba_selection"]["default"].split("(")[-1].rstrip(")")
    data_dir = os.path.dirname(CMP_SUMMARY)
    # per-amplicon cells: both-answer F1 overall and by region, legacy population F1, and the
    # full-precision conditional means quoted in the notes (the summary keeps 4 decimals)
    cells, pop_leg, mean_ok = {}, [], {}
    with gzip.open(os.path.join(data_dir, "comparison_cells.csv.gz"), "rt") as fh:
        rdr = csv.reader(fh)
        ix = {c: i for i, c in enumerate(next(rdr))}
        for r in rdr:
            if r[ix["outcome"]] == "truth_empty":
                continue
            meth, ok = r[ix["method"]], r[ix["outcome"]] == "ok"
            if ok and meth in ("picrust2", "prfba_exclude"):
                for q in ("precision", "f1"):
                    mean_ok.setdefault((r[ix["space"]], meth, q), []).append(float(r[ix[q]]))
            if r[ix["space"]] != "ec":
                continue
            if meth == "prfba_exclude_legacy":
                pop_leg.append(float(r[ix["f1"]]) if ok else 0.0)
            elif meth in ("picrust2", "prfba_exclude"):
                cells.setdefault(r[ix["amplicon"]], {"region": r[ix["region"]]})[meth] = (
                    float(r[ix["f1"]]) if ok else None)
    both = {}
    for c in cells.values():
        if c.get("prfba_exclude") is not None and c.get("picrust2") is not None:
            both.setdefault(c["region"], []).append((c["prfba_exclude"], c["picrust2"]))
    nb = sum(len(v) for v in both.values())
    bf = sum(a for v in both.values() for a, _ in v) / nb
    bp = sum(b for v in both.values() for _, b in v) / nb
    # outright abstentions: the selection keeps no genome for an EC-scoreable amplicon
    sel = json.load(open(os.path.join(data_dir, "prfba_selection.json")))
    n_abst = sum(1 for a in cells if sel[a]["exclude_0.987"]["n_selected"] == 0)
    no_ec, abst = 1 - pf["coverage"], n_abst / pf["n"]
    gap = pf["population"]["f1"]["value"] - p2["population"]["f1"]["value"]
    mo = {key: sum(v) / len(v) for key, v in mean_ok.items()}
    cw = (W - 2 * M - 2 * 0.3) / 3
    for i, (v, lab, col) in enumerate((
            (f"{rd(bf, 3)} vs {rd(bp, 3)}", f"EC F1 when both answer, n = {nb:,}", TEAL),
            (f"{rd(100 * no_ec, 1)}%", f"get no EC prediction; {rd(100 * abst, 1)}% abstain",
             AMBER_TXT),
            (sg(gap, 3), "whole-population F1 gap to PICRUSt2", SEA_TXT))):
        stat(s, M + i * (cw + 0.3), BODY_TOP, cw, v, lab, color=col, h=1.15)
    v1 = [r for r in regs if r == "FullLength16S" or r.startswith("V1")]
    v3 = [r for r in regs if r not in v1]
    cond = [sum(a - b for a, b in both[r]) / len(both[r]) for r in v1 + v3]
    popd = [ec[r]["all"]["prfba_exclude"]["population"]["f1"]["value"]
            - ec[r]["all"]["picrust2"]["population"]["f1"]["value"] for r in v1 + v3]
    ahead = [REG_LABEL[r] for r, d in zip(v1 + v3, popd) if d > 0]
    takeaway(s, ("Whole population: ahead in " + ", ".join(ahead) + " only")
             if ahead else "Whole population: PICRUSt2 ahead in every region",
             M, BODY_TOP + 1.35, W - 2 * M, h=0.45)
    fmt = '+0.00;"−"0.00;0.00'                 # quoted: Excel allows few bare literals
    series = [("prFBA − PICRUSt2, both answer", cond),
              ("prFBA − PICRUSt2, whole population", popd)]
    ch = vbar(s, M, BODY_TOP + 1.8, W - 2 * M, 3.45, [REG_LABEL[r] for r in v1 + v3], series,
              fmt=fmt, vmin=-0.08, vmax=0.12, colors=[TEAL, MUTED])
    cd = CategoryChartData(number_format=fmt)
    for grp, rs in (("Full-length + V1-anchored", v1), ("V3 onward", v3)):
        c = cd.add_category(grp)
        for r in rs:
            c.add_sub_category(REG_LABEL[r])
    for name, vals in series:
        cd.add_series(name, vals, number_format=fmt)
    ch.replace_data(cd)
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    footer(s, f"Shipped {rule} · {DEPTH}-hit records · EC space, n = {pf['n']:,} scoreable "
              "amplicons")
    pt = cm["paired_tests_ec_f1"]
    pc = [pt[f"{r}|prfba_exclude_vs_picrust2|conditional"] for r in regs]
    wins = [t["prfba_wins"] / t["picrust2_wins"] for t in pc]
    pr = {r: ec[r]["all"] for r in regs}
    prose = lambda r: "full-length 16S" if r == "FullLength16S" else REG_LABEL[r]      # noqa: E731
    lead = ", ".join(f"{prose(r)} ({rd(pr[r]['prfba_exclude']['population']['f1']['value'], 3)}"
                     f" against {rd(pr[r]['picrust2']['population']['f1']['value'], 3)})"
                     for r in regs if pr[r]["prfba_exclude"]["population"]["f1"]["value"]
                     > pr[r]["picrust2"]["population"]["f1"]["value"])
    notes(s, (
        f"As the two systems ship, scored in the shared EC space against BV-BRC truth "
        f"(n = {pf['n']:,} scoreable amplicons), they trade coverage against precision. "
        f"PICRUSt2 answers all but {p2['n'] - p2['n_called']:,} amplicons, at EC precision "
        f"{rd(mo[('ec', 'picrust2', 'precision')], 3)}; prFBA, under its shipped {rule} with "
        f"the source organism excluded, gives an EC prediction for {rd(100 * pf['coverage'], 1)}% "
        f"at precision {rd(mo[('ec', 'prfba_exclude', 'precision')], 3)}: "
        f"{rd(100 * no_ec, 1)}% get none, {rd(100 * abst, 1)}% because the selection abstains "
        f"outright. On the {nb:,} amplicons both answer, conditional EC F1 is {rd(bf, 3)} against "
        f"{rd(bp, 3)}, prFBA leads in all {len(regs)} regions (paired Wilcoxon, median ΔF1 "
        f"{sg(min(t['median_diff'] for t in pc), 3)} to {sg(max(t['median_diff'] for t in pc), 3)}"
        f", Holm p ≈ 0) and wins {rd(min(wins), 0)} to {rd(max(wins), 0)} amplicons per loss, "
        "depending on the region. Over the whole population, scoring prFBA's non-answers as "
        f"zero, PICRUSt2's mean F1 is narrowly higher ({rd(p2['population']['f1']['value'], 3)} "
        f"against {rd(pf['population']['f1']['value'], 3)}, a gap of {sg(gap, 3)}; under the "
        f"legacy reducer prFBA reaches {rd(sum(pop_leg) / len(pop_leg), 3)}, within resampling "
        "noise), while the median amplicon still favours prFBA. "
        + (f"By region the population mean favours prFBA in {lead}, and PICRUSt2 elsewhere. "
           if lead else "By region the population mean favours PICRUSt2 in every region. ")
        + "ModelSEED reaction space tracks EC space within 0.01 (conditional F1 "
        f"{rd(mo[('rxn', 'prfba_exclude', 'f1')], 3)} against "
        f"{rd(mo[('rxn', 'picrust2', 'f1')], 3)}). Both built-in asymmetries, a "
        "tenfold larger reference and a truth in prFBA's own vocabulary, favour prFBA here, so "
        "this measures the systems as shipped rather than the method; the controlled comparison "
        "removes both and the ranking reverses. Source: comparison_summary.json, "
        f"comparison_cells.csv.gz and prfba_selection.json, {DEPTH}-hit records."))
    return s


def s12_novelty(prs, D):
    import csv
    import gzip
    from pptx.enum.chart import XL_MARKER_STYLE
    s = blank(prs)
    title(s, "Precision collapses for PICRUSt2, bends for prFBA",
          sub="EC precision across PICRUSt2's own NSTI distance bins")
    cm = D["cmp"]
    rule = cm["_meta"]["prfba_selection"]["default"].split("(")[-1].rstrip(")")
    bins = ["<0.01", "0.01-0.05", "0.05-0.15", "0.15-0.5", ">=0.5"]
    st = cm["by_stratum"]["ec"]["nsti_bin"]
    # per-amplicon cells: per-bin conditional means at full precision (the summary keeps 4
    # decimals), and what each method asserts where the reference is far
    sums, far = {}, {}
    with gzip.open(os.path.join(os.path.dirname(CMP_SUMMARY), "comparison_cells.csv.gz"),
                   "rt") as fh:
        rdr = csv.reader(fh)
        ix = {c: i for i, c in enumerate(next(rdr))}
        for r in rdr:
            meth = r[ix["method"]]
            if (r[ix["space"]] != "ec" or r[ix["outcome"]] != "ok"
                    or meth not in ("picrust2", "prfba_exclude")):
                continue
            for q in ("precision", "recall"):
                sums.setdefault((r[ix["nsti_bin"]], meth, q), []).append(float(r[ix[q]]))
            if r[ix["nsti_bin"]] == ">=0.5":
                far.setdefault(r[ix["amplicon"]], {})[meth] = (
                    int(r[ix["n_pred"]]), int(r[ix["n_truth"]]), float(r[ix["precision"]]),
                    r[ix["ident_bin"]])
    get = lambda b, m, q: sum(sums[(b, m, q)]) / len(sums[(b, m, q)])                   # noqa: E731
    p2 = [get(b, "picrust2", "precision") for b in bins]
    pf = [get(b, "prfba_exclude", "precision") for b in bins]
    cd = CategoryChartData()
    cd.categories = ["< 0.01", "0.01–0.05", "0.05–0.15", "0.15–0.5", "≥ 0.5"]
    cd.add_series("prFBA", pf, number_format="0.00")
    cd.add_series("PICRUSt2", p2, number_format="0.00")
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(M), Inches(BODY_TOP),
                            Inches(7.85), Inches(5.25), cd)
    ch = gf.chart
    style_chart(ch, size=12)
    for sr, col, txt, mk, pos in zip(ch.series, (TEAL, AMBER), (TEAL, AMBER_TXT),
                                     (XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.DIAMOND),
                                     (XL_LABEL_POSITION.ABOVE, XL_LABEL_POSITION.BELOW)):
        sr.format.line.color.rgb = col
        sr.format.line.width = Pt(3.25)
        sr.smooth = False
        sr.marker.style = mk
        sr.marker.size = 9 if mk == XL_MARKER_STYLE.CIRCLE else 11
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = col
        sr.marker.format.line.color.rgb = col
        sr.has_data_labels = True
        dl = sr.data_labels
        dl.show_value = True
        dl.number_format, dl.number_format_is_linked = "0.00", False
        dl.position = pos
        dl.font.size, dl.font.bold, dl.font.name = Pt(13), True, BODY
        dl.font.color.rgb = txt
    va, ca = ch.value_axis, ch.category_axis
    va.minimum_scale, va.maximum_scale = 0.0, 1.0
    va.major_unit = 0.2
    va.has_major_gridlines = False          # every point is labelled
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = INK
    ans = [v for v in far.values() if "prfba_exclude" in v and "picrust2" in v]
    mean = lambda xs: sum(xs) / len(xs)                                                 # noqa: E731
    e_pf = mean([v["prfba_exclude"][0] for v in ans])
    e_p2 = mean([v["picrust2"][0] for v in ans])
    e_tr = mean([v["prfba_exclude"][1] for v in ans])
    prec_p2 = mean([v["picrust2"][2] for v in ans])
    near995 = sum(1 for v in ans if v["prfba_exclude"][3] == ">=0.995")
    cov = lambda b: st[b]["prfba_exclude"]["coverage"]                                  # noqa: E731
    x0, w0 = 8.85, W - M - 8.85
    stat(s, x0, BODY_TOP + 0.3, w0, f"{rd(e_pf, 0)} vs {rd(e_p2, 0)}",
         f"ECs per answer at NSTI ≥ 0.5, prFBA vs PICRUSt2; the truth has {rd(e_tr, 0)}",
         h=1.5)
    stat(s, x0, BODY_TOP + 2.2, w0, f"{rd(100 * cov('>=0.5'), 0)}%",
         "of NSTI ≥ 0.5 amplicons still answered by prFBA", h=1.5)
    takeaway(s, "prFBA bends by asserting smaller gene sets", x0, BODY_TOP + 4.1, w0, h=0.9)
    footer(s, f"Shipped {rule} · {DEPTH}-hit records · NSTI = nearest-sequenced-taxon index, "
              "PICRUSt2's own distance measure")
    fl = cm["by_stratum"]["ec"]["ident_bin"]["<0.865"]
    notes(s, (
        f"Across PICRUSt2's own NSTI distance bins, its EC precision falls from {rd(p2[0], 3)} to "
        f"{rd(p2[-1], 3)} while its recall barely moves "
        f"({rd(get(bins[0], 'picrust2', 'recall'), 3)} to "
        f"{rd(get(bins[-1], 'picrust2', 'recall'), 3)}): distant placements keep asserting "
        "gene content, and most of the extra assertions are wrong. prFBA's precision falls only "
        f"from {rd(pf[0], 3)} to {rd(pf[-1], 3)}, and not because it abstains: at NSTI ≥ 0.5 it "
        f"still answers {rd(100 * cov('>=0.5'), 1)}% of amplicons, against "
        f"{rd(100 * cov('<0.01'), 1)}% in the nearest bin. Instead it asserts far smaller gene "
        f"sets there, about {rd(e_pf, 0)} ECs per answered amplicon against PICRUSt2's "
        f"{rd(e_p2, 0)} on the same {len(ans):,} amplicons, whose truth sets average "
        f"{rd(e_tr, 0)} and on which PICRUSt2's precision is {rd(prec_p2, 3)}. High NSTI partly "
        f"reflects PICRUSt2's sparser tree: for {near995:,} of the {len(ans):,}, prFBA's database "
        "holds a different organism at 0.995 identity or more, which self-exclusion strips before "
        "selection. Below the 0.865 family floor prFBA declines to answer at all: for the "
        f"{fl['prfba_exclude']['n']:,} amplicons whose nearest different organism sits there, "
        "PICRUSt2 predicts at precision "
        f"{rd(fl['picrust2']['conditional']['precision']['value'], 3)}. "
        "That is the point of the slide: an abstention tells you not to build a model, while a "
        "model built from a confidently wrong genome is a plausible artefact that propagates "
        f"into every flux prediction. Shipped {rule}, {DEPTH}-hit records; source "
        "comparison_summary.json (by_stratum) and comparison_cells.csv.gz."))
    return s


def s13_ceiling(prs, D):
    """What the control measures: two annotations of one genome, no prediction."""
    import csv
    import gzip
    s = blank(prs)
    background(s, DEEP)
    title(s, "Two annotations of one genome agree only partly", dark=True,
          sub="BV-BRC's annotation vs PICRUSt2's own reference row, no prediction")
    ce = D["ceil"]
    sh = ce["ceiling"]["shared"]
    tint = RGBColor(0x0A, 0x4A, 0x55)
    ch = hbar(s, M, BODY_TOP, 7.75, 4.0,
              ["BV-BRC ECs found in PICRUSt2's row", "PICRUSt2 ECs found in BV-BRC's set"],
              [sh["recall"]["mean"], sh["precision"]["mean"]], colors=[MINT, SEAFOAM],
              fmt="0.0%", vmin=0, vmax=1, dark=True, gap=45)
    ch.value_axis.tick_labels.number_format = "0%"
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.major_unit = 0.25
    for i, (v, lab) in enumerate(((rd(sh["jaccard"]["mean"], 3), "Jaccard overlap of the two "
                                   "EC sets"),
                                  (f"{sh['recall']['n']:,}", "genomes held in both databases"))):
        card = stat(s, 8.8, BODY_TOP + i * 2.15, W - M - 8.8, v, lab, color=MINT,
                    tint=tint, h=1.85)
        for r in card.text_frame.paragraphs[1].runs:      # stat() labels in MUTED: too dark here
            r.font.color.rgb = WHITE
    takeaway(s, "Two pipelines read the same DNA differently", M,
             BODY_TOP + 4.4, W - 2 * M, dark=True, h=0.5)
    footer(s, f"Per-genome means over the {sh['recall']['n']:,} benchmark genomes in PICRUSt2's "
              f"reference; {ce['_meta']['ec_vocab']['shared']:,}-EC shared vocabulary",
           dark=True)
    # how often PICRUSt2's prediction beats that genome's own reference row on precision
    row = {g["genome_id"]: g["shared"]["precision"] for g in ce["per_genome"]}
    beat = tot = 0
    nsti, absent, far = [], [], []
    with gzip.open(os.path.join(os.path.dirname(CMP_SUMMARY), "comparison_cells.csv.gz"),
                   "rt") as fh:
        rdr = csv.reader(fh)
        ix = {c: i for i, c in enumerate(next(rdr))}
        for r in rdr:
            if (r[ix["space"]] != "ec" or r[ix["method"]] != "picrust2"
                    or r[ix["outcome"]] != "ok"):
                continue
            if r[ix["nsti_bin"]] == ">=0.5":
                far.append(float(r[ix["precision"]]))
            if r[ix["src_in_picrust2_ref"]] == "0":
                absent.append(float(r[ix["precision"]]))
            elif r[ix["src_in_picrust2_ref"]] == "1":
                g = r[ix["amplicon"]].split("__")[0]
                nsti.append(float(r[ix["nsti"]]))
                if g in row:
                    tot += 1
                    beat += float(r[ix["precision"]]) > row[g]
    nsti.sort()
    med_nsti = (nsti[(len(nsti) - 1) // 2] + nsti[len(nsti) // 2]) / 2
    inref = D["cmp"]["by_stratum"]["ec"]["src_in_picrust2_ref"]["True"]["picrust2"]["conditional"]
    notes(s, (
        "This is a control, not a result, and no prediction is involved. "
        f"{sh['recall']['n']:,} scoreable benchmark genomes sit in both databases, annotated "
        "independently by BV-BRC and by the PICRUSt2 project with different gene callers and EC "
        "rules. If PICRUSt2 placed an amplicon perfectly, on its own source genome's tip, it "
        "would return that genome's stored row, so the two annotations of the one genome are "
        f"compared directly. PICRUSt2's row recovers {rd(100 * sh['recall']['mean'], 1)}% of the "
        f"BV-BRC EC set at {rd(100 * sh['precision']['mean'], 1)}% precision (Jaccard "
        f"{rd(sh['jaccard']['mean'], 3)}): a quarter of PICRUSt2's enzyme calls and a tenth of "
        "BV-BRC's have no counterpart in the other, and neither is wrong, they are different "
        "annotation pipelines. This sets the reference level for what PICRUSt2 can score "
        f"against our truth, but it is not a strict bound: for {rd(100 * beat / tot, 0)}% of "
        "in-reference amplicons PICRUSt2's prediction is more precise than that genome's own "
        "row. On those amplicons, which land close to the source (median NSTI "
        f"{rd(med_nsti, 3)}), PICRUSt2's measured recall "
        f"({rd(inref['recall']['value'], 3)}) is "
        f"{rd(100 * inref['recall']['value'] / sh['recall']['mean'], 0)}% of this level and its "
        f"precision ({rd(inref['precision']['value'], 3)}) essentially at it, so the rest of the "
        "gap to prFBA under a BV-BRC truth "
        "comes mostly from the two projects annotating the same genomes differently. That does "
        "not extend to novel or distant amplicons, where PICRUSt2's precision falls well below "
        f"this level ({rd(sum(absent) / len(absent), 3)} when the source is absent from its "
        f"reference, {rd(sum(far) / len(far), 3)} at NSTI ≥ 0.5). Source: "
        "annotation_ceiling.json (shared vocabulary; independent of record depth) and "
        "comparison_cells.csv.gz."))
    return s


def s13b_dialect(prs, D):
    """What the control means: each method scores best in its own dialect."""
    from pptx.enum.chart import XL_MARKER_STYLE
    s = blank(prs)
    title(s, "Each method wins at home",
          sub="Identical predictions, two answer keys for the same genome")
    dl = D["dialect"]
    sc = dl["by_scope"]["conditional"]
    dm = dl["_meta"]["prfba_minus_picrust2_f1"]
    n = sc["picrust2"]["bvbrc"]["f1"]["n"]
    rule = D["cmp"]["_meta"]["prfba_selection"]["default"].split("(")[-1].rstrip(")")
    arms = (("PICRUSt2", "picrust2", AMBER, AMBER_TXT),
            ("prFBA, source removed", "prfba_exclude", TEAL, TEAL),
            ("prFBA, source present", "prfba_include", MUTED, MUTED))
    f1 = lambda k, key: sc[k][key]["f1"]["mean"]                                       # noqa: E731
    cd = CategoryChartData()
    cd.categories = ["BV-BRC's annotation", "PICRUSt2's annotation"]
    for name, k, _, _ in arms:
        cd.add_series(name, [f1(k, "bvbrc"), f1(k, "picrust2")], number_format="0.000")
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(M), Inches(BODY_TOP),
                            Inches(7.85), Inches(5.25), cd)
    ch = gf.chart
    style_chart(ch, size=12)
    for sr, (_, _, col, txt) in zip(ch.series, arms):
        sr.format.line.color.rgb = col
        sr.format.line.width = Pt(2.0 if col == MUTED else 3.25)
        sr.smooth = False
        sr.marker.style = (XL_MARKER_STYLE.DIAMOND if col == AMBER else XL_MARKER_STYLE.CIRCLE)
        sr.marker.size = 12 if col == AMBER else 10
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = col
        sr.marker.format.line.color.rgb = col
        for i, pos in enumerate((XL_LABEL_POSITION.LEFT, XL_LABEL_POSITION.RIGHT)):
            lab = sr.points[i].data_label
            lab.position = pos
            lab.font.size, lab.font.bold, lab.font.name = Pt(14), True, BODY
            lab.font.color.rgb = txt
    va, ca = ch.value_axis, ch.category_axis
    va.minimum_scale, va.maximum_scale = 0.7, 1.0
    va.major_unit = 0.05
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.00", False
    va.has_major_gridlines = False          # every point is labelled
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.size = Pt(13)
    ca.tick_labels.font.color.rgb = INK
    for ax, text in ((ca, "Answer key"), (va, "Conditional EC F1")):
        ax.has_title = True
        ax.axis_title.text_frame.text = text
        atf = ax.axis_title.text_frame.paragraphs[0].runs[0].font
        atf.size, atf.bold, atf.color.rgb = Pt(12), False, MUTED
    x0, w0 = 8.85, W - M - 8.85
    stat(s, x0, BODY_TOP + 0.3, w0, sg(dm["conditional|truth=bvbrc"], 3),
         "prFBA (source removed) − PICRUSt2, BV-BRC's key", h=1.5)
    stat(s, x0, BODY_TOP + 2.2, w0, sg(dm["conditional|truth=picrust2"], 3),
         "same predictions, PICRUSt2's key", h=1.5, color=AMBER_TXT)
    takeaway(s, "Cannot rank the methods", x0, BODY_TOP + 4.1, w0, h=0.9)
    footer(s, f"Conditional on prFBA answering: n = {n:,} amplicons whose source genome is in "
              f"both databases · shipped {rule}")
    home = {"picrust2": "picrust2", "prfba_exclude": "bvbrc", "prfba_include": "bvbrc"}
    away = {"picrust2": "bvbrc", "prfba_exclude": "picrust2", "prfba_include": "picrust2"}
    ratio = [f1(k, away[k]) / f1(k, home[k]) for k in home]
    notes(s, (
        "The same predictions, on the same amplicons and in the same EC vocabulary, are scored "
        "against two answer keys for the same genome: BV-BRC's annotation and PICRUSt2's own. "
        f"Under our key prFBA leads by {rd(dm['conditional|truth=bvbrc'], 3)} F1; under "
        f"PICRUSt2's it trails by {rd(-dm['conditional|truth=picrust2'], 3)}, so each method "
        "scores best in its own dialect. The diagonal is near-tautological, since each system is "
        "handed back its own annotation of a genome it already contains, and each off-diagonal "
        f"cell is {rd(min(ratio), 2)}–{rd(max(ratio), 2)} of the home score. So do not claim "
        "that prFBA predicts function more accurately than PICRUSt2: this benchmark cannot "
        "support that in either direction, which is why the controlled comparison was built, "
        "and its result does not favour us. What does survive here is coverage and abstention, a "
        "design difference rather than a measurement; the shape of precision decay with novelty, "
        "within each method; and prFBA's own validation, whose output is a BV-BRC genome, so "
        "that truth is the operative one there. A fair cross-method ranking needs a truth "
        "neither system authors, which is exactly what the shared-reference k-fold benchmark "
        f"provides. Conditional on prFBA answering: {n:,} of the {dl['_meta']['n_amplicons']:,} "
        f"amplicons whose source genome is in both databases; shipped {rule}, {DEPTH}-hit "
        "records; source annotation_dialect.json."))
    return s


def s14_depth(prs, D):
    """Depth 100 is the standard: what the 20 -> 100 hit record changed, and what it left."""
    s = blank(prs)
    title(s, "Depth 100 is now the standard",
          sub="Every record keeps 100 hits, not 20 · chart: self-exclusion, exact-tie union")
    regs = list(REG_LABEL)
    # abstention decomposition, shipped exact-tie union: (amplicons, abstained, emptied)
    (n20, a20, e20), (n100, a100, e100) = D["absd"]["d20"], D["absd"]["d100"]
    # both bars per depth as % of amplicons: the drop in abstention is the drop in emptying
    vbar(s, M, BODY_TOP, 5.3, 4.95, ["Abstaining", "Record emptied"],
         [("Depth 20", [100.0 * a20 / n20, 100.0 * e20 / n20]),
          ("Depth 100", [100.0 * a100 / n100, 100.0 * e100 / n100])],
         colors=[MUTED, TEAL], fmt='0.0"%"', vmin=0, vmax=30, gap=70)

    def per_region(tree, key):
        return [D["taxd"][tree]["regions"][r]["anchor"]["exclude_0.987"]["all"]["all"]
                ["per_rank"]["Genus"][key] for r in regs]

    def genus(tree, key):
        v = per_region(tree, key)
        return f"{rd(min(v), 2)}–{rd(max(v), 2)}"

    def genes(tree):
        v = [D["capd"][tree]["_meta"]["headline_macro_recall_pct_population"]
             ["exclude_0.987"][r] for r in regs]
        return f"{rd(min(v), 0)}–{rd(max(v), 0)}%"

    def gap(dep, key="prfba_exclude"):
        c = D["cmpd"][dep]["by_region"]["ec"]["all"]["all"]
        return sg(c[key]["population"]["f1"]["value"]
                  - c["picrust2"]["population"]["f1"]["value"])

    def pairs(tree):
        p = D["mcnd"][tree]["by"]["anchor|Genus|exclude_0.987"]["pairs"]
        return sum(1 for x in p if x["significant"]), len(p), p

    m20, m100 = pairs("d20"), pairs("L100")
    table(s, 6.3, BODY_TOP + 0.05, W - M - 6.3, [
        ["Self-exclusion metric", "Depth 20", "Depth 100"],
        ["Population genus", genus("d20", "correct_population"),
         genus("L100", "correct_population")],
        ["Population gene capture", genes("d20"), genes("L100")],
        ["Gap to PICRUSt2 (union)", gap("d20"), gap("d100")],
        ["Region pairs that differ", f"{m20[0]}/{m20[1]}", f"{m100[0]}/{m100[1]}"],
    ], col_w=[3.03, 1.7, 1.7], size=15, row_h=0.66, emph_col=2)
    takeaway(s, f"Still emptied at depth 100: {e100:,} amplicons", 6.3, BODY_TOP + 3.75,
             W - M - 6.3, h=0.5)
    footer(s, f"Original panel, {n100:,} amplicons, τ = 0.987 · table: legacy reducer, the "
              "only rule with a depth-20 headline scoring, unless marked")
    same = [f"{REG_LABEL[x['region_a']]} with {REG_LABEL[x['region_b']]}"
            for x in m100[2] if not x["significant"]]
    shift = max(abs(a - b) for a, b in zip(per_region("d20", "correct"),
                                           per_region("L100", "correct")))
    cf = [D["cmpd"][d]["by_region"]["ec"]["all"]["all"]["prfba_exclude"]["conditional"]
          ["f1"]["value"] for d in ("d20", "d100")]
    pr = D["probe"]
    notes(s,
          f"Since the depth-100 rerun every amplicon record keeps the best 100 of the 500 "
          f"shortlisted references instead of 20, and the depth-20 trees survive only as the "
          f"baseline. Under the shipped exact-tie union, self-exclusion abstention fell from "
          f"{a20:,} amplicons ({100 * a20 / n20:.1f}%) to {a100:,} ({100 * a100 / n100:.1f}%), "
          f"and the abstentions with surviving hits barely moved ({a20 - e20} against "
          f"{a100 - e100}), so the whole drop is records that exclusion no longer empties; "
          f"their share of the abstentions fell from {100 * e20 / a20:.1f}% to "
          f"{100 * e100 / a100:.1f}% (the chart gives both as shares of all amplicons). "
          f"Under the legacy reducer that tab:main uses, population "
          f"genus accuracy rose from {genus('d20', 'correct_population')} to "
          f"{genus('L100', 'correct_population')} and population gene capture from "
          f"{genes('d20')} to {genes('L100')}, while conditional genus accuracy moved by at "
          f"most {rd(shift)} in any region ({genus('d20', 'correct')} against "
          f"{genus('L100', 'correct')}) and the as-shipped conditional EC F1 only from "
          f"{cf[0]:.3f} to {cf[1]:.3f}. The as-shipped population gap to PICRUSt2 shrank from "
          f"{gap('d20')} to {gap('d100')} under the union, and under the legacy reducer it is "
          f"{gap('d100', 'prfba_exclude_legacy')}, within resampling noise (next slide). The "
          f"paired McNemar test now separates {m100[0]} of {m100[1]} region pairs at genus "
          f"under the legacy reducer instead of all {m20[1]}; the ones that no longer differ "
          f"are {', '.join(same[:-1])} and {same[-1]}. What remains: {e100:,} amplicons are "
          f"still emptied at depth 100, and for {SHORTLIST_ALL_TIED[0]} of the "
          f"{SHORTLIST_ALL_TIED[1]:,} distinct sequences all 500 shortlisted references tie at "
          f"the top score, so their tie cluster is wider than the shortlist itself. The depth "
          f"probe that motivated the rerun ({pr['_meta']['n_sampled']} emptied amplicons "
          f"re-aligned at depth 500, {100 * pr['recovery']['rate']:.1f}% recovered) had "
          f"predicted most of this. Sources: prfba_selection.json, "
          f"taxacc_region_summary.json, gene_capture_by_region.json, mcnemar_regions.json and "
          f"comparison_summary.json in region_validation and region_validation_d100[_legacy]; "
          f"region_validation_d100/D100_VS_D20.md sections 2 and 5.")
    return s


def s15_next(prs, D):
    s = blank(prs)
    background(s, DEEP)
    title(s, "Where this leaves us", dark=True)
    # C12 under amendment A-C12: only once controls.json is bound to the rerun's summary
    c12 = D["ctl"]["C12"]
    cm = c12.get("measured") or {}
    rescored = D["c12_rescored"]
    if rescored:
        c12_item = "C12 amended: " + ("passes" if c12["status"] == "pass" else "still fails")
        valid = D["kg"]["headline"].get("valid")
        c12_note = (
            f"C12 {'passes' if c12['status'] == 'pass' else 'still fails'} on the rerun "
            f"({cm['n_truth_models_gated']} gated truth rows, {cm['n_models_gated']} of "
            f"{cm['n_models']:,} rows gated), so the leave-genus-out artefact is "
            + ("valid under the amendment. " if valid else
               "still marked invalid; as pre-committed, the failure is reported and nothing "
               "is tuned. "))
    else:
        c12_item = "C12 amendment: rerun under way"
        c12_note = "the rerun is under way. "
    cols = [("Decided", AMBER, ["Exact-tie union throughout", "GTDB map: future work",
                                "ModelSEED PRs and issue #10 filed", c12_item]),
            ("Next steps", MINT, ["Push the revised manuscript", "Deployment run at depth 100",
                                  "Resolve two open faults", "Close the loop to flux"])]
    gap = 0.4
    cw = (W - 2 * M - gap) / 2
    card = RGBColor(0x0A, 0x4A, 0x55)
    for c, (head, accent, items) in enumerate(cols):
        x = M + c * (cw + gap)
        chip = rect(s, x, BODY_TOP, cw, 0.58, accent)
        tf = chip.text_frame
        tf.margin_left = Inches(0.25)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, head, size=18, bold=True, color=INK, font=HEAD, first=True, space_after=0)
        for i, it in enumerate(items):
            y = BODY_TOP + 0.85 + i * 1.15
            box = rect(s, x, y, cw, 0.85, card)
            btf = box.text_frame
            btf.margin_left = Inches(0.95)
            btf.margin_right = Inches(0.2)
            btf.margin_top = btf.margin_bottom = 0
            btf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(btf, it, size=18, color=WHITE, first=True, space_after=0)
            dot = rect(s, x + 0.25, y + 0.175, 0.5, 0.5, accent, shape=MSO_SHAPE.OVAL)
            dtf = dot.text_frame
            dtf.margin_left = dtf.margin_right = dtf.margin_top = dtf.margin_bottom = 0
            dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(dtf, str(i + 1), size=15, bold=True, color=INK, font=HEAD,
                 align=PP_ALIGN.CENTER, first=True, space_after=0)
    footer(s, "github.com/freiburgermsu/prFBA  ·  github.com/freiburgermsu/prFBA_paper",
           dark=True)
    notes(s, (
        "Four decisions are made. First, one selection rule: the exact-tie union is used "
        "throughout the paper and the deck; the legacy (ordered) reducer stays a user option "
        "and appears only in the labelled depth-20 versus depth-100 comparisons. Second, the "
        "GTDB accession map is built but deferred to future work, because wiring it in "
        "changes which genomes a GTDB hit selects. Third, the ModelSEED template pull "
        "requests and the direction-audit issue (ModelSEED/ModelSEEDTemplates#10) were filed "
        "on 2026-10-27. Fourth, the C12 fix was adopted as SPEC amendment A-C12, post hoc and "
        "disclosed: templates 21a8084 plus the gapfiller's reversals, with both phenotype "
        "passes rerun on ModelSEEDpy pinned to the pre-registered commit"
        + ("; amendment A-C13 then reran both passes on ModelSEEDpy c2e31a0, which fixes the "
           "builder's demand reactions and an ATP-correction step that could undo its own "
           "loop filter" if D["amend_label"] == "A-C13" else "") + "; " + c12_note +
        "Next: push the revised manuscript branch, re-run the real-amplicon deployment at "
        "record depth 100 (it still uses the earlier depth and is labelled so in the paper), "
        "resolve the two open faults on the audit slide, and close the loop to flux by "
        "propagating probability-weighted genomes through community prFBA and validating "
        "against paired metagenomes."))
    return s


# ------------------------------------------------ the MAG arm (synthetic MAGs)
def s16_mag_arm(prs, D):
    s = blank(prs)
    m, g = D["mag"]["summary"], D["gtdb"]["summary"]
    tiers = [1, 0.9, 0.7, 0.5, 0.3]
    ex = [m[f"c{c:g}_x0_exclude"] for c in tiers]
    gi = [g[f"c{c:g}_x0_include"] for c in tiers]
    ge = [g[f"c{c:g}_x0_exclude"] for c in tiers]
    n_mags = ex[0]["n"]
    title(s, "Completeness is a null axis for the MAG arm",
          sub=f"Synthetic MAGs matched by whole-genome ANI, 100% → {tiers[-1]:.0%} complete")
    ch = vbar(s, M, BODY_TOP + 0.05, 7.5, 4.45, [f"{c:.0%}" for c in tiers],
              [("MAG (query) aligned", [r["median_af_query"] / 100 for r in gi]),
               ("Reference aligned", [r["median_af_ref"] / 100 for r in gi])],
              fmt="0.0%", vmin=0, vmax=1.16, gap=60)
    chart_title(ch, "Median aligned fraction by MAG completeness, source present")

    def arrow_value(card):   # the arrow in the body font, which every renderer carries
        p = card.text_frame.paragraphs[0]
        r0 = p.runs[0]
        a, b = r0.text.split(" → ")
        r0.text = a
        for t, fn in (("  →  ", BODY), (b, HEAD)):
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = r0.font.size, True, fn
            r.font.color.rgb = r0.font.color.rgb

    x2 = 8.45
    w2 = W - M - x2
    f1 = [r["genecap_cond_f1"]["mean"] for r in ex]
    cov = [r["coverage_answered"] for r in ex]
    gen = [r["Genus"]["correct"] for r in ex]
    for i, (a, b, lab) in enumerate([
            (f1[0], f1[-1], "gene-capture F1, 100% → 30%"),
            (cov[0], cov[-1], "resolve rate, 100% → 30%"),
            (gen[0], gen[-1], "genus accuracy, 100% → 30%")]):
        arrow_value(stat(s, x2, BODY_TOP + 0.05 + i * 1.5, w2, f"{rd(a)} → {rd(b)}", lab,
                         h=1.2))
    takeaway(s, "Selection gates on the MAG's aligned fraction, which never moves",
             M, BODY_TOP + 4.8, W - 2 * M, h=0.45)
    footer(s, f"Bars: {D['gtdb']['n_genomes']:,}-organism GTDB subsample, self-inclusion  "
              "·  cards: panel, self-exclusion at ANI 95, 0% contamination")
    n_cells = sum(v["n"] for k, v in m.items() if "_x0_" in k)
    notes(s, (
        f"The MAG arm re-runs the same panel organisms as {n_mags:,} synthetic MAGs, matched "
        "to a skani "
        "sketch of the panel by whole-genome ANI instead of 16S identity, and scores them with "
        "the amplicon arm's own truth, metric definitions and conditional-versus-population "
        "denominators, so the two arms are directly comparable. "
        "Self-exclusion removes the source genome, its taxon and every hit at 95 ANI or above "
        "(the species boundary, the analogue of the amplicon arm's τ = 0.987); the "
        f"0%-contamination grid has {n_cells:,} scored cells, {n_mags:,} MAGs at five "
        "completeness tiers under both regimes. "
        f"From 100% to 30% completeness nothing the pipeline decides moves: conditional "
        f"gene-capture F1 {rd(f1[0])} to {rd(f1[-1])}, resolve rate {rd(cov[0])} to "
        f"{rd(cov[-1])}, genus accuracy {rd(gen[0])} to {rd(gen[-1])}. "
        "The bars show why: with the source present the MAG's own aligned fraction stays at "
        f"{rd(gi[0]['median_af_query'], 1)}% at every completeness, because a clean subset of "
        "contigs still aligns end to end, while the reference's aligned fraction falls in step "
        f"with completeness, {rd(gi[0]['median_af_ref'], 1)}% to "
        f"{rd(gi[-1]['median_af_ref'], 1)}%; under self-exclusion the MAG's fraction is flat "
        f"too ({rd(min(r['median_af_query'] for r in ge), 1)}–"
        f"{rd(max(r['median_af_query'] for r in ge), 1)}%) and the reference's still scales "
        f"({rd(ge[0]['median_af_ref'], 1)}% to {rd(ge[-1]['median_af_ref'], 1)}%). "
        "The selector gates on the MAG's aligned fraction and ignores the reference's, so the "
        "pipeline already computes, and discards, an essentially unbiased completeness "
        "estimate. "
        "The prediction is the gene content of the selected reference, not of the MAG, so "
        "completeness can matter only by changing which reference is chosen, and ANI is "
        "robust enough that it rarely does."))
    return s


def s17_crossover(prs, D):
    from pptx.enum.chart import XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
    s = blank(prs)
    rows = D["magx"]
    top, bot = rows[0], rows[-1]
    rule = "exact-tie union" if MAGX_JSON.endswith("_d100.json") else "legacy reducer"
    title(s, "No completeness at which the amplicon catches up",
          sub=f"Paired within organism, from complete genomes down to "
              f"{bot['completeness']:.0%} drafts")
    cd = CategoryChartData()
    cd.categories = [f"{r['completeness']:.0%}" for r in rows]
    cd.add_series("vs full-length 16S", [r["FullLength16S_delta"] for r in rows],
                  number_format="0.000")
    cd.add_series("vs V4", [r["V4_delta"] for r in rows], number_format="0.000")
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(M), Inches(BODY_TOP + 0.05),
                            Inches(7.55), Inches(4.85), cd)
    ch = gf.chart
    style_chart(ch, size=12)
    ch.legend.font.size = Pt(12)
    va, ca = ch.value_axis, ch.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0.0, 0.25, 0.05
    va.has_major_gridlines = False          # both lines are labelled at both ends
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.00", False
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.color.rgb = INK
    for sr, col, tcol, key in zip(ch.series, (TEAL, AMBER), (TEAL, AMBER_TXT),
                                  ("FullLength16S_delta", "V4_delta")):
        sr.format.line.color.rgb = col
        sr.format.line.width = Pt(3)
        sr.smooth = False
        sr.marker.style = XL_MARKER_STYLE.CIRCLE
        sr.marker.size = 7
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = col
        sr.marker.format.line.color.rgb = col
        for idx in (0, len(rows) - 1):
            dl = sr.points[idx].data_label
            dl.text_frame.text = sg(rows[idx][key])
            dl.position = XL_LABEL_POSITION.ABOVE
            f = dl.text_frame.paragraphs[0].runs[0].font
            f.size, f.bold, f.name = Pt(14), True, BODY
            f.color.rgb = tcol
    chart_title(ch, "MAG − 16S gene-capture F1 by completeness (uneven steps; 0 = a tie)")

    def arrow_value(card):   # the arrow in the body font, which every renderer carries
        p = card.text_frame.paragraphs[0]
        r0 = p.runs[0]
        a, b = r0.text.split(" → ")
        r0.text = a
        for t, fn in (("  →  ", BODY), (b, HEAD)):
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = r0.font.size, True, fn
            r.font.color.rgb = r0.font.color.rgb

    P = D["magp"]
    regs = list(REG_LABEL)
    cond = {r: P[f"exclude|c1_x0|{r}|f1"]["conditional"] for r in regs}
    dlt = {r: round(cond[r]["mag"] - cond[r]["amp"], 4) for r in regs}
    pop = P["exclude|c1_x0|FullLength16S|f1"]["population"]
    x2 = 8.45
    w2 = W - M - x2
    for i, (v, lab, col) in enumerate([
            (f"{sg(min(dlt.values()))} to {sg(max(dlt.values()))}",
             f"lead over all {len(regs)} regions at 100%", TEAL),
            (f"{rd(top['resolve_rate'])} → {rd(bot['resolve_rate'])}",
             f"resolve rate, 100% → {bot['completeness']:.0%}: coverage degrades",
             SEA_TXT),
            (f"{rd(pop['mag'])} vs {rd(pop['amp'])}",
             "whole-population F1, MAG vs full-length 16S", INK)]):
        c = stat(s, x2, BODY_TOP + 0.05 + i * 1.62, w2, v, lab, color=col, h=1.32)
        if "→" in v:
            arrow_value(c)
    ns = [r[k] for r in rows for k in ("FullLength16S_n", "V4_n")]
    footer(s, f"Amplicon cells: depth-{DEPTH} records, {rule}  ·  conditional on both "
              f"arms answering, n = {min(ns):,}–{max(ns):,} per point")

    v1 = [r for r in regs if r.startswith("V1") or r == "FullLength16S"]
    v3 = [r for r in regs if r not in v1]
    widest = max(dlt, key=dlt.get)
    p_max = max(r[k] for r in rows for k in ("FullLength16S_p", "V4_p"))
    notes(s, (
        f"Paired within organism against the amplicon arm's own cells (depth-{DEPTH} records, "
        f"{rule}, as in tab:main), the MAG arm leads full-length 16S by "
        f"{sg(top['FullLength16S_delta'])} in gene-capture F1 at full completeness and still by "
        f"{sg(bot['FullLength16S_delta'])} at {bot['completeness']:.0%}, fifty-fold below "
        "MIMAG's medium-quality floor and not a bin any study would keep; against V4 the lead "
        f"is {sg(top['V4_delta'])} and {sg(bot['V4_delta'])}. "
        f"At {bot['completeness']:.0%} the MAG arm still wins "
        f"{bot['FullLength16S_mag_wins']:,} organisms to {bot['FullLength16S_amp_wins']:,} "
        f"against full-length 16S, and every one of the {2 * len(rows)} paired tests has "
        f"Wilcoxon p ≤ {p_max:.0e}. "
        f"Across all {len(regs)} regions at full completeness the lead follows the variable "
        f"region, not amplicon length: {sg(min(dlt[r] for r in v1))} to "
        f"{sg(max(dlt[r] for r in v1))} against full-length 16S and the two V1-anchored "
        f"regions, {sg(min(dlt[r] for r in v3))} to {sg(max(dlt[r] for r in v3))} against the "
        f"five that start at V3 or later, with {REG_LABEL[widest]} the widest gap. "
        "The reason is structural: the prediction is the gene content of the selected "
        "reference genome, not of the MAG, so completeness matters only if it changes which "
        "reference is chosen, and ANI is robust enough that it rarely does. "
        f"What degrades is coverage, not quality: the resolve rate falls from "
        f"{rd(top['resolve_rate'])} to {rd(bot['resolve_rate'])} as the remaining sequence "
        f"stops clearing skani's screen, while conditional F1 drops by only "
        f"{rd(top['cond_f1'] - bot['cond_f1'])}. "
        "All of this is conditional on both arms answering; over the whole population, scoring "
        f"an abstention as zero, the amplicon arm wins, {rd(pop['amp'])} against "
        f"{rd(pop['mag'])} for full-length 16S, because the MAG arm answers for only about "
        f"{top['resolve_rate']:.0%} of organisms (see the abstention slide)."))
    return s


def s18_abstain(prs, D):
    s = blank(prs)
    ex = D["mag"]["summary"]["c1_x0_exclude"]
    res = ex["coverage_answered"]
    top = D["magx"][0]
    rule = "exact-tie union" if MAGX_JSON.endswith("_d100.json") else "legacy reducer"
    title(s, f"The MAG arm abstains on {100 * (1 - res):.0f}% — and should",
          sub="One gate, not the ANI floor, accounts for nearly all of it")
    labels = [t.split(" — ")[0] for t, _ in AF_SWEEP]
    labels = [t[0].upper() + t[1:] for t in labels]
    ch = hbar(s, M, BODY_TOP + 0.05, 7.45, 4.75, labels, [v for _, v in AF_SWEEP],
              colors=[TEAL, MUTED, AMBER, AMBER, AMBER, RGBColor(0x7F, 0x96, 0x9C)],
              fmt="0.0%", vmin=0, vmax=0.75, gap=50)
    chart_title(ch, "Share with a qualifying surviving hit, as each gate is relaxed")

    fl = top["FullLength16S_amp"]
    if fl not in AF2_BEAT_FL:
        raise SystemExit("s18_abstain: no AF2_BEAT_FL count for the full-length 16S level %s"
                         % fl)
    beat = 100.0 * AF2_BEAT_FL[fl] / AF2_NEW[0]
    f_ok = ex["genecap_cond_f1"]["mean"]
    x2 = 8.45
    w2 = W - M - x2
    takeaway(s, "A looser gate buys worse answers", x2, BODY_TOP + 0.1, w2, h=0.5)
    stat(s, x2, BODY_TOP + 0.8, w2, f"{rd(AF2_NEW[1])} vs {rd(f_ok)}",
         "gene-capture F1, newly admitted vs kept", color=AMBER_TXT, h=1.6)
    stat(s, x2, BODY_TOP + 2.7, w2, f"{beat:.1f}%",
         f"of newly admitted beat full-length 16S ({rd(fl)})", h=1.6)
    footer(s, f"Self-exclusion at ANI 95, 100% complete  ·  16S level: depth-{DEPTH} "
              f"records, {rule}")

    n_ok = ex["n_ok"]
    pooled = (n_ok * f_ok + AF2_NEW[0] * AF2_NEW[1]) / (n_ok + AF2_NEW[0])
    band = D["cmp"]["by_stratum"]["ec"]["ident_bin"]["0.945-0.987"]
    amp_band = band["prfba_exclude" if rule == "exact-tie union" else "prfba_exclude_legacy"]
    sweep = dict(AF_SWEEP)
    keys = [t for t, _ in AF_SWEEP]
    gt = D["gtdb"]
    notes(s, (
        f"At full completeness the MAG arm's resolve rate is {rd(res)}, and the cause of the "
        "abstentions is marker reach rather than MAG quality: among the "
        f"{MAG_REACH[0]:,} organisms that do have a same-genus congener in the panel, the "
        f"best surviving hit aligns over a median of only {MAG_REACH[1]}% of the query, and "
        f"{MAG_REACH[2]}% fall under the 15% AF gate. "
        "An amplicon aligns end to end by construction, so in the 0.945–0.987 16S-identity "
        f"band the amplicon arm answers {amp_band['coverage']:.1%} of the time (depth-{DEPTH} "
        f"records, {rule}) where whole-genome ANI answers {BAND_MAG[0] * 100:.0f}–"
        f"{BAND_MAG[1] * 100:.0f}%. "
        "Counterfactually the ANI floor is almost irrelevant and the AF gate is almost "
        f"everything: dropping the floor from 83 to 0 moves the share with a qualifying "
        f"surviving hit from {rd(sweep[keys[0]])} to {rd(sweep[keys[1]])}, while relaxing the "
        f"gate from 15% to 2% moves it to {rd(sweep[keys[3]])}; the resolve rate after "
        f"selection is {rd(res)}. "
        f"The gate is nonetheless the method working: the {AF2_NEW[0]:,} organisms a 2% gate "
        f"would newly answer carry a mean gene-capture F1 of {rd(AF2_NEW[1])} against "
        f"{rd(f_ok)} for those already answered, and only {beat:.1f}% of them exceed the "
        f"full-length 16S level of {rd(fl)}, so relaxing it would buy "
        f"{AF2_NEW[0] / n_ok:.0%} more coverage while pulling the pooled F1 down to "
        f"{rd(pooled)}. "
        f"Reference density also matters: against GTDB r226, about {GTDB_FOLD}-fold denser, "
        f"the same pipeline resolves {rd(gt['summary']['c1_x0_exclude']['resolve_rate'])} of a "
        f"{gt['n_genomes']:,}-organism subsample, with completeness-invariance intact. "
        "Retention depth, the dominant cause of the amplicon arm's self-exclusion "
        f"abstentions, plays almost no part: only {MAG_CAPPED_EMPTIED_X0[0]} of the "
        f"{MAG_CAPPED_EMPTIED_X0[1]:,} uncontaminated records "
        f"({100 * MAG_CAPPED_EMPTIED_X0[0] / MAG_CAPPED_EMPTIED_X0[1]:.2f}%; "
        f"{MAG_CAPPED_EMPTIED[0]} of {MAG_CAPPED_EMPTIED[1]:,} counting the contaminated grid) "
        f"were both capped at 100 hits and emptied by exclusion, and re-searching {MAG_DEPTH_PROBE[0]} capped "
        f"queries and {MAG_DEPTH_PROBE[1]} controls at depth 2,000 changed {MAG_DEPTH_PROBE[2]} "
        "selections, every one from abstention to an answer."))
    return s


def s19_two_panel(prs, D):
    s = blank(prs)
    title(s, "The coverage gap is emptied records",
          sub=f"Full-length 16S · self-exclusion at τ = 0.987 · a {DEPTH}-hit record")
    P = ("original", "kmeans", "random")
    sfx = D["tp3"]["meta"].get("suffix") or ""          # which rule scored the three trees
    rule = ("legacy reducer" if sfx.endswith("_legacy") else
            "exact-tie union" if sfx else "mixed rules")
    cov = tp3_cell(D, "FullLength16S", "exclude_0.987", "genecap_coverage")
    unc = {p: 100 * (1 - cov[p]["crude"]) for p in P}
    E = D["tp3_emptied"]
    surv = ({p: 100 * E["panels"][p]["survivors"]["non_ok_rate"] for p in P} if E else
            {p: 100 * (unc[p] - EMPTIED[p]) / (100 - EMPTIED[p]) for p in P})
    ch = vbar(s, M, BODY_TOP + 0.05, 7.75, 5.0, ["Taxonomy-stratified", "k-means",
                                                  "Uniform random"],
              [("Uncovered (1 − coverage)", [unc[p] for p in P]),
               ("Record emptied by exclusion", [EMPTIED[p] for p in P])],
              fmt='0.0"%"', vmin=0, vmax=40, colors=[MUTED, DEEP])
    ch.value_axis.tick_labels.number_format = '0"%"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    x, w = M + 7.75 + 0.3, W - M - (M + 7.75 + 0.3)
    stat(s, x, BODY_TOP + 0.3, w, f"{min(surv.values()):.1f}–{max(surv.values()):.1f}%",
         "uncovered where the record survives, on every panel", h=1.85)
    stat(s, x, BODY_TOP + 0.3 + 2.15, w,
         f"{SELF_TOP20['original']:.1f}% vs {SELF_TOP20['random']:.1f}%",
         f"source genome among its {DEPTH} hits, stratified vs random", h=1.85)
    footer(s, f"Record depth {DEPTH} · {rule} · coverage = gene capture, crude")
    # uncovered = abstained unless a selected reference carries no gene families
    roots = (D["tp3"]["meta"].get("roots") or {}).values()
    nogf = sum(json.load(open(os.path.join(REPO, os.path.basename(r), "data",
                                           "gene_capture_by_region.json")))
               ["_meta"]["outcome_counts_total"]["exclude_0.987"][k]
               for r in roots for k in ("pred_empty", "pgfam_missing"))
    gf = (" Uncovered is one minus crude gene-capture coverage, which on these trees is the "
          "abstentions alone: " + ("no cell" if nogf == 0 else f"only {nogf:,} cells")
          + " on any panel selects a reference without gene families.") if roots else ""
    clo = "; ".join(f"1 − {rd(EMPTIED[p] / 100)} = {rd(1 - EMPTIED[p] / 100)} against "
                    f"{rd(cov[p]['crude'])}" for p in ("random", "kmeans", "original"))
    cmh = (" The surviving-record rates run opposite to the coverage gap, the best-covered "
           "stratified panel leaving the most survivors uncovered (Cochran-Mantel-Haenszel "
           f"p = {rd(E['tests']['cmh_survivor_strata']['original_vs_random']['p'])} "
           "stratified against random, "
           f"{rd(E['tests']['cmh_survivor_strata']['original_vs_kmeans']['p'], 2)} against "
           "k-means).") if E else ""
    notes(s, (
        "On full-length 16S under self-exclusion at τ = 0.987, with all three panels scored "
        f"at record depth {DEPTH} under the {rule}, self-exclusion deletes all {DEPTH} "
        f"retained hits for {EMPTIED['original']}% of stratified-panel queries, "
        f"{EMPTIED['kmeans']}% of k-means and {EMPTIED['random']}% of random. Where the "
        f"record survives, only {surv['original']:.1f}%, {surv['kmeans']:.1f}% and "
        f"{surv['random']:.1f}% go uncovered; where it is emptied, every query goes uncovered "
        f"by construction, with no evidence either way about rank {DEPTH + 1}.{gf} So the "
        "cross-panel coverage difference is, to within a point and a half, the difference in "
        f"how often the record empties: {clo}.{cmh} Self-inclusion is panel-dependent for a "
        "related "
        f"reason: the source genome is among its amplicon's {DEPTH} retained hits for "
        f"{SELF_TOP20['original']}% of amplicons on the stratified panel but "
        f"{SELF_TOP20['kmeans']}% and {SELF_TOP20['random']}% on the k-means and random "
        "panels. The misses are exact ties filling the whole record "
        f"({SELF_SAT['original']}%, {SELF_SAT['kmeans']}% and {SELF_SAT['random']}% of "
        "amplicons) or the source's sequence credited to another genome by the index, which "
        f"keeps one genome per distinct sequence ({SELF_CRED['original']}%, "
        f"{SELF_CRED['kmeans']}% and {SELF_CRED['random']}%). A miss still finds an "
        "exact-identity relative, but a different genome, so self-inclusion on a "
        "database-representative panel mixes true self-recovery with recovery of a "
        "16S-identical neighbour."))
    return s


def s20_panel_space(prs, D):
    s = blank(prs)
    title(s, "Stratified over-samples sparse space; random tracks it",
          sub="Local density in the full 1024-d embedding, not in the projection")
    # three_panel/plot_space_three.py --slide: drawn at this size, so its type reads at
    # 11-13 pt (composition only, so independent of record depth)
    img = os.path.join(REPO, "three_panel", "F_panel_space_three_slide.png")
    s.shapes.add_picture(img, Inches(M), Inches(1.55), width=Inches(W - 2 * M))
    d = D["dec3"]
    sp = lambda k: 100 * sum(d[k][:2])        # noqa: E731  two sparsest deciles
    takeaway(s, f"Two sparsest deciles: {sp('og'):.1f}% of stratified picks, "
                f"{sp('rnd'):.1f}% of random, {sp('population'):.1f}% of sequences",
             M, 5.95, W - 2 * M)
    footer(s, f"Shares of on-frame picks: {d['n_og']:,} of the stratified panel's 10,000; "
              f"all {d['n_km']:,} k-means and {d['n_rnd']:,} random picks")
    pop = [100 * v for v in d["population"]]
    notes(s, (
        "Each panel's picks are placed on density deciles of the "
        f"{D['knn']['n']:,} clusterable full-length 16S sequences. Density is the cosine "
        "distance to the 21st nearest neighbour in the full 1024-dimensional embedding, not "
        "in the UMAP projection, which keeps neighbourhoods but not densities; the ranking is "
        f"stable in k (Spearman {D['knn']['spearman_21_50']:.3f} between k = 21 and k = 50). "
        f"The stratified panel over-weights the sparse space hardest: {sp('og'):.1f}% of its "
        f"on-frame picks come from the two sparsest deciles, which hold "
        f"{sp('population'):.1f}% of sequences, against {sp('km'):.1f}% for k-means. The "
        f"random panel (C) tracks the population: {sp('rnd'):.1f}% from the two sparsest "
        f"deciles, and {100 * d['rnd'][9]:.1f}% of its picks in the densest decile against "
        f"the population's {100 * d['population'][9]:.1f}% (stratified "
        f"{100 * d['og'][9]:.1f}%, k-means {100 * d['km'][9]:.1f}%). The dense deciles are "
        "where a handful of species carry thousands of near-identical operons, which is "
        "where self-exclusion empties a record, as the next slides show. The black ticks in "
        "D are each decile's share of all sequences: the density measure is heavily tied, so "
        f"the deciles hold {min(pop):.1f}–{max(pop):.1f}% of sequences rather than 10% each. "
        f"Only {d['n_og']:,} of the stratified panel's organisms are on frame (the rest have "
        "short or atypical 16S, or a sequence the index credits to no admissible taxon). The "
        "manuscript's version of this figure (F16) still shows only the stratified and k-means "
        "panels."))
    return s


# ----------------------------------------- the controlled, shared-reference arm
def s15_kfold_design(prs, D):
    """The controlled design as one flow: panel, blocked folds, one reference, one key."""
    s = blank(prs)
    title(s, "A controlled test: one reference, one answer key",
          sub="Neither system authors the truth it is scored against")
    g, f, kp = D["kg"], D["kf"], D["kplan"]
    mg = g["meta"]
    cg = g["contrasts"]["prfba_nestedcv__vs__picrust2"]["ec"]
    cf = f["contrasts"]["prfba_nestedcv__vs__picrust2"]["ec"]
    ref = g["headline"]["reference_density_genomes"]
    sesoi = mg["frozen"]["sesoi_delta_f1"]
    n_base = sum(1 for k in g["arms"] if k.startswith("base_"))
    flow(s, M, BODY_TOP + 0.4, W - 2 * M, 1.6, [
        (f"{mg['n_eligible']:,} genomes", "dereplicated panel"),
        (f"{mg['n_folds']} folds", "genera held out"),
        (f"{ref:,} references", "median, shared by both"),
        ("2 methods", f"+ {n_base} baselines"),
        ("1 answer key", "BV-BRC ECs")], size=17, gap=0.36)
    for i, (v, lab, col) in enumerate([
            (f"{mg['n_cells_scored']:,}", "scored cells", TEAL),
            (f"{cg['bootstrap_n_groups']:,}", "holdout groups", TEAL),
            (f"{sesoi}", "pre-registered effect size", AMBER_TXT)]):
        card = stat(s, M + i * 4.144, BODY_TOP + 2.6, 3.845, v, lab, color=col, h=2.0,
                    vsize=40)
        for p in card.text_frame.paragraphs:          # centred, like the flow above
            p.alignment = PP_ALIGN.CENTER
    footer(s, f"Leave-genus-out, the primary rank; leave-family-out adds "
              f"{f['meta']['n_cells_scored']:,} cells in {cf['bootstrap_n_groups']} groups")
    winners = {v["inner_winner"] for blk in (g, f) for v in blk["nested_cv"].values()
               if isinstance(v, dict)}
    notes(s, (
        "This is the experiment the dialect slide said was needed: both methods are confined, "
        f"fold by fold, to one species-dereplicated reference ({ref:,} genomes, median per "
        "fold; PICRUSt2's alignment, tree and trait table rebuilt from it with its own tooling) "
        "and scored in one EC vocabulary, BV-BRC's, against the same truth. "
        "Holdout blocks are the connected components of a leakage graph joining genomes that "
        "share a genus, an identical amplicon or a 16S identity of at least "
        f"{mg['frozen']['tau_amplicon']}, packed into {kp['k']} folds; the planned "
        f"{mg['frozen']['k_folds']} were impossible because the largest component holds "
        f"{100 * kp['largest_component'] / kp['n_eligible']:.1f}% of the "
        f"{kp['n_eligible']:,} eligible genomes. "
        f"Leave-genus-out is primary ({mg['n_cells_scored']:,} scored cells, cluster bootstrap "
        f"over {cg['bootstrap_n_groups']:,} holdout groups, effective n {cg['n_eff']:,}); "
        f"leave-family-out is secondary ({f['meta']['n_cells_scored']:,} cells, "
        f"{cf['bootstrap_n_groups']} groups, effective n {cf['n_eff']:,}). "
        f"The smallest effect of interest ({sesoi} population F1), the selection mode, the "
        "controls and the exit criteria were fixed before the run; leave-one-fold-out "
        "selection chose the "
        + (f"{winners.pop()} reducer in every fold at both ranks"
           if len(winners) == 1 else "mode per fold") +
        ", so 'prFBA' on the next slide is that rule. "
        "The baselines are a constant (the most prevalent ECs), nn1 (copy the top-ranked "
        "surviving reference), a family donor (a random same-family genome), random abstain "
        "(prFBA's forced answers declined at random to its coverage) and clade-max (the best "
        "genome of the nearest surviving clade, chosen against the truth: an oracle bound). "
        f"Panel caveat: an HMM check finds {D['hmm']['n_not_16s']} of "
        f"{D['hmm']['n_records']:,} panel records are not 16S and {D['hmm']['n_low_cover']} "
        "cover less than 80% of the model. Sources: region_validation/shared_reference/data/"
        "summary_genus.json, data_family/summary_family.json and data/fold_plan.json."))
    return s


def s16_kfold_result(prs, D):
    """The controlled result: one chart, the gap, and what the design's rule says of it."""
    s = blank(prs)
    title(s, "Under a neutral answer key, the ranking reverses",
          sub="Population EC F1 · leave-genus-out · one shared reference")
    g, f = D["kg"], D["kf"]
    GREY = RGBColor(0xA7, 0xB8, 0xBC)
    arms = [("base_clade_max", "Clade-max oracle", GREY),
            ("picrust2", "PICRUSt2", AMBER),
            ("prfba_nestedcv_forced", "prFBA, forced to answer", TEAL),
            ("base_nn1", "Copy nearest reference", GREY),
            ("prfba_nestedcv", "prFBA, pre-registered (legacy)", TEAL),
            ("prfba_exact_tie", "prFBA, exact-tie (shipped)", TEAL),
            ("base_random_abstain", "Random abstain", GREY),
            ("base_constant", "Constant floor", GREY)]
    pop = lambda blk, k: blk["arms"][k]["ec"]["population"]["f1"]["value"]   # noqa: E731
    arms.sort(key=lambda a: -pop(g, a[0]))
    ch = hbar(s, M - 0.1, BODY_TOP - 0.05, 8.15, 5.2, [lab for _, lab, _ in arms],
              [pop(g, k) for k, _, _ in arms], colors=[c for _, _, c in arms], fmt="0.000",
              vmin=0.0, vmax=0.9, gap=45)
    ch.value_axis.tick_labels.number_format = "0.0"
    ch.value_axis.tick_labels.number_format_is_linked = False

    cg = g["contrasts"]["prfba_nestedcv__vs__picrust2"]["ec"]
    cf = f["contrasts"]["prfba_nestedcv__vs__picrust2"]["ec"]
    fg = g["contrasts"]["prfba_nestedcv_forced__vs__picrust2"]["ec"]
    cnn = g["contrasts"]["prfba_nestedcv__vs__base_nn1"]["ec"]
    p2nn = g["contrasts"]["picrust2__vs__base_nn1"]["ec"]
    rx = 8.85
    rw = W - M - rx
    stat(s, rx, BODY_TOP, rw, sg(cg["delta_f1_population"]),
         f"prFBA − PICRUSt2 · 95% CI {sg(cg['ci95'][0])} to {sg(cg['ci95'][1])}",
         color=AMBER_TXT, h=1.4)
    stat(s, rx, BODY_TOP + 1.58, rw, sg(fg["delta_f1_population"]),
         "forced to answer: still behind", color=TEAL, h=1.4)

    # validity: the include-self ceiling control and the phenotype gate
    hd = g["headline"]
    c12_failed = "C12" in (hd.get("hard_controls_failed") or [])
    c2h = D["c2"]["headline"]
    c2_pass = (D["ctl"]["C2"]["status"] == "pass" and c2h["valid"]
               and "C2" in c2h["controls_evaluated_this_run"])
    vy = BODY_TOP + 3.16
    card = rect(s, rx, vy, rw, 1.98, LIGHT)
    for i, (ok, head, sub) in enumerate([
            (c2_pass, "C2 include-self ceiling",
             f"source recovery {D['ctl']['C2']['measured']['prfba_source_recovery']:.2f}"),
            (not c12_failed, "C12 gapfill validity",
             "run marked invalid · EC unchanged" if c12_failed else
             "amended · EC unchanged" if D["c12_rescored"] else None)]):
        py = vy + 0.2 + i * 0.86
        pill = rect(s, rx + 0.2, py + 0.04, 0.78, 0.34, MINT if ok else AMBER, adj=0.3)
        ptf = pill.text_frame
        ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
        ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(ptf, "PASS" if ok else "FAIL", size=11.5, bold=True, color=INK,
             align=PP_ALIGN.CENTER, first=True)
        tf = textbox(s, rx + 1.15, py, rw - 1.3, 0.8)
        para(tf, head, size=14, bold=True, color=INK, font=HEAD, first=True, space_after=1)
        if sub:
            para(tf, sub, size=12, color=MUTED, space_after=0)
    rule = {v["inner_winner"] for v in g["nested_cv"].values() if isinstance(v, dict)}
    footer(s, f"prFBA = {'/'.join(sorted(rule))} rule, chosen in every fold · cluster "
              f"bootstrap, {cg['bootstrap_n_groups']:,} groups")

    ff = f["contrasts"]["prfba_nestedcv_forced__vs__picrust2"]["ec"]
    c2m = D["ctl"]["C2"]["measured"]
    notes(s, (
        f"Every arm is scored on the same {g['meta']['n_cells_scored']:,} cells against the "
        "same answer key, with abstentions counted as zero. "
        f"prFBA reaches population EC F1 {rd(pop(g, 'prfba_nestedcv'))} against PICRUSt2's "
        f"{rd(pop(g, 'picrust2'))} ({sg(cg['delta_f1_population'])}, 95% CI "
        f"{sg(cg['ci95'][0])} to {sg(cg['ci95'][1])}); forced to answer it reaches "
        f"{rd(pop(g, 'prfba_nestedcv_forced'))} and still trails by "
        f"{rd(abs(fg['delta_f1_population']))}, so about two thirds of the deficit is "
        "abstention and the rest is in the answers. "
        "The bar the design registered is nn1, copying the single top-ranked surviving "
        f"reference: forced prFBA only ties it ({rd(pop(g, 'prfba_nestedcv_forced'))} against "
        f"{rd(pop(g, 'base_nn1'))}, a margin with no paired test), the abstaining arm falls "
        f"{rd(abs(cnn['delta_f1_population']))} below it, and PICRUSt2 clears it by "
        f"{rd(p2nn['delta_f1_population'])}. "
        f"At leave-family-out the gap widens to {sg(cf['delta_f1_population'])} (forced "
        f"{sg(ff['delta_f1_population'])}) and prFBA's skill score falls to "
        f"{sg(f['arms']['prfba_nestedcv']['ec']['skill_score'], 2)}, a point estimate whose "
        "interval against the prevalence floor spans zero; clade-max is an oracle chosen "
        "against the truth, a ceiling and not a method to beat. "
        f"These numbers do not depend on the as-shipped record depth: each fold ran its own "
        f"alignment at k2 = {g['meta']['frozen']['prfba']['k2']}, and 'prFBA' is the "
        f"{'/'.join(sorted(rule))} rule that leave-one-fold-out selection picked in every fold "
        f"(the shipped exact-tie rule scores {rd(pop(g, 'prfba_exact_tie'))}). "
        f"Validity: the include-self ceiling control C2 passes (prFBA source recovery "
        f"{c2m['prfba_source_recovery']:.2f}, PICRUSt2 {rd(c2m['picrust2_mean_ec_f1'])} "
        f"against a soft {c2m['picrust2_soft_threshold']} flag; data_c2/summary_incself.json "
        "and data/controls.json), " + (
            "but the gapfill validity gate C12 fails, so the design's rule marks this "
            "leave-genus-out artefact invalid; no EC number changes, and the next three slides "
            "show why it fails and how the fix would clear it. " if c12_failed else
            f"and the gapfill validity gate C12 passes under {D['amend_text']} (post hoc and "
            "disclosed: corrected templates, the gapfiller's reversals"
            + (" and an updated ModelSEEDpy" if D["amend_label"] == "A-C13" else "") + "), so the "
            "leave-genus-out artefact is valid; rescoring changed no EC number, and the next "
            "three slides show why C12 failed in the pre-registered run and how the fix "
            "cleared it. ") +
        "Sources: shared_reference/data/summary_genus.json and data_family/summary_family.json."))
    return s


def s17_phenotype(prs, D):
    """The phenotype readout, dark: media called wrong, the tests, and the two controls."""
    s = blank(prs)
    background(s, DEEP)
    title(s, "Where prFBA answers, ATP barely registers the gap", dark=True,
          sub=f"Media called wrong per genome · EC F1 gap here only {sg(PHENO_EC_GAP, 3)}")
    pa, pg = D["phatp"], D["phgf"]
    A = {k: v["controlled"] for k, v in pa["arms"].items()}
    c = pa["contrasts"]["prfba_exact_tie__vs__picrust2"]["controlled"]
    cnn = pa["contrasts"]["prfba_exact_tie__vs__base_nn1"]["controlled"]
    mcn_all = c["mcnemar_per_medium"]
    mcn = {m: v for m, v in mcn_all.items() if v.get("q_bh") is not None and v["q_bh"] < 0.05}
    mcn_pic = sum(1 for v in mcn.values() if v["c"] > v["b"])      # c: PICRUSt2 right, prFBA wrong
    n_tested = sum(1 for v in mcn_all.values() if v.get("p") is not None)
    CARD = RGBColor(0x0A, 0x4A, 0x55)
    GREY = RGBColor(0x6B, 0x87, 0x8D)
    rows = [("PICRUSt2", "picrust2", AMBER), ("prFBA", "prfba_exact_tie", MINT),
            ("Copy nearest reference", "base_nn1", GREY), ("Constant floor", "base_constant", GREY)]
    ch = hbar(s, M - 0.1, BODY_TOP - 0.05, 6.9, 3.85, [r[0] for r in rows],
              [A[k]["atp_hamming"]["mean"] for _, k, _ in rows], colors=[r[2] for r in rows],
              fmt="0.00", vmin=0.0, vmax=8.0, dark=True, gap=55)
    ch.value_axis.tick_labels.number_format = "0"
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.major_unit = 2

    gx = 7.55
    more = lambda v: "more" if v > 0 else "fewer"                          # noqa: E731
    for i, (value, label, col) in enumerate([
            (sg(c["delta_atp_hamming"], 2), f"{more(c['delta_atp_hamming'])} media wrong than "
             f"PICRUSt2 · p = {c['wilcoxon_p']:.2f}", AMBER),
            (sg(cnn["delta_atp_hamming"], 2), f"{more(cnn['delta_atp_hamming'])} than the "
             f"nearest copy · q = {cnn['q_bh']:.3f}", MINT)]):
        card = stat(s, gx, BODY_TOP + 0.1 + i * 1.9, W - M - gx, value, label, color=col,
                    h=1.6, tint=CARD)
        card.text_frame.paragraphs[1].runs[0].font.color.rgb = WHITE

    # the two phenotype controls, then the pointer to the diagnosis
    gated = len(D["phgf_gated"])
    truth_rows = pg["meta"]["n_gated_out"].get("truth", 0)
    ident = (pg["meta"].get("identity_control") or {}).get("pass")
    by = BODY_TOP + 3.98
    cws = [4.0, W - 2 * M - 0.3 - 4.0]                    # the C12 card carries most text
    for i, (status, head, sub) in enumerate([
            ("PASS" if ident else "FAIL", "C11 identity", "reproducible builds"),
            ("FAIL" if gated else "PASS", "C12 gapfill validity",
             f"{truth_rows} truth rows in {gated} genomes, zero allowed → cause and fix next"
             if gated else
             f"0 of {pg['meta']['n_models']:,} rows gated · "
             f"{sum(D['phgf_pre']['meta']['n_gated_out'].values())} before the amendments")]):
        x, cw = M + sum(cws[:i]) + 0.3 * i, cws[i]
        rect(s, x, by, cw, 1.08, CARD)
        tx = x + 0.2
        if status:
            pill = rect(s, x + 0.2, by + 0.33, 0.78, 0.38, MINT if status == "PASS" else AMBER,
                        adj=0.3)
            ptf = pill.text_frame
            ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
            ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(ptf, status, size=11.5, bold=True, color=INK, align=PP_ALIGN.CENTER, first=True)
            tx = x + 1.15
        tf = textbox(s, tx, by + 0.17, x + cw - tx - 0.15, 0.8, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, head, size=15, bold=True, color=WHITE if status else MINT, font=HEAD,
             first=True, space_after=1)
        if sub:
            para(tf, sub, size=12.5, color=WHITE, space_after=0)
    n_atp = pa["meta"]["n_candidates_eligible"]
    footer(s, f"{c['n_pairs']} ATP-informative of {n_atp:,} genomes · prFBA = shipped "
              "exact-tie rule · Wilcoxon signed-rank; BH-adjusted McNemar", dark=True)

    med = lambda k, m: A[k][m]["median"]                                  # noqa: E731
    c12 = D["ctl"]["C12"]["measured"]
    notes(s, (
        "Each arm's predicted EC set goes through one fixed ModelSEED reconstruction, and the "
        f"model is asked on which of {len(mcn_all)} media it can make ATP; only {n_tested} ever "
        "carry signal, and "
        "the score is the number of media called wrong against the model built from the truth "
        "genome's own annotation (which, rebuilt, scores "
        f"{A['truth']['atp_hamming']['mean']:.2f}). "
        f"Over the {c['n_pairs']} ATP-informative genomes (prFBA answers all "
        f"{n_atp:,} assayed), where its EC-F1 gap to PICRUSt2 is "
        f"only {sg(PHENO_EC_GAP, 3)}, prFBA calls {abs(c['delta_atp_hamming']):.2f} "
        f"{'more' if c['delta_atp_hamming'] > 0 else 'fewer'} media wrong than PICRUSt2 "
        f"(Wilcoxon p = {c['wilcoxon_p']:.2f}, BH q = {c['q_bh']:.2f}): no detectable mean "
        "difference, a null and not an equivalence. "
        f"Per medium, PICRUSt2 is ahead on {mcn_pic} of {n_tested} testable media "
        f"({', '.join(sorted(mcn))}; McNemar q = {rd(max(v['q_bh'] for v in mcn.values()))}), "
        f"and against copying the nearest reference prFBA calls "
        f"{abs(cnn['delta_atp_hamming']):.2f} fewer media wrong (q = {cnn['q_bh']:.3f}), the "
        "baseline it ties on ECs. "
        "PICRUSt2's drafts share more of the truth's reactions (median reaction Jaccard "
        f"{med('picrust2', 'rxn_jaccard_draft'):.3f} against "
        f"{med('prfba_exact_tie', 'rxn_jaccard_draft'):.3f}); on reaction Jaccard, "
        "essentiality and growth prFBA is indistinguishable from the nearest-reference "
        "baseline, and those readouts are unpaired medians with no paired test specified. "
        f"The identity control C11 {'passes' if ident else 'FAILS'}: one EC set built twice "
        "gives identical models. "
        + ((f"The validity gate C12 still fails its zero-tolerance leg: {gated} of "
            f"{pg['meta']['n_candidates_eligible']} gapfilled truth models ({truth_rows} rows, "
            "one build shared by both pipelines) make ATP from water and protons, while its "
            f"fraction leg passes ({sum(pg['meta']['n_gated_out'].values())} of "
            f"{pg['meta']['n_models']:,} rows, {100 * c12['gated_fraction']:.2f}% against a "
            f"{100 * c12['max_gated_fraction']:.0f}% ceiling); those "
            "genomes are dropped from every arm's gapfilled comparison, and the next two "
            "slides show the cause and a validated fix. ") if gated else
           (f"The validity gate C12 passes under {D['amend_text']} (templates 21a8084 with the "
            "gapfiller's reversals applied"
            + (", on ModelSEEDpy c2e31a0" if D["amend_label"] == "A-C13" else "")
            + ", post hoc and disclosed): no gapfilled model of any "
            f"arm makes ATP or biomass from water and protons (0 of {pg['meta']['n_models']:,} "
            f"rows, against {sum(D['phgf_pre']['meta']['n_gated_out'].values())} in the "
            f"pre-registered run, {D['phgf_pre']['meta']['n_gated_out'].get('truth', 0)} of "
            "them truth rows), so every genome stays in every arm's gapfilled comparison; the "
            "ATP pass on the corrected templates reproduced the pre-registered pass row for "
            "row, so the ATP numbers here are unchanged, and the next two slides show the "
            "pre-registered failure's cause and the fix. ")) +
        f"The phenotype arm uses the shipped exact-tie rule on {n_atp:,} "
        f"pre-specified genomes (ATP) and a {pg['meta']['n_candidates_eligible']}-genome "
        "gapfilled subsample; sources: shared_reference/data/phenotype_atp"
        + (f"{D['amend_tag']}/ and phenotype_gf{D['amend_tag']}/")
        + "phenotype_summary.json, data/controls.json."))
    return s


# ------------------------------------------------ why the MAG arm exists at all
def s19_mag_why(prs, D):
    from decimal import Decimal, ROUND_HALF_UP

    def pc(x, d=1):   # a stored fraction as a percent, half away from zero
        v = (Decimal(repr(x)) * 100).quantize(Decimal(1).scaleb(-d), rounding=ROUND_HALF_UP)
        return f"{v:.{d}f}%"

    s = blank(prs)
    ssu, fd = D["ssu"], D["feed"]
    ret = {r["completeness"]: r for r in D["ret"]["clean"]}
    iso, mag = ssu["by_kind"]["isolate"], ssu["by_kind"]["MAG"]
    mag_hi = ssu["by_kind_completeness"]["MAG"]["95-100"]
    title(s, "Most real MAGs have no 16S to give",
          sub="Share of genomes carrying a full-length 16S (≥ 1,200 bp)")
    hbar(s, M, BODY_TOP + 0.05, 7.45, 4.75,
         ["Isolate genomes", "Real MAGs", "Real MAGs, 95–100% complete",
          "Our synthetic MAGs, complete", "Our synthetic MAGs, 30% complete"],
         [iso["has_ssu_ge1200"], mag["has_ssu_ge1200"], mag_hi["has_ssu_ge1200"],
          ret[1.0]["full"], ret[0.3]["full"]],
         colors=[TEAL, AMBER, AMBER, SEAFOAM, SEAFOAM], fmt="0%", vmin=0, vmax=1.0, gap=55)

    x2 = 8.45
    w2 = W - M - x2
    uc = fd["union_over_contigs"]
    takeaway(s, "Fed to PICRUSt2, each contig becomes an organism", x2, BODY_TOP + 0.1, w2,
             h=0.85)
    stat(s, x2, BODY_TOP + 1.15, w2, f"{uc['n_ec']:,} vs {fd['n_truth_ec']}",
         "EC numbers asserted vs true, one MAG", color=AMBER_TXT, h=1.6)
    stat(s, x2, BODY_TOP + 3.0, w2, rd(uc["precision"]),
         f"precision; its one 16S alone scores {rd(fd['fulllength16s_amplicon']['precision'])}",
         h=1.6)
    footer(s, f"GTDB r232: {mag['n']:,} MAGs, {iso['n']:,} isolates  ·  synthetic: "
              f"{ret[1.0]['n']:,} panel genomes")
    nc = fd["n_contigs"]
    notes(s, (
        f"Across GTDB r232 ({ssu['n_genomes']:,} genomes; a metadata census, while every "
        f"search in this deck uses the r226 sketch), {pc(mag['has_ssu_ge1200'])} of the "
        f"{mag['n']:,} MAGs carry a full-length 16S (at least 1,200 bp) against "
        f"{pc(iso['has_ssu_ge1200'])} of the {iso['n']:,} isolate genomes, and only "
        f"{pc(mag['has_ssu'])} of MAGs carry any SSU fragment at all, so a 16S method cannot "
        "be run on most of the genomes people actually assemble. "
        f"Even the best real MAGs, 95–100% complete, carry one in only "
        f"{pc(mag_hi['has_ssu_ge1200'])} of cases. "
        f"Our synthetic MAGs are kinder: full-length 16S survives in {pc(ret[1.0]['full'])} of "
        f"complete drafts, about {ret[1.0]['full'] / mag_hi['has_ssu_ge1200']:.0f} times the "
        "rate in real MAGs at 95–100% completeness, and in "
        f"{pc(ret[0.3]['full'])} at 30% completeness, so the synthetic arm understates the "
        "problem. "
        f"Nor are contigs a substitute: fed one bacterial genome's {nc} contigs, PICRUSt2's "
        f"default run refuses all {nc} for aligning poorly to its reference, and forced "
        f"through it treats each contig as its own organism — {fd['distinct_closest_refs']} "
        f"different closest references, {fd['domains']['arc']} of {nc} contigs called "
        f"archaeal, and a union asserting {uc['n_ec']:,} EC numbers against "
        f"{fd['n_truth_ec']} real ones (precision {rd(uc['precision'])}). "
        f"{fd['no16s_contigs']['nsti_lt_0_15']} of the {fd['no16s_contigs']['n']} contigs that "
        "carry no 16S are still reported at NSTI below 0.15, which reads as a confident "
        "placement. "
        "The one full-length 16S that genome does carry scores precision "
        f"{rd(fd['fulllength16s_amplicon']['precision'])}: the marker works when it is there; "
        "the problem is that usually it is not."))
    return s


def s20_mag_build(prs, D):
    s = blank(prs)
    clean = D["ret"]["clean"]
    lo = clean[-1]["completeness"]
    title(s, "Fragment first, so every genome degrades alike",
          sub="Bases never mutated: a draft loses sequence, not identity")
    steps = [("Fragment", f"~{FRAG_MEAN_KB} kb, {FRAG_MIN_KB} kb floor"),
             ("Drop", f"whole pieces, 100% → {lo:.0%}"),
             ("Label", "MAG id on every header")]
    gap, fw = 0.42, W - 2 * M
    flow(s, M, BODY_TOP + 0.15, fw, 1.3, steps, size=18, gap=gap)
    bw = (fw - gap * (len(steps) - 1)) / len(steps)

    # a storyboard under each step: one deposited contig, cut, thinned, labelled
    lens = [5, 2, 4, 3, 6, 2, 3, 5, 4, 3]          # schematic piece lengths
    keep = [1, 0, 1, 1, 0, 1, 0, 1, 1, 0]          # which pieces survive the drop
    pad, pg, ph = 0.3, 0.06, 0.3
    sw = bw - 2 * pad
    unit = (sw - pg * (len(lens) - 1)) / sum(lens)

    def strip(x0, y, mode):
        x = x0
        for ln, k in zip(lens, keep):
            pw = ln * unit
            if mode == "cut" or k:
                rect(s, x, y, pw, ph, TEAL, shape=MSO_SHAPE.RECTANGLE)
                if mode == "label":         # a narrow flag on each surviving piece
                    rect(s, x, y - 0.13, 0.05, 0.13, AMBER, shape=MSO_SHAPE.RECTANGLE)
            elif mode == "drop":
                rect(s, x, y, pw, ph, WHITE, shape=MSO_SHAPE.RECTANGLE, line=MUTED)
            x += pw + pg

    ys = BODY_TOP + 1.85
    xs = [M + i * (bw + gap) + pad for i in range(3)]
    rect(s, xs[0], ys, sw, ph, MUTED, shape=MSO_SHAPE.RECTANGLE)    # the deposited contig
    strip(xs[0], ys + 0.55, "cut")
    strip(xs[1], ys + 0.275, "drop")
    strip(xs[2], ys + 0.275, "label")

    def arrow_value(card):   # the arrow in the body font, which every renderer carries
        p = card.text_frame.paragraphs[0]
        r0 = p.runs[0]
        a, b = r0.text.split(" → ")
        r0.text = a
        for t, fn in (("  →  ", BODY), (b, HEAD)):
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = r0.font.size, True, fn
            r.font.color.rgb = r0.font.color.rgb

    n_clean = sum(r["n"] for r in clean)
    cards = [(f"{PANEL_CLOSED_PCT}%", "panel genomes: one closed contig", TEAL),
             (f"{n_clean:,}", f"drafts, {len(clean)} completeness levels", TEAL),
             (f"{SKANI_COLLIDE[0]} → {SKANI_COLLIDE[1]}",
              "same-header queries → skani results", AMBER_TXT)]
    for i, (v, lab, col) in enumerate(cards):
        c = stat(s, M + i * (bw + gap), BODY_TOP + 3.3, bw, v, lab, color=col, h=1.5)
        if "→" in v:
            arrow_value(c)
    footer(s, "Random generator seeded by a CRC32 of the genome id · mag_validation/synth_mags.py")
    cr = D["ret"]["c100_reasons"]
    notes(s, (
        "Each panel genome is degraded in two steps rather than one. "
        f"First every deposited contig is cut into synthetic contigs of about {FRAG_MEAN_KB} kb "
        f"mean length with a {FRAG_MIN_KB} kb floor (MetaBAT2's default minimum contig "
        "length); this is required rather than cosmetic, because "
        f"{PANEL_CLOSED_PCT}% of panel genomes are closed and {PANEL_LT10_PCT}% have fewer "
        "than ten contigs, so dropping deposited contigs could not reach a fine completeness "
        "grid and would make degradation severity depend on each reference's assembly "
        "quality, across a panel whose assemblies run from one contig to thousands. "
        f"Then whole fragments are dropped at random until the target completeness is met, "
        f"giving {n_clean:,} uncontaminated drafts at {len(clean)} levels from 100% down to "
        f"{lo:.0%}; bases are never mutated, because skani compares k-mer sketches and what "
        "a MAG suffers is incompleteness, not divergence. "
        "Every contig header is prefixed with the MAG id because skani keys a query on its "
        f"first header and silently collapses collisions: {SKANI_COLLIDE[0]} variants "
        f"sharing one header came back as {SKANI_COLLIDE[1]} results, with no warning. "
        "The synthesizer can also splice in fragments from a different-phylum donor, but it "
        "overshot its contamination dose at low completeness, so every contamination claim "
        "was withdrawn and all results use the 0%-contamination grid. "
        f"The {FRAG_MIN_KB} kb floor has a cost: at full completeness "
        f"{cr['16s_only_on_contigs_dropped_by_min_frag']:,} MAGs lose their only 16S to it, "
        f"and {cr['genome_has_no_16s']} source genomes never had one. "
        "The random generator is seeded from a CRC32 of the genome id "
        "(mag_validation/synth_mags.py), so a given draft is byte-identical across runs and "
        "machines."))
    return s



# ------------------------------------------- the synthetic-genome step, re-examined
def s06b_weighting(prs, D):
    s = blank(prs)
    title(s, "The weight is active on a minority of queries",
          sub="What the weighting can, and cannot, claim")

    titled = chart_title

    cw = (W - 2 * M - 0.4) / 2
    titled(hbar(s, M, BODY_TOP + 0.05, cw, 3.9, ["Controlled benchmark", "Deployment run"],
                [BENCH_N1 / 100, DEPLOY_N1 / 100], colors=[TEAL, MUTED], fmt="0.0%", vmin=0,
                vmax=1, gap=130),
           "One genome selected: p = 1 for all")
    titled(hbar(s, M + cw + 0.4, BODY_TOP + 0.05, cw, 3.9,
                ["Benchmark ties", "Deploy, 2 genomes", "Deploy, 3+ genomes"],
                [XPHYLUM_BENCH / 100, XPHYLUM_DEPLOY[0] / 100, XPHYLUM_DEPLOY[1] / 100],
                colors=[TEAL, MUTED, MUTED], fmt="0.0%", vmin=0, vmax=1, gap=55),
           "Multi-genome sets that span phyla")
    tf = takeaway(s, "No experiment yet runs the weighted objective", M, 5.75, W - 2 * M)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    footer(s, f"Benchmark: shipped exact-tie rule · deployment: {DEPLOY_ASVS:,} real ASVs on "
              "20-hit records, before the 100-hit standard")
    notes(s, "In practice the selector usually returns a single reference, so p = 1 for every "
             "function and prFBA reduces to a two-level parsimony objective: that holds for "
             f"{BENCH_N1}% of answered amplicons on the controlled benchmark under the shipped "
             f"exact-tie rule, and for {DEPLOY_N1}% of answered ASVs in a {DEPLOY_ASVS:,}-ASV "
             "deployment run that was made on 20-hit records, before the 100-hit standard. "
             "Where the weight is active its commonest fractional value is 0.5, and the "
             "selection size tracks database redundancy rather than query confidence: it is "
             "largest where one reference would suffice and falls to one where the query is "
             "novel, so the hedge is allocated inversely to need. Where several genomes are "
             f"merged they are often not a pangenome: {XPHYLUM_BENCH}% of the benchmark's "
             f"multi-genome tie sets span phyla, against {XPHYLUM_DEPLOY[0]}% of two-genome and "
             f"{XPHYLUM_DEPLOY[1]}% of larger deployment merges. Every validation arm in this "
             "deck scores the unweighted union of the selected references' content; no "
             "experiment builds the weighted model or solves the prFBA objective, so no "
             "accuracy number here tests the weighting. The next slide tests p itself against "
             "the truth genomes. On prior art: attaching a likelihood to each reaction on the "
             "same ModelSEED/KBase stack is likelihood-based gapfilling (Benedict 2014), the "
             "objective is GIMME's with prevalence in place of expression (Becker and Palsson "
             "2008), and pangenome profiles from a 16S call (PanFP) and nearest-genome copying "
             "(Piphillin) precede it; what is new is the combination of exact-tie selection at "
             "database depth with abstention.")
    return s


# ------------------------------------------------------------ the random panel
def s21_random_panel(prs, D):
    s = blank(prs)
    title(s, "Three panels, from flattened to the database's mix",
          sub="Same size and pipeline; sampled on three different principles")
    P = ("original", "kmeans", "random")
    C = {k: composition(v) for k, v in D["sel3"].items()}
    ids = {k: {e["genome_id"] for e in v} for k, v in D["sel3"].items()}
    rows = [["", "Taxonomy-stratified", "k-means (16S)", "Uniform random"],
            ["Archaea"] + [f"{C[p]['arc_pct']:.1f}%" for p in P],
            ["Phyla"] + [f"{C[p]['phyla']}" for p in P],
            ["Named genera"] + [f"{C[p]['genera']:,}" for p in P],
            ["Genus Gini"] + [f"{C[p]['gini']:.3f}" for p in P]]
    tbl = table(s, M, BODY_TOP + 0.15, 8.3, rows, col_w=[1.85, 2.35, 2.05, 2.05], size=17,
                row_h=0.8)
    for j, p in enumerate(P, 1):                 # each panel's colour, as in its figure
        tbl.cell(0, j).fill.fore_color.rgb = PANEL_COLORS[p]
    pool = D["random_meta"]["n_admissible_genomes"]
    shared = len(ids["original"] & ids["kmeans"] & ids["random"])
    x = M + 8.3 + 0.35
    stat(s, x, BODY_TOP + 0.15, W - M - x, f"{pool:,}",
         "admissible organisms: the pool behind k-means and random", h=1.85)
    stat(s, x, BODY_TOP + 0.15 + 2.15, W - M - x, f"{shared}",
         "organisms in all three panels: largely disjoint samples", h=1.85)
    takeaway(s, "Random is the null: no correction, the pool's own skew", M, 6.0,
             W - 2 * M)
    footer(s, f"Each panel 10,000 organisms · random panel drawn uniformly, seed "
              f"{D['random_meta']['seed']}")
    ov = lambda a, b: len(ids[a] & ids[b])   # noqa: E731
    notes(s, (
        "We asked whether the benchmark panel decides the answer, so the 16S validation and "
        "the core of the MAG validation were repeated on two further panels of the same size. "
        "The taxonomy-stratified original square-root-flattens availability, up-weights rare "
        "phyla and archaea and keeps one organism per species; the k-means panel takes the "
        "organism nearest each of 10,000 k-means++ centroids over the "
        f"{D['random_meta']['n_admissible_sequences']:,} full-length 16S embeddings; the random "
        f"panel draws 10,000 organisms uniformly (seed {D['random_meta']['seed']}) from the same "
        f"{pool:,} admissible organisms, with the original panel's representative-16S rule. "
        "The random panel is the null: it applies no correction, so Archaea fall to "
        f"{C['random']['arc_pct']:.1f}% against {100 * POOL_ARCHAEA / pool:.1f}% of the pool, "
        f"its largest genus contributes {C['random']['largest']} organisms "
        f"({C['original']['largest']} on the stratified panel), and it holds "
        f"{C['random']['species']:,} species against {C['kmeans']['species']:,} for k-means "
        f"and {C['original']['species']:,} for the stratified panel. The three schemes order "
        "identically on every diversity axis, the stratified panel most even and the random "
        f"panel most concentrated. They overlap little: {ov('original', 'kmeans')} organisms "
        f"stratified and k-means, {ov('original', 'random')} stratified and random, "
        f"{ov('kmeans', 'random')} k-means and random, {shared} in all three. Before any "
        "comparison, each comparison's crude column was required to reproduce the panel's own "
        f"scorer exactly: {3 * len(D['tp3']['cells'])} of {3 * len(D['tp3']['cells'])} "
        f"amplicon cells and {3 * len(D['tp3g']['cells'])} of {3 * len(D['tp3g']['cells'])} "
        "GTDB cells."))
    return s


def s22_three_panel(prs, D):
    s = blank(prs)
    title(s, "Coverage falls on the random panel; quality holds",
          sub="Phylum-standardized, 113 shared phyla · self-exclusion at τ = 0.987")
    P = ("original", "kmeans", "random")
    labs = ("Taxonomy-stratified", "k-means", "Uniform random")
    regs = list(REG_LABEL)
    sfx = D["tp3"]["meta"].get("suffix") or ""          # which rule scored the three trees
    rule = ("legacy reducer" if sfx.endswith("_legacy") else
            "exact-tie union" if sfx else "mixed rules")
    cell = lambda r, ep: tp3_cell(D, r, "exclude_0.987", ep)            # noqa: E731
    eps = [("genus_coverage", "Genus coverage"),
           ("genecap_coverage", "Gene-capture coverage"),
           ("genus_cond_acc", "Genus accuracy (answered)"),
           ("genecap_cond_f1", "Gene-capture F1 (answered)")]
    vbar(s, M, BODY_TOP + 0.05, 7.75, 5.0, [lab for _, lab in eps],
         [(lab, [cell("V4", ep)[p]["phystd"] for ep, _ in eps]) for p, lab in zip(P, labs)],
         fmt="0.00", vmin=0, vmax=1.0, colors=[PANEL_COLORS[p] for p in P])
    est = lambda r, ep, k: cell(r, ep)[k]["est"]                          # noqa: E731
    rmo, kmo, rmk = ("delta_random_minus_original", "delta_kmeans_minus_original",
                     "delta_random_minus_kmeans")
    covc = [cell(r, ep)[rmo] for r in regs for ep in ("genus_coverage", "genecap_coverage")]
    cov = [c["est"] for c in covc]
    gen = [est(r, "genus_cond_acc", k) for r in regs for k in (rmo, kmo)]
    fam = max(abs(est(r, "family_cond_acc", k)) for r in regs for k in (rmo, kmo, rmk))
    f1 = max(abs(est(r, "genecap_cond_f1", k)) for r in regs for k in (rmo, kmo, rmk))
    every = "every region" if all(c["excludes_zero"] for c in covc) else "by region"
    gtxt = ((f"up to {sg(max(gen), 2)}", "conditional genus accuracy, alternative panels higher")
            if min(gen) > 0 else (f"≤ {rd(max(map(abs, gen)), 2)}",
                                  "conditional genus accuracy, either direction"))
    x, w, h, gap = M + 7.75 + 0.3, W - M - (M + 7.75 + 0.3), 1.45, 0.3
    for i, (v, lab, col) in enumerate([
            (f"{sg(max(cov), 2)} to {sg(min(cov), 2)}",
             f"coverage, random − stratified, {every}", AMBER_TXT),
            gtxt + (TEAL,),
            (f"≤ {rd(f1, 3)}", "conditional gene-capture F1, any two panels", TEAL)]):
        stat(s, x, BODY_TOP + 0.05 + i * (h + gap), w, v, lab, color=col, h=h)
    footer(s, f"Chart: V4 · record depth {DEPTH} · {rule} · 1,000 phylum-stratified "
              "bootstrap replicates")
    v4g = cell("V4", "genus_coverage")
    keep = [cell(r, ep)[rmo]["est"] / (cell(r, ep)["random"]["crude"]
                                       - cell(r, ep)["original"]["crude"])
            for r in regs for ep in ("genus_coverage", "genecap_coverage")]
    gci = sum(cell(r, "genus_cond_acc")[rmo]["excludes_zero"] for r in regs)
    f1r = [-est(r, "genecap_cond_f1", rmo) for r in regs]
    notes(s, (
        "Every rate here is phylum-standardized over the 113 phyla present in all three "
        "panels, restricted to organisms with a full-length best 16S, under self-exclusion "
        f"at τ = 0.987, with all three panels scored at record depth {DEPTH} under the {rule}; "
        "intervals are 1,000 phylum-stratified bootstrap replicates. The chart shows "
        f"V4: genus coverage is {rd(v4g['original']['phystd'])}, "
        f"{rd(v4g['kmeans']['phystd'])} and {rd(v4g['random']['phystd'])} on the stratified, "
        f"k-means and random panels, a random-minus-stratified difference of "
        f"{sg(v4g[rmo]['est'])} ({sg(v4g[rmo]['ci'][0])} to {sg(v4g[rmo]['ci'][1])}) against "
        f"{sg(v4g['random']['crude'] - v4g['original']['crude'])} crude. In every region, "
        f"genus and gene-capture coverage fall by {rd(-max(cov), 2)} to {rd(-min(cov), 2)} on "
        "the random panel, every interval excludes zero, and the standardized gap keeps "
        f"{100 * min(keep):.0f}–{100 * max(keep):.0f}% of the crude one, so composition is a "
        "minority of it. Quality given an answer moves far less. Conditional genus accuracy "
        "is higher on both alternative panels than on the stratified one in every region, by "
        f"up to {rd(max(gen), 2)} (random minus stratified excludes zero in {gci} of "
        f"{len(regs)} regions). Conditional family accuracy differs by at most {rd(fam, 2)} "
        "between any two panels, and conditional gene-capture F1 by at most "
        f"{rd(f1, 3)}: it is lower on the random panel by {rd(min(f1r), 3)} to "
        f"{rd(max(f1r), 3)}, an order of magnitude below the coverage gap; the MAG arm, "
        "scored against each panel as its own reference, stays within the same bound."))
    return s


def s23_reversal(prs, D):
    import csv
    s = blank(prs)
    title(s, "Against a dense reference the sign reverses",
          sub="Same panels: 16S coverage falls, genome resolve rate rises")
    P = ("original", "kmeans", "random")
    labs = ("Taxonomy-stratified", "k-means", "Uniform random")
    sfx = D["tp3"]["meta"].get("suffix") or ""          # which rule scored the three trees
    rule = ("legacy reducer" if sfx.endswith("_legacy") else
            "exact-tie union" if sfx else "mixed rules")
    v4 = tp3_cell(D, "V4", "exclude_0.987", "genus_coverage")
    gt = next(c for c in D["tp3g"]["cells"]
              if c["regime"] == "exclude" and c["completeness"] == 1.0)
    vbar(s, M, BODY_TOP + 0.05, 7.75, 4.5,
         ["16S V4 · genus coverage", "MAG vs GTDB · resolve rate"],
         [(lab, [v4[p]["phystd"], gt[p]["phystd"]]) for p, lab in zip(P, labs)],
         fmt="0.00", vmin=0, vmax=0.9, colors=[PANEL_COLORS[p] for p in P])
    takeaway(s, "The coverage gap tracks the record, not the organisms", M, 6.18, 7.75)

    def emptied(path):
        """% of MAGs whose GTDB record is capped (100 hits) under inclusion and empty
        under species exclusion -- the genome-scale counterpart of an emptied record."""
        rows = list(csv.DictReader(open(path)))
        inc = {r["mag"]: int(r["n_hits"]) for r in rows if r["regime"] == "include"}
        exc = [r for r in rows if r["regime"] == "exclude"]
        return 100 * sum(inc.get(r["mag"], 0) >= 100 and int(r["n_hits"]) == 0
                         for r in exc) / len(exc)
    ge = {p: emptied(os.path.join(REPO, *f)) for p, f in (
        ("original", ("mag_validation", "out_gtdb", "gtdb_cells_2000.csv")),
        ("kmeans", ("mag_validation_km", "out_gtdb", "gtdb_cells.csv")),
        ("random", ("mag_validation_rnd", "out_gtdb", "gtdb_cells.csv")))}
    dg = gt["delta_random_minus_original"]
    x, w = M + 7.75 + 0.3, W - M - (M + 7.75 + 0.3)
    stat(s, x, BODY_TOP + 0.3, w, sg(dg["est"], 2),
         f"random − stratified against GTDB, 95% CI {sg(dg['ci'][0], 2)} to "
         f"{sg(dg['ci'][1], 2)}", color=SEA_TXT, h=1.85)
    stat(s, x, BODY_TOP + 0.3 + 2.15, w,
         "< 1%" if max(ge.values()) < 1 else f"{max(ge.values()):.1f}%",
         f"GTDB records emptied by exclusion; 16S: {EMPTIED['original']}–"
         f"{EMPTIED['random']}%", color=TEAL, h=1.85)
    footer(s, "Phylum-standardized · GTDB r226, exclusion at ANI ≥ 95 · 16S at record "
              f"depth {DEPTH}, {rule}")
    n = D["tp3g"]["meta"]["n_genomes"]
    dex = [c["delta_random_minus_original"]["est"] for c in D["tp3g"]["cells"]
           if c["regime"] == "exclude"]
    mag = max(abs(c[k]["est"]) for c in D["tp3p"]["cells"] if c["regime"] == "exclude"
              for k in c if k.startswith("delta_cond_f1"))
    amp = max(abs(c[k]["est"]) for c in D["tp3"]["cells"]
              if c["regime"] == "exclude_0.987" and c["endpoint"] == "genecap_cond_f1"
              for k in c if k.startswith("delta_"))
    notes(s, (
        "The GTDB arm settles whether the amplicon coverage gap reflects the organisms or the "
        "record. It searches synthetic MAGs of an approximately 20% phylum-proportional "
        f"subsample of each panel ({n['original']:,}, {n['kmeans']:,} and {n['random']:,} "
        f"scored organisms) against the {GTDB_REPS:,} species-representative genomes of GTDB "
        "r226, excluding the source species at ANI 95, with a record 100 hits deep"
        + (", as deep as the amplicon record" if DEPTH == 100 else "")
        + ". Species exclusion empties a depth-capped GTDB record for "
        f"under 1% of MAGs ({ge['original']:.1f}%, {ge['kmeans']:.1f}% and "
        f"{ge['random']:.1f}%), against {EMPTIED['original']}%, {EMPTIED['kmeans']}% and "
        f"{EMPTIED['random']}% of full-length 16S queries. With identity measured over the "
        "genome, the alternative panels resolve more novel organisms, not fewer: "
        f"{rd(gt['random']['phystd'])} and {rd(gt['kmeans']['phystd'])} against "
        f"{rd(gt['original']['phystd'])} at full completeness, standardized over the "
        f"{D['tp3g']['meta']['n_common_phyla']} shared phyla, a random-minus-stratified "
        f"difference of {sg(dg['est'])} ({sg(dg['ci'][0])} to {sg(dg['ci'][1])}) that stays "
        f"between {sg(min(dex))} and {sg(max(dex))} at every completeness down to 30%. The "
        f"median best surviving relative sits at {GTDB_BEST_ANI['random']} and "
        f"{GTDB_BEST_ANI['kmeans']} ANI for the random and k-means panels against "
        f"{GTDB_BEST_ANI['original']} for the stratified one: "
        "organisms drawn in proportion to the database come from well-sampled clades, while "
        "the stratified panel reaches into sparse ones. What this licenses: the headline "
        "self-exclusion coverage is a property of the stratified panel at record depth "
        f"{DEPTH}, a lower bound whose tightness depends on the panel, and the claim that "
        "replicates on all three panels is the conditional one: standardized conditional "
        f"gene-capture F1 differs by at most {rd(max(amp, mag), 3)} between any two panels in "
        "either arm."))
    return s


# --------------------------------------------------------------- the audit
def s24_audit(prs, D):
    s = blank(prs)
    title(s, "How the work was checked",
          sub="Since the last deck · each claim re-derived from the primary files")
    ok, n = D100_REPORT_CHECK
    cw, gap = (W - 2 * M - 2 * 0.3) / 3, 0.3
    for i, (v, lab, col) in enumerate([
            (f"{ok} / {n}", "depth-100 report claims confirmed", TEAL),
            (" · ".join(map(str, MS_D100_ROUNDS)),
             f"manuscript findings in {len(MS_D100_ROUNDS)} rounds, all fixed", AMBER_TXT),
            (f"{TEMPLATE_PR_ROUNDS} rounds", "ModelSEED PR reviews; one caught a NadX block",
             TEAL)]):
        stat(s, M + i * (cw + gap), BODY_TOP + 0.05, cw, v, lab, color=col, h=1.5)
    tax = sum(r["abstention"]["exclude_0.987"]["all"]["all"]["n_abstained"]
              for r in D["tax"]["regions"].values())
    sel = D["sel_legacy_abst"]
    rows = [["Fault caught", "Status"],
            ["ModelSEEDpy fork drops gapfill reversals", "fixed · PR ready"],
            ["ModelSEEDpy builds free demand sources", "rerun on the fix (A-C13)"],
            [f"Legacy abstentions: {tax:,} vs {sel:,}", "open"],
            ["Source-taxon mismatch, k-means/random", "open"]]
    tbl = table(s, M, BODY_TOP + 1.85, W - 2 * M, rows, col_w=[8.6, W - 2 * M - 8.6],
                size=16, row_h=0.62, emph_col=1)
    for i, row in enumerate(rows[1:], 1):
        if row[1] == "open":
            for r in tbl.cell(i, 1).text_frame.paragraphs[0].runs:
                r.font.color.rgb = AMBER_TXT
    footer(s, "Report check: region_validation_d100/D100_VS_D20.md section 5 · PR texts: "
              "region_validation/shared_reference/")
    mm = SRC_TAXON_MISMATCH
    notes(s, (
        "Since the last deck, the depth-100 rerun and everything built on it went through "
        "adversarial checks that re-derive each claim from the primary files. The report "
        "comparing the depth-100 rerun with the earlier results "
        "(region_validation_d100/D100_VS_D20.md) was "
        f"checked claim by claim by independent verifiers: {ok} of {n} claims confirmed, "
        "with the corrections listed in its section 5. The manuscript's depth-100 revision "
        f"went through three verification rounds, with {MS_D100_ROUNDS[0]}, "
        f"{MS_D100_ROUNDS[1]} and {MS_D100_ROUNDS[2]} findings, every one applied; a fourth "
        "re-derived every number the switch to the exact-tie union changed or added, about "
        f"580, found none wrong, and found {MS_D100_ROUNDS[3]} text and figure defects, all "
        "fixed (region_validation_d100/UNION_VERIFICATION.md). The "
        f"ModelSEED template PR texts went through {TEMPLATE_PR_ROUNDS} rounds of review; one "
        "caught that the first fix set NadX to the database's reverse direction, which "
        "blocks the first step of NAD synthesis in organisms without NadB, so the fix became "
        "five reactions with L-aspartate oxidase (rxn08167) restricted forward. Faults "
        "caught on the way: the ModelSEEDpy fork's MSBuilder.build_and_gapfill drops the "
        "gapfiller's reversals; amendment A-C12 applies them in prFBA's own gapfill step, "
        "and the fork is fixed on a branch whose own build_and_gapfill now reproduces the "
        "harness's models, while the unfixed builder returned all four test models unable to "
        "grow (PR text: region_validation/shared_reference/PR_ModelSEEDpy_gapfill_reversals.md). "
        "The editable ModelSEEDpy install moved under the first A-C12 launch, so both passes "
        "were rerun pinned to the pre-registered commit, which reproduces every stored row; "
        "the commits it had moved to fix two real build defects (demand reactions built as "
        "free sources of six drain compounds, invisible to the validity gate, and an "
        "ATP-correction step that could undo its own loop filter), so amendment A-C13 reran "
        "both passes on them; four manuscript "
        "figures were still the earlier renders under the revised captions and were "
        "re-rendered. "
        "Two stay open: the legacy reducer's two scoring passes disagree on self-exclusion "
        f"abstentions ({tax:,} in score_taxacc's estimator path against {sel:,} in the "
        f"gene-capture selection), and in {mm[0]}–{mm[1]} full-length rows per k-means and "
        "random tree the cells' source taxon differs from the truth file's, an effect of at "
        f"most {mm[2]} points on the emptied share. Earlier audits checked "
        f"{sum(AUDIT_CLAIMS)} manuscript claims ({sum(AUDIT_DEFECTS)} defects corrected) and "
        f"{DECK_AUDIT[0]} deck claims plus {FIX_AUDIT[0]} over the fixes ({DECK_AUDIT[1]} and "
        f"{FIX_AUDIT[1]} defects), all fixed, and our own runs caught "
        f"{len(PIPELINE_CAUGHT)} pipeline faults before that."))
    return s


def s06c_calibration(prs, D):
    import csv
    import re
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.oxml.ns import qn
    if not D["calib_ok"]:
        raise SystemExit("calibration/assertions.json: a required check failed; not reporting")
    s = blank(prs)
    title(s, "p(function) ranks well but is overconfident",
          sub="Is p a calibrated probability that the function is present?")
    # observed presence at each exact p: multi-genome selections, and n = 1 for contrast
    pts = {}
    with open(os.path.join(SR, "data", "calibration", "reliability.csv")) as fh:
        for r in csv.DictReader(fh):
            if (r["analysis"], r["run"], r["space"], r["stratum_kind"], r["denominator"],
                    r["bin_scheme"]) != ("primary_controlled", "exact_tie|d100", "ec", "all",
                                         "union", "exact_or_decile"):
                continue          # "d100" here is reference density 1.0, not record depth
            m = re.fullmatch(r"n=(\d+),p=(\d+)/(\d+)", r["bin"])
            if m and r["pop"] == ("single" if m.group(1) == "1" else "multi"):
                pts.setdefault(int(m.group(1)), []).append(
                    (float(r["mean_p"]), float(r["obs_freq"])))
    for v in pts.values():
        v.sort()
    E = D["calib"]["ec"]
    mu, n2 = E["multi|all=all"], E["n2|all=all"]
    slope = mu["union"]["calib_slope_wls"]
    skill = n2["union"]["bss_vs_const"]
    dF1 = E["all|all=all"]["binary"]["p_ge_0.5"]["d_macro_F1_vs_union"]
    sesoi = D["calib"]["meta"]["sesoi_delta_f1"]
    share2 = n2["n_cells"] / mu["n_cells"]
    # the slide's claims, checked against the data they summarize
    for ok, claim in (
            (all(b[1] > a[1] for nn in (2, 3, 4) for a, b in zip(pts[nn], pts[nn][1:])),
             "presence rises with p within n = 2, 3, 4"),
            (slope["hi"] < 1, "the calibration slope's CI lies below 1 (overconfident)"),
            (skill["lo"] < 0 < skill["hi"], "the n = 2 skill CI spans 0"),
            (abs(dF1["est"]) < sesoi, "thresholding moves macro F1 by less than the SESOI")):
        if not ok:
            raise SystemExit("s06c_calibration: no longer true: " + claim)
    num = lambda x, d: sg(x, d) if x < 0 else rd(x, d)   # noqa: E731
    # the reliability diagram
    cd = XyChartData()
    srs = ([("calibrated", [(0.0, 0.0), (1.0, 1.0)])]
           + [(f"{nn} genomes", pts[nn]) for nn in (2, 3, 4)] + [("one genome", pts[1])])
    longest = max(len(xy) for _, xy in srs)
    for name, xy in srs:
        se = cd.add_series(name)
        # repeat the last point: LibreOffice drops XY series shorter than the longest one
        for xv, yv in list(xy) + [xy[-1]] * (longest - len(xy)):
            se.add_data_point(xv, yv)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(M), Inches(BODY_TOP),
                            Inches(7.35), Inches(5.25), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=12)
    ch.legend.font.size = Pt(12)
    for ax, lab in ((ch.category_axis, "p(function) = k / n"),
                    (ch.value_axis, "observed presence")):
        ax.minimum_scale, ax.maximum_scale, ax.major_unit = 0, 1, 0.25
        ax.has_major_gridlines = True
        ax.major_gridlines.format.line.color.rgb = RGBColor(0xDC, 0xE6, 0xE6)
        ax.major_gridlines.format.line.width = Pt(0.75)
        ax.format.line.color.rgb = RGBColor(0xC8, 0xD4, 0xD6)
        ax.tick_labels.font.size = Pt(12)
        ax.tick_labels.font.color.rgb = MUTED
        ax.tick_labels.number_format, ax.tick_labels.number_format_is_linked = "0.00", False
        ax.has_title = True
        atf = ax.axis_title.text_frame
        atf.text = lab
        for r in atf.paragraphs[0].runs:
            r.font.size, r.font.bold, r.font.name = Pt(13), False, BODY
            r.font.color.rgb = INK
    dot, diamond = XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.DIAMOND
    styles = [(MUTED, 1.5, None, True), (SEAFOAM, 2.5, dot, False), (TEAL, 2.5, dot, False),
              (DEEP, 2.5, dot, False), (AMBER, 0, diamond, False)]
    for se, (col, wid, mk, dash) in zip(ch.series, styles):
        se.smooth = False
        if wid:
            se.format.line.color.rgb = col
            se.format.line.width = Pt(wid)
            if dash:
                se.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            se.format.line.fill.background()
        if mk is None:
            se.marker.style = XL_MARKER_STYLE.NONE
        else:
            se.marker.style = mk
            se.marker.size = 10 if mk == dot else 13
            se.marker.format.fill.solid()
            se.marker.format.fill.fore_color.rgb = col
            se.marker.format.line.color.rgb = WHITE
    # label the two points the argument rests on: p = 1/2 of a two-way tie, and n = 1
    half = next(i for i, (xv, _) in enumerate(pts[2]) if abs(xv - 0.5) < 1e-6)
    for se, i, yv, pos in ((ch.series[1], half, pts[2][half][1], XL_LABEL_POSITION.LEFT),
                           (ch.series[4], 0, pts[1][0][1], XL_LABEL_POSITION.BELOW)):
        dl = se.points[i].data_label
        dl.has_text_frame = True
        dl.text_frame.text = rd(yv, 3)
        dl.position = pos
        for r in dl.text_frame.paragraphs[0].runs:
            r.font.size, r.font.bold = Pt(13), True
            r.font.color.rgb = INK
        dLbl = dl._dLbl                  # a white backing, so the label never sits on a line
        sp = dLbl.makeelement(qn("c:spPr"), {})
        fill = sp.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": "FFFFFF"}))
        sp.append(fill)
        dLbl.find(qn("c:tx")).addnext(sp)
    # three numbers
    sx, sw = 8.3, W - M - 8.3
    for i, (v, lab, col) in enumerate([
            (rd(slope["est"], 2), "calibration slope (1 = calibrated)", TEAL),
            (num(skill["est"], 3), "skill over a constant, 2-way ties: CI spans 0", AMBER_TXT),
            (sg(dF1["est"], 4), f"macro F1 gain from p ≥ 0.5 (bar: {sesoi})", AMBER_TXT)]):
        stat(s, sx, BODY_TOP + 0.05 + i * 1.45, sw, v, lab, color=col, h=1.15)
    takeaway(s, "Use p to rank, not as a probability", sx, 5.95, sw)
    footer(s, "Controlled genus-blocked benchmark · exact-tie rule, full reference density · EC "
              "level · intervals in the notes")
    pooled = E["all|all=all"]["union"]          # p = 1: agreeing references vs a single one
    p1m, p1s = pooled["obs_p1_multi"], pooled["obs_p1_single"]
    three = ", ".join(rd(yv, 3) for _, yv in pts[3])
    big = mu["clusters"]["group"]["largest_share_points"]
    sk3, sk4 = (E[k]["union"]["bss_vs_const"]["est"] for k in ("n3|all=all", "n4+|all=all"))
    notes(s, "This asks whether p(function) = k/n behaves as a probability that the function is "
             "present in the organism the amplicon came from, scored against the truth genomes "
             "of the controlled genus-blocked benchmark (shipped exact-tie selector, full "
             f"reference density, EC level; the {mu['n_cells']:,} multi-genome selections among "
             f"{E['all|all=all']['n_cells']:,} scored cells). It ranks correctly: within "
             "selections of two, three and four genomes, observed presence rises strictly with "
             f"p (for three genomes {three} at p = 1/3, 2/3 and 1). It is overconfident: the "
             f"probability-scale calibration slope is {rd(slope['est'], 2)} "
             f"[{rd(slope['lo'], 2)}, {rd(slope['hi'], 2)}], a function at p = 1/2 is present "
             f"only {rd(pts[2][half][1], 3)} of the time, and p = 1 from two or more agreeing "
             f"references is present {rd(p1m['est'], 3)}, against {rd(p1s['est'], 3)} from a "
             "single reference. Over a constant confidence, p adds skill only for ties of "
             f"three or more genomes ({num(sk3, 3)} for three, {num(sk4, 3)} for more); for "
             f"two-way ties, {share2:.0%} of multi-genome cells, the Brier skill is "
             f"{num(skill['est'], 3)} [{num(skill['lo'], 3)}, {num(skill['hi'], 3)}], an "
             "interval that includes zero. Keeping only functions with p ≥ 0.5 changes the "
             f"benchmark's macro F1 by {sg(dF1['est'], 4)} [{rd(dF1['lo'], 4)}, "
             f"{rd(dF1['hi'], 4)}], far below the pre-registered {sesoi}. Caveats: the folds "
             "are genus-blocked, so every number is cross-genus transfer and pessimistic for "
             "deployment; intervals are cluster bootstraps over holdout groups, one of which "
             f"carries {big:.0%} of the multi-genome points; every join, invariant and leakage "
             "check in assertions.json passes. The revised manuscript draft reports this "
             "calibration.")
    return s


def s17b_c12_cause(prs, D):
    """Why C12 fails: five loop families, each closed by one MSGapfill reaction."""
    import csv
    with open(os.path.join(SR, "EGC_C12_gated_models.csv")) as fh:
        rows = list(csv.DictReader(fh))
    n_leak, n_genomes = len(rows), len({r["genome"] for r in rows})
    s = blank(prs)
    title(s, f"Why C12 {'failed' if D['c12_rescored'] else 'fails'}: ATP from nothing in "
             f"{n_leak} models",
          sub=("Pre-registered run · " if D["c12_rescored"] else "")
              + "each a closed loop that one gapfilled reaction completes")
    n_models = D["vts"]["baseline"]["n"]
    rev = [r for r in rows if r["cut_bounds"] == "[-1000.0, 0.0]"]     # added in reverse
    fams = {}
    for r in rows:
        k = r["msgapfill_cut"].replace("_c0", "")
        fams.setdefault(k, [0, r["cut_bounds"] == "[-1000.0, 0.0]", r["uphill_step"]])[0] += 1
    order = sorted(fams, key=lambda k: -fams[k][0])
    label = {k: f"{k} reversed" + (" + NadX" if "NadX" in fams[k][2] else "") if fams[k][1]
             else f"{k} + catalase" if "rxn00006" in fams[k][2] else k
             for k in order}
    hbar(s, M - 0.1, BODY_TOP - 0.05, 7.0, 4.9, [label[k] for k in order],
         [fams[k][0] for k in order], colors=[TEAL if fams[k][1] else DEEP for k in order],
         fmt="0", vmin=0, vmax=17, gap=50)

    rx = 7.85
    rw = W - M - rx
    stat(s, rx, BODY_TOP, rw, f"{n_leak} / {n_models}",
         f"rebuilt models leak · {n_genomes} genomes", color=AMBER_TXT, h=1.35)
    for i, (sw, head, sub) in enumerate([
            (DEEP, "Template direction",
             "catalase only 2 H₂O + O₂ → 2 H₂O₂"),
            (TEAL, "Negative reverse penalty",
             f"reverse made cheaper · {len(rev)} of {n_leak}"),
            (RGBColor(0xA7, 0xB8, 0xBC), "Fork drops reversals", "growth only, no leak")]):
        y = BODY_TOP + 1.65 + i * 1.17
        rect(s, rx, y, rw, 0.87, LIGHT, shape=MSO_SHAPE.RECTANGLE)
        rect(s, rx, y, 0.14, 0.87, sw, shape=MSO_SHAPE.RECTANGLE)
        tf = textbox(s, rx + 0.32, y + 0.06, rw - 0.45, 0.75, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, head, size=14.5, bold=True, color=INK, font=HEAD, first=True, space_after=1)
        para(tf, sub, size=12.5, color=MUTED, space_after=0)
    footer(s, f"Bars: the MSGapfill reaction whose removal ends each loop · {n_models} models, "
              f"{n_genomes} genomes")

    base = D["vts"]["baseline"]
    kcal = next(r["uphill_step"].split(" kcal/mol")[0].split()[-1].lstrip("+") for r in rows
                if "rxn00006" in r["uphill_step"])                 # the catalase step's dG
    n_truth = sum(r["arm"] == "truth" for r in rows)   # one truth build serves both pipelines
    n_rows = sum(D["phgf_pre"]["meta"]["n_gated_out"].values())
    assert n_rows == n_leak + n_truth, "leaking models no longer match C12's gated rows"
    notes(s, (
        f"C12 fails on gapfilled models only: rebuilding every gapfilled model of the "
        f"{n_genomes} affected genomes ({n_models} models, truth first, exactly as the pipeline "
        f"orders them) reproduces the stored rows, and {n_leak} of them make ATP on water and "
        f"protons alone: these are C12's {n_rows} gated rows, since each of the {n_truth} truth "
        "models is shared by both pipelines and counts twice; "
        + ("none makes biomass. " if all(float(r["max_empty_biomass"]) == 0 for r in rows)
           else "some also make biomass. ")
        + "Every leak is a closed internal cycle (zero net exchange, persisting with every "
        "exchange closed), and in each leaking model exactly one reaction added by the final "
        "MSGapfill step is a single cut that abolishes the loop: the bars count those cuts, "
        + ", ".join(f"{k} {fams[k][0]}" for k in order) + ". "
        f"The catalase loop accounts for {n_leak - len(rev)}: every template allows catalase "
        f"(rxn00006) only as 2 H2O + O2 to 2 H2O2, uphill by {kcal} kcal/mol, and rxn09180 "
        "(anaerobic protoporphyrinogen oxidase) closes it. "
        f"In the other {len(rev)} the culprit (rxn10966, rxn08167 with draft NadX, rxn00381, "
        "rxn00084) is reversible in the template and carries a negative reverse penalty, so "
        "the gapfilling objective prices the reverse below the forward direction and adds it "
        "reversed. "
        "Separately, the ModelSEEDpy fork this pipeline runs (MSBuilder.build_and_gapfill, "
        "inlined in phenotype_fold.gapfill_block) integrates the gapfiller's new reactions but "
        f"drops its reversals; that causes no leak but left {base['no_growth']} of the "
        f"{n_models} models unable to grow (upstream ModelSEEDpy applies them). "
        "The failure is a reconstruction-pipeline artefact that touches every arm, truth "
        "included, not a property of any method. "
        "Sources: region_validation/shared_reference/EGC_C12_REPORT.md and "
        "EGC_C12_gated_models.csv (column msgapfill_cut)."))
    return s


def s17c_c12_fix(prs, D):
    """The validated fix, what running it involves, and what it would change."""
    import csv
    n_fix = {3: "Three", 4: "Four", 5: "Five", 6: "Six"}.get(len(C12_FIX_RXNS),
                                                           str(len(C12_FIX_RXNS)))
    s = blank(prs)
    title(s, f"{n_fix} template fixes close every leak",
          sub=", ".join(C12_FIX_RXNS) + ": forward only")
    V = D["vts"]
    b0, br, fx = V["baseline"], V["base+rev"], V["PR1v2+rev"]
    commit = fx["templates"].rsplit("(", 1)[1].rstrip(")")      # "five-reaction fix (21a8084)"
    with open(os.path.join(SR, "EGC_C12_gated_models.csv")) as fh:
        n_genomes = len({r["genome"] for r in csv.DictReader(fh)})
    n_gf = D["phgf_pre"]["meta"]["n_candidates_eligible"]
    n_atp = D["phatp_pre"]["meta"]["n_candidates_eligible"]
    f1 = lambda k: rd(D["kg"]["arms"][k]["ec"]["population"]["f1"]["value"])   # noqa: E731
    table(s, M, BODY_TOP, 7.2, [
        [f"{fx['n']} models", "Shipped", "Reversals only", "Fix + reversals"],
        ["Leaking", f"{b0['leaking']}", f"{br['leaking']}", f"{fx['leaking']}"],
        ["New leaks", "n/a", f"{br['new_leaks']}", f"{fx['new_leaks']}"],
        ["Not growing", f"{b0['no_growth']}", f"{br['no_growth']}", f"{fx['no_growth']}"],
        ["Growth ratio", "n/a", "reference", f"{fx['passing_growth_median_ratio']:.3f}"]],
        col_w=[2.1, 1.5, 1.75, 1.85], size=15, row_h=0.56, emph_col=3)

    ox = 8.1                      # what passing the gate would mean, as one card
    card = rect(s, ox, BODY_TOP, W - M - ox, 2.8, DEEP)
    tf = card.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    done = D["c12_rescored"]                  # the amended run has been scored
    c12 = D["ctl"]["C12"]
    passed = done and c12["status"] == "pass"
    if done:
        gm = D["phgf"]["meta"]
        para(tf, "C12 passes" if passed else "C12 still fails", size=28, bold=True,
             color=MINT if passed else AMBER, font=HEAD, first=True, space_after=6)
        para(tf, "EC numbers unchanged", size=16, color=WHITE, space_after=4)
        para(tf, f"{sum(gm['n_gated_out'].values())} of {gm['n_models']:,} gapfilled rows "
                 f"gated, all {n_gf} genomes", size=16, color=WHITE, space_after=0)
    else:
        para(tf, "C12 should pass", size=28, bold=True, color=MINT, font=HEAD, first=True,
             space_after=6)
        para(tf, "EC numbers unchanged", size=16, color=WHITE, space_after=4)
        para(tf, f"{n_gf - n_genomes} gapfilled genomes not yet tested", size=16, color=AMBER,
             space_after=0)

    flow(s, M, BODY_TOP + 3.3, W - 2 * M, 1.4, [
        ("Amend SPEC", "post hoc, disclosed"),
        ("ATP pass", f"{n_atp:,} genomes · " + ("unchanged" if done else C12_RUN_HOURS['atp'])),
        ("Gapfilled pass", f"{n_gf} genomes · " + (
            f"{sum(D['phgf']['meta']['n_gated_out'].values())} gated" if done
            else C12_RUN_HOURS['gf'])),
        ("Rescore", ("C11, C12 pass · valid" if passed else "C11, C12, folds"))],
        size=17, gap=0.4)
    footer(s, f"ModelSEEDTemplates {commit}, validated first on the {n_genomes} affected "
              "genomes · " + (("rerun 2026-10-29 on ModelSEEDpy c2e31a0 (A-C13)"
                               if D["amend_label"] == "A-C13" else
                               "run 2026-10-27/28, ModelSEEDpy pinned at the pre-registered commit")
                              if done else "run not started; needs the authors' go-ahead"))

    outcome = ""
    if done:
        cg = D["c12cmp"]["all"]
        g_after, g_before = cg["gate"]["n_gated_after"], cg["gate"]["n_gated_before"]
        n_new = sum(cg["gate"]["new_leaks"].values())
        ng_b, ng_a = cg["growth"]["non_growing_before"], cg["growth"]["non_growing_after"]
        outcome = (
            "The amendment has since run (SPEC section 14, with ModelSEEDpy pinned to the "
            "commit the pre-registered run used, after the first launch was found importing a "
            "moved editable install)"
            + (", and amendment A-C13 (section 15) reran it on ModelSEEDpy c2e31a0, which "
               "fixes the builder's demand reactions and an ATP-correction step; the figures "
               "below are A-C13's" if D["amend_label"] == "A-C13" else "")
            + ". The ATP pass on the corrected templates reproduced the "
            "pre-registered pass row for row, so its readouts did not move. The gapfilled pass "
            f"gated {g_after} of {cg['n_rows']:,} rows against {g_before}, with {n_new} new "
            f"leaks in all {n_gf} genomes, the {n_gf - n_genomes} outside the validation "
            f"included, and {ng_a} non-growing models against {ng_b}; rescored, C11 and C12 "
            f"{'pass' if passed else 'do not both pass'} and the leave-genus-out artefact is "
            f"{'valid' if D['kg']['headline'].get('valid') else 'still invalid'}, with every EC "
            f"number unchanged ({f1('prfba_nestedcv')} against {f1('picrust2')}). ")
    notes(s, (
        f"The fix restricts {n_fix.lower()} template reactions to their physiological direction "
        f"({', '.join(C12_FIX_RXNS)}; ModelSEEDTemplates {commit}) and applies the reversals "
        "the gapfiller proposes, which the pipeline's ModelSEEDpy fork drops. "
        f"Rebuilt from the patched template files over all {fx['n']} gapfilled models of the "
        f"{n_genomes} affected genomes, it takes the leaking models from {br['leaking']} to "
        f"{fx['leaking']} with {fx['new_leaks']} new leaks; every model grows "
        f"({fx['no_growth']} non-growers, against {b0['no_growth']} as shipped, all of them "
        "dropped reversals), and the models that already passed keep their growth (median "
        f"ratio {fx['passing_growth_median_ratio']:.3f} against the original templates with "
        "reversals applied). "
        + (outcome if done else (
        "Adopting it changes pre-registered pipeline settings, so it needs a SPEC amendment "
        "that records it was decided after C12 failed, and both phenotype passes must be "
        f"re-run on one template version: the ATP pass ({n_atp:,} genomes, "
        f"{D['phatp_pre']['meta']['n_models']:,} models, no gapfilling, about "
        f"{C12_RUN_HOURS['atp']}), because the change also alters its un-gapfilled drafts, "
        f"and the gapfilled pass ({n_gf} genomes, {D['phgf_pre']['meta']['n_models']:,} models, "
        f"about {C12_RUN_HOURS['gf']}), then the C11 identity control and the fold rescoring. "
        f"C12 should then pass, since all {D['phgf_pre']['meta']['n_gated_out']['truth']} gated "
        "truth rows sit in the validated genomes, "
        "and the leave-genus-out artefact becomes valid if every other hard control holds, "
        f"with its EC numbers unchanged ({f1('prfba_nestedcv')} against {f1('picrust2')}). "
        "The phenotype readouts move: small shifts in the ATP readout through the drafts, "
        f"larger ones in the gapfilled readouts as the {n_genomes} genomes return to every "
        "arm's denominator and growth counts the reversals; the headline null is unlikely to "
        "flip but must be re-checked. "
        f"The risk is the other {n_gf - n_genomes} gapfilled genomes, which were not part of "
        "the validation: a new leak there would leave C12 failing. ")) +
        "Sources: region_validation/shared_reference/C12_FIX_PLAN.md, C12_SCOPE.md and data/"
        "template_direction_audit/validation_summary.json (entries baseline, base+rev and "
        "PR1v2+rev)."))
    return s


def s14b_rule(prs, D):
    """The selection-rule decision that depth 100 forces: union vs legacy, side by side."""
    s = blank(prs)
    background(s, DEEP)
    title(s, "One rule throughout: the exact-tie union", dark=True,
          sub="Decided 2026-10-27 · the legacy reducer stays a user option")
    regs = list(REG_LABEL)

    def kept(tree):
        a = [D["taxd"][tree]["regions"][r]["self_recovery"]["include"]["all"]["all"]
             for r in regs]
        return sum(x["self_selected"] for x in a), sum(x["self_in_top20"] for x in a)

    def v4(tree):
        return D["capd"][tree]["regions"]["V4"]["include"]["macro"]["precision"]["value"]

    c = D["cmpd"]["d100"]["by_region"]["ec"]["all"]["all"]

    def gap(key):
        return sg(c[key]["population"]["f1"]["value"]
                  - c["picrust2"]["population"]["f1"]["value"])

    def sig(rule):
        lo, hi = GAP_CI_D100[rule]
        return "significant" if hi < 0 or lo > 0 else "n.s."

    folds = [f for blk in (D["kg"], D["kf"]) for k, f in sorted(blk["nested_cv"].items())
             if k.startswith("fold_")]
    wins = sum(f["inner_winner"] == "legacy" for f in folds)
    marg = [f["inner_f1"]["legacy"] - f["inner_f1"]["exact_tie"] for f in folds]
    ku, kl = kept("U100"), kept("L100")
    tbl = table(s, M, BODY_TOP + 0.05, W - 2 * M, [
        ["Original panel, depth 100", "Exact-tie union", "Legacy reducer"],
        ["Source kept when in the record", f"{rd(100 * ku[0] / ku[1], 1)}%",
         f"{rd(100 * kl[0] / kl[1], 1)}%"],
        ["Population gap to PICRUSt2", f"{gap('prfba_exclude')} ({sig('union')})",
         f"{gap('prfba_exclude_legacy')} ({sig('legacy')})"],
        ["Controlled benchmark, folds won", f"{len(folds) - wins} of {len(folds)}",
         f"{wins} of {len(folds)} (by {rd(min(marg))}–{rd(max(marg))})"],
    ], col_w=[5.0, 3.5, 3.63], size=19, row_h=0.9, dark=True, header_fill=TEAL)
    # on the dark slide, draw the cell edges in the background colour (as the light
    # slides' white edges do) rather than the table style's white grid
    from lxml import etree
    from pptx.oxml.ns import qn
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                for old in tcPr.findall(qn(tag)):
                    tcPr.remove(old)
                ln = etree.Element(qn(tag), w=str(Pt(1.5)), cap="flat", cmpd="sng",
                                   algn="ctr")
                fill = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(fill, qn("a:srgbClr"), val=str(DEEP))
                tcPr.insert(i, ln)
    tf = takeaway(s, "Broader and less precise, but what users run", M, BODY_TOP + 4.25,
                  W - 2 * M, dark=True, h=0.5)
    r = tf.paragraphs[0].add_run()             # the lean, on the same line, in amber
    r.text = "  ·  legacy stays an option"
    r.font.size, r.font.bold, r.font.name = Pt(18), True, HEAD
    r.font.color.rgb = AMBER
    footer(s, "Fold choice: leave-one-fold-out population EC F1 on the shared-reference "
              "benchmark, 5 folds × genus and family holdout", dark=True)
    sel = INCL_N_SELECTED
    ci = {r: f"{sg(lo)} to {sg(hi)}" for r, (lo, hi) in GAP_CI_D100.items()}
    word = {"significant": "significant", "n.s.": "not significant"}
    notes(s,
          f"Depth 100 makes the choice of selection rule matter. The shipped exact-tie union "
          f"keeps every reference tied at the best identity, so among the self-inclusion "
          f"amplicons whose source genome is in the 100-hit record it keeps the source in "
          f"{rd(100 * ku[0] / ku[1], 1)}% ({ku[0]:,} of {ku[1]:,}); the earlier ordered "
          f"(legacy) reducer keeps it in {rd(100 * kl[0] / kl[1], 1)}% and discards it in "
          f"{kl[1] - kl[0]:,}. The price shows when the source is present: the union grows to "
          f"{sel['union']:.2f} genomes per answer against {sel['legacy']:.2f}, and V4 "
          f"gene-capture precision falls to {rd(v4('U100'))} against {rd(v4('L100'))}. On the "
          f"as-shipped PICRUSt2 axis the population EC F1 gap is {gap('prfba_exclude')} under "
          f"the union (95% genome-clustered bootstrap CI {ci['union']}, "
          f"{word[sig('union')]}) and {gap('prfba_exclude_legacy')} under the legacy reducer "
          f"(CI {ci['legacy']}, {word[sig('legacy')]}). In the controlled shared-reference "
          f"benchmark, leave-one-fold-out selection picked legacy over the exact-tie union in "
          f"{wins} of {len(folds)} folds (genus and family holdouts), by {rd(min(marg))} to "
          f"{rd(max(marg))} F1, an order of magnitude below the pre-registered 0.02 smallest "
          f"effect of interest; the genus summary marks itself invalid because the C12 "
          f"gapfill validity gate failed, which does not touch these EC fold scores. On 2026-10-27 "
          f"the authors chose the union for the whole paper, because it is the rule users run, "
          f"it keeps the source when the source is there, and legacy's edge in the controlled "
          f"benchmark is far below the smallest effect of interest; its cost is the larger "
          f"self-inclusion union and lower short-amplicon precision. Until then the manuscript "
          f"used each rule where it was like for like (legacy for tab:main, McNemar, MAG "
          f"against 16S and the cross-panel numbers). The legacy reducer stays an option users "
          f"can select, and the controlled benchmark keeps its pre-registered arm, which "
          f"leave-one-fold-out selection set to legacy, beside the exact-tie arm. "
          f"Sources: taxacc_region_summary.json and gene_capture_by_region.json in "
          f"region_validation_d100 and region_validation_d100_legacy; gene_capture_cells.csv "
          f"(n_selected); comparison_summary.json and D100_VS_D20.md section 2D; "
          f"shared_reference summary_genus.json and summary_family.json (nested_cv).")
    return s


def s17d_template_audit(prs, D):
    """Every v7.0 template direction against the ModelSEED database: curate, do not align."""
    import collections
    import csv
    s = blank(prs)
    ad = os.path.join(SR, "data", "template_direction_audit")
    with open(os.path.join(ad, "tierB_readings.tsv")) as fh:
        reads = collections.Counter(r["reading"] for r in csv.DictReader(fh, delimiter="\t"))
    n_opp = sum(reads.values())
    aligned = collections.defaultdict(set)
    with open(os.path.join(ad, "direction_audit.tsv")) as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            if r["template"] != "Core-V6.json" and r["kind"] in ("template_opposite",
                                                                 "template_reversible"):
                aligned[r["template"]].add(r["reaction"].rsplit("_", 1)[0])
    n_al = sorted(len(v) for v in aligned.values())
    with open(os.path.join(ad, "validation_per_model.csv")) as fh:
        n_gen = len({r["model"].split("__")[0] for r in csv.DictReader(fh)})
    with open(os.path.join(ad, "alignment_growth_probe.json")) as fh:
        arch = {m["model"].split("__")[0] for m in json.load(fh) if "Archaea" in m["template"]}
    title(s, f"Only {reads['template error']} of {n_opp} direction conflicts are template "
             "errors",
          sub=f"{n_opp} reactions allowed only in the direction the database forbids")
    bars = [("Database error", "database error", TEAL),
            ("Near equilibrium", "reversible or near equilibrium", SEAFOAM),
            ("Template error", "template error", AMBER),
            ("Mis-specified", "mis-specified", MUTED),
            ("Unresolved", "not resolved", MUTED)]
    ch = hbar(s, M, BODY_TOP, 5.5, 4.95, [b[0] for b in bars], [reads[b[1]] for b in bars],
              colors=[b[2] for b in bars], fmt="0", gap=45)
    v = D["tdv"]
    cur = [v[k] for k in ("PR1v2+rev", "Bcur2+rev", "PENcur2+rev")]
    full, base = v["BC_full+rev"], v["base+rev"]
    mv = sorted(x["passing_changed_gt10pct"] for x in cur)
    cur_ratio = min(x["passing_growth_median_ratio"] for x in cur)
    lost = full["lost_growth_genomes"]
    who = ("one archaeon" if len(lost) == 1 and lost[0] in arch
           else f"{len(lost)} genome{'' if len(lost) == 1 else 's'}")
    table(s, 6.4, BODY_TOP + 0.45, W - M - 6.4, [
        [f"{full['n']} gapfilled models", "Curated PRs", "Database-aligned"],
        ["Leaking models", f"{base['leaking']} → {max(x['leaking'] for x in cur)}",
         f"{base['leaking']} → {full['leaking']}"],
        ["Cannot grow", f"{max(x['no_growth'] for x in cur)}", f"{full['no_growth']} ({who})"],
        ["Median growth ratio", f"{cur_ratio:.2f}",
         f"{full['passing_growth_median_ratio_still_growing']:.2f}"],
    ], col_w=[2.45, 1.75, 2.13], size=15, row_h=0.7, emph_col=1)
    takeaway(s, "Curated PRs submitted · issue #10 filed", 6.4, BODY_TOP + 3.75,
             W - M - 6.4, h=0.8)
    footer(s, f"{full['n']} gapfilled models of {n_gen} genomes, gapfiller reversals applied · "
              "growth relative to the original templates")
    notes(s,
          f"We compared every reaction direction in the v7.0 ModelSEED templates with the "
          f"ModelSEED database: {n_opp} reactions per genome-scale template are allowed only in "
          f"the direction the database forbids, and {n_al[0]:,}–{n_al[-1]:,} differ from it "
          f"once the reactions the templates leave reversible are counted. Read one by one "
          f"against known physiology, only {reads['template error']} of the {n_opp} are "
          f"template errors; for {reads['database error']} the database direction is the wrong "
          f"one, {reads['reversible or near equilibrium']} are reversible or near equilibrium, "
          f"{reads['mis-specified']} are mis-specified reactions to remove or block, and "
          f"{reads['not resolved']} are unresolved. The curated set is the five-reaction fix "
          f"that closes the C12 leaks ({', '.join(FIVE_RXN_FIX)}), the "
          f"{reads['template error']} template errors, and an optional reverse-penalty sign "
          f"clean-up, {len(FIVE_RXN_FIX) + reads['template error']} direction changes in all "
          f"against {n_al[0]:,}–{n_al[-1]:,} per template for the database-aligned set; on the {full['n']} gapfilled models of the {n_gen} affected genomes it "
          f"removes all {base['leaking']} leaks with no new one, every model grows, the median "
          f"growth ratio is {cur_ratio:.3f}, and {mv[0]}–{mv[-1]} of the "
          f"{full['n_passing_growing']} passing models move by more than 10% because the "
          f"gapfill MILP lands on another near-equal optimum. Applying the database directions "
          f"wholesale also closes every leak, but archaeon {lost[0]} then has no gapfill "
          f"solution, so all {full['no_growth']} of its models stop growing "
          f"({full['gated_now_ok_via_infeasible_gapfill']} of the {base['leaking']} leaking "
          f"models count as fixed only because gapfilling failed); the models that still grow "
          f"keep a median {full['passing_growth_median_ratio_still_growing']:.3f} of their "
          f"growth ({full['passing_growth_median_ratio']:.3f} counting the non-growers as "
          f"zero), and {full['passing_changed_gt10pct']} of {full['n_passing_growing']} move "
          f"by more than 10%. So the database directions must not be applied wholesale; the "
          f"{reads['database error']} database errors and "
          f"{reads['reversible or near equilibrium']} near-equilibrium entries go to "
          f"ModelSEEDDatabase as an issue instead. The branches are pushed to "
          f"freiburgermsu/ModelSEEDTemplates, and the pull requests and the audit issue "
          f"(ModelSEED/ModelSEEDTemplates#10) were filed on 2026-10-27 from the texts in "
          f"region_validation/shared_reference/ (SUBMIT_ModelSEEDTemplates.md). Sources: "
          f"data/template_direction_audit/tierB_readings.tsv, "
          f"direction_audit.tsv, validation_per_model.csv and validation_summary.json (entries "
          f"PR1v2+rev, Bcur2+rev, PENcur2+rev, BC_full+rev, base+rev).")
    return s


def s18b_accession_map(prs, D):
    s = blank(prs)
    am = D["accmap"]
    ov, mt, pcache = am["coverage"]["overall"], am["match_type"], am["chosen_in_pgfam_cache"]
    real = am["coverage"]["mag_selections"]["real_mags_gtdb_r226"]
    title(s, "A GTDB-to-BV-BRC map would let GTDB predict function",
          sub="GTDB hits: GCA/GCF only; genes need BV-BRC IDs")
    n, hit = ov["n"], ov["resolved"]
    flow(s, M, BODY_TOP + 0.15, W - 2 * M, 1.4, [
        ("GTDB hit", "GCA/GCF: classify only"),
        ("BV-BRC ID", f"new map: {100 * hit / n:.1f}%"),
        ("PGFams", f"{pcache['nonempty']:,} cached; fetch rest"),
        ("Model", "gene union → FBA")], size=18)
    other = mt["wildcard_core"] + mt["metadata_gca"] + mt["taxid_length"]
    rows = [["Linked by", "Representatives", "Share"]]
    for lab, k in (("Exact accession", mt["exact"]), ("GCF ↔ GCA swap", mt["prefix_swap"]),
                   ("BioSample, other", mt["biosample"] + other), ("Unmapped", n - hit)):
        rows.append([lab, f"{k:,}", f"{100 * k / n:.1f}%"])
    fb = (W - 2 * M - 3 * 0.42) / 4               # the flow's box width: align to its grid
    table(s, M, BODY_TOP + 2.0, 2 * fb + 0.42, rows, col_w=[2.5, 1.9, 2 * fb + 0.42 - 4.4],
          size=14, row_h=0.5, emph_col=1)
    x2 = M + 2 * (fb + 0.42)
    w2 = W - M - x2
    stat(s, x2, BODY_TOP + 2.0, w2, f"{real['mags_anchor_mapped']} / {real['n_with_selection']}",
         "real-MAG GTDB picks already mapped", h=1.4)
    takeaway(s, "Deferred to future work: it changes GTDB picks", x2, BODY_TOP + 3.62,
             w2, color=AMBER_TXT, h=0.8)
    footer(s, f"GTDB r226: {n:,} species representatives  ·  map built "
              f"{am['built_at'][:10]} (reference_maps/gtdb_r226_acc_to_bvbrc.tsv.gz)")
    lo, hi = ov["wilson95"]
    notes(s, (
        f"The MAG arm's skani search runs against the GTDB r226 sketch of {n:,} species "
        "representatives as readily as against BV-BRC genomes, but a GTDB hit carries a "
        "GCA/GCF assembly accession, whereas PGFam "
        "gene content — and with it the synthetic-genome merge and the model — needs a "
        "BV-BRC genome ID, so a MAG searched against GTDB can today be classified and "
        "assessed for novelty but not given a predicted metabolism. "
        f"The map built on {am['built_at'][:10]} links {hit:,} of the {n:,} representatives "
        f"({100 * hit / n:.2f}%, Wilson 95% interval {lo:.2f}–{hi:.2f}%) to a BV-BRC "
        f"genome: {mt['exact']:,} by exact accession, {mt['prefix_swap']:,} by swapping the "
        f"GCF/GCA prefix of the same assembly, {mt['biosample']:,} through the BioSample and "
        f"{other:,} through looser routes; the {n - hit} left unmapped are "
        f"{ov['length_rejected']} rejected on genome length, {ov['biosample_rejected']} on "
        f"BioSample and {ov['deprecated_only']} with only deprecated records. "
        f"Of the {real['n_mags']} real MAGs, all {real['n_with_selection']} that have a GTDB "
        "selection already have their anchor reference mapped. "
        f"Integration still needs a PGFam fetch: only {pcache['nonempty']:,} of the mapped "
        f"genomes have non-empty cached PGFams ({pcache['empty']} are cached empty and "
        f"{pcache['absent']:,} are not cached). "
        "Wiring it in also changes the reported GTDB selections — numeric BV-BRC IDs "
        "escape the selector's MAG check, and an unmapped anchor changes the taxonomic call "
        "— so the integration, and any revision of the manuscript's GTDB results, is the "
        "authors' decision rather than a silent fix."))
    return s


def main():
    global CURRENT_SLIDE
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", default=None,
                    help="default presentations/prFBA_updates_2026-10.pptx "
                         "(prFBA_updates_2026-10_d20.pptx at --depth 20)")
    ap.add_argument("--depth", type=int, default=100, choices=[20, 100],
                    help="amplicon record depth whose results the deck reports (default 100, "
                         "the standard since the depth-100 rerun; 20 rebuilds the earlier deck)")
    ap.add_argument("--rule", default="union", choices=["legacy", "union"],
                    help="--depth 100 only: selection rule of the original-panel and "
                         "three-panel inputs (default union, the exact-tie rule the paper "
                         "uses throughout since 2026-10-27; legacy is the optional reducer)")
    ap.add_argument("--constants", default=None,
                    help="--depth 100 only: JSON with SELF_TOP20/SELF_SAT/SELF_CRED/EMPTIED")
    args = ap.parse_args()
    configure_depth(args.depth, args.rule, args.constants)
    if args.out is None:
        args.out = os.path.join(REPO, "presentations", "prFBA_updates_2026-10%s.pptx"
                                % ("" if args.depth == 100 else "_d%d" % args.depth))
    D = load()
    cm = D["ctl"]["C12"].get("measured") or {}
    D["c12_rescored"] = bool(D["c12fix"] and ("phenotype_gf" + D["amend_tag"] + "/") in
                             (cm.get("binding") or {}).get("phenotype_summary", ""))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for fn in (s01_title, s02_problem, s03_pipeline, s04_alignment, s05_selection,
               s06_synthetic, s06b_weighting, s06c_calibration, s07_validation, s08_regions,
               s09_denominators, s10_benchmark_design, s11_results, s12_novelty, s13_ceiling,
               s13b_dialect, s15_kfold_design, s16_kfold_result, s17_phenotype, s17b_c12_cause,
               s17c_c12_fix, s17d_template_audit, s14_depth, s14b_rule, s19_mag_why,
               s20_mag_build, s16_mag_arm, s17_crossover, s18_abstain, s18b_accession_map,
               s21_random_panel, s20_panel_space, s22_three_panel, s19_two_panel, s23_reversal,
               s24_audit, s15_next):
        CURRENT_SLIDE = fn.__name__
        _STATS.clear()
        fn(prs, D)
        if _STATS:                      # one value size for all of a slide's stat cards
            size = min(v for _, v in _STATS)
            for run, _ in _STATS:
                run.font.size = Pt(size)
    prs.core_properties.title = "prFBA project update"
    prs.core_properties.author = "Andrew Freiburger"
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    prs.save(args.out)
    print(f"[make_deck] {len(prs.slides.__iter__.__self__._sldIdLst)} slides -> {args.out}")


if __name__ == "__main__":
    main()
