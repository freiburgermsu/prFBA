#!/usr/bin/env python
"""
md_to_docx.py — render a Markdown manuscript to a clean academic .docx (python-docx).
Handles: # title, ##/###/#### headings, paragraphs, **bold**/*italic* inline, bullet
('- '/'* ') and numbered ('1. ') lists, and pipe tables (| a | b | with a |---| rule).

    md_to_docx.py INPUT.md OUTPUT.docx [--title "Override Title"]

Interpreter: ~/Documents/py_venv/bin/python
"""
from __future__ import annotations
import argparse
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITAL = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_CODE = re.compile(r"`([^`]+)`")


def _add_runs(paragraph, text):
    """Add inline-formatted runs (**bold**, *italic*, `code`) to a paragraph."""
    # tokenize on the three inline markers, preserving order
    pos = 0
    pattern = re.compile(r"\*\*(.+?)\*\*|(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)|`([^`]+)`")
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            paragraph.add_run(m.group(1)).bold = True
        elif m.group(2) is not None:
            paragraph.add_run(m.group(2)).italic = True
        else:
            r = paragraph.add_run(m.group(3)); r.font.name = "Consolas"; r.font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _is_table_row(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|")


def _is_table_rule(line):
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", line)) and "-" in line


def _cells(line):
    s = line.strip().strip("|")
    return [c.strip() for c in s.split("|")]


def render(md_path, docx_path, title_override=None):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(11)

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # --- tables ---
        if _is_table_row(line) and i + 1 < len(lines) and _is_table_rule(lines[i + 1]):
            header = _cells(line)
            i += 2
            rows = []
            while i < len(lines) and _is_table_row(lines[i]):
                rows.append(_cells(lines[i])); i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                c = t.rows[0].cells[j]
                c.paragraphs[0].text = ""
                _add_runs(c.paragraphs[0], h)
                for r in c.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cs = t.add_row().cells
                for j, val in enumerate(row[:len(header)]):
                    cs[j].paragraphs[0].text = ""
                    _add_runs(cs[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        # --- headings ---
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1)); text = m.group(2).strip()
            if level == 1 and not title_done:
                h = doc.add_heading("", level=0)
                _add_runs(h, title_override or text)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_done = True
            else:
                h = doc.add_heading("", level=min(level - 1, 4) if title_done else min(level, 4))
                _add_runs(h, text)
            i += 1
            continue

        # --- lists ---
        mb = re.match(r"^\s*[-*]\s+(.*)$", line)
        mn = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if mb:
            _add_runs(doc.add_paragraph(style="List Bullet"), mb.group(1)); i += 1; continue
        if mn:
            _add_runs(doc.add_paragraph(style="List Number"), mn.group(1)); i += 1; continue

        # --- horizontal rule / blank ---
        if re.match(r"^\s*([-*_])\1{2,}\s*$", line):
            i += 1; continue
        if not s:
            i += 1; continue

        # --- paragraph (gather wrapped lines until blank) ---
        para = [s]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|\s*[-*]\s|\s*\d+\.\s|\|)", lines[i]):
            para.append(lines[i].strip()); i += 1
        _add_runs(doc.add_paragraph(), " ".join(para))

    doc.save(docx_path)
    wc = len(open(md_path, encoding="utf-8").read().split())
    print(f"wrote {docx_path} (~{wc} source words)")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("md"); ap.add_argument("docx"); ap.add_argument("--title", default=None)
    a = ap.parse_args(argv)
    render(a.md, a.docx, a.title)


if __name__ == "__main__":
    sys.exit(main())
